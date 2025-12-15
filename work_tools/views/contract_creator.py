"""
合同创建人修改视图模块
修改tphct01表的创建人信息：BPO_EDIT_PERSON, CREATE_USER, CREATE_USER_NAME
"""
from django.shortcuts import render
from django.http import JsonResponse, FileResponse
from ..forms import ContractCreatorUpdateForm
from ..models import UserOrgDetail
from ..logger_utils import log_view_input
from .base import parse_ops_remark, save_sql_file
from ..validation_utils import (
    validate_required,
    generate_validation_failure_excel,
    get_temp_filename_from_path
)
import io
import logging

try:
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

logger = logging.getLogger(__name__)


def validate_contract_creator_records(records):
    """校验合同创建人修改记录"""
    results = []
    passed_count = 0
    failed_count = 0
    
    for idx, record in enumerate(records):
        row_number = idx + 2  # 跳过表头
        errors = []
        
        # 必填项校验：合同号必填
        error = validate_required(record.get('bpo_id'), '合同号')
        if error:
            errors.append(error)
        
        # 必填项校验：创建人账号必填
        error = validate_required(record.get('creator_account'), '创建人账号')
        if error:
            errors.append(error)
        
        # 数据有效性校验：如果查询不到创建人信息，记录错误
        if record.get('creator_account'):
            user_org = UserOrgDetail.objects.filter(login_name=record['creator_account']).first()
            if not user_org:
                errors.append(f"创建人账号查询失败，请检查账号是否正确")
            else:
                # 将查询结果存储到record中供后续使用
                record['creator_name'] = user_org.user_name or ''
        
        # 记录结果
        if errors:
            results.append({
                'row_number': row_number,
                'valid': False,
                'errors': errors
            })
            failed_count += 1
        else:
            results.append({
                'row_number': row_number,
                'valid': True,
                'errors': []
            })
            passed_count += 1
    
    return {
        'valid': failed_count == 0,
        'total': len(records),
        'passed': passed_count,
        'failed': failed_count,
        'results': results
    }


def _process_excel_import(excel_file, dynamic_id, ops_remark):
    """处理Excel批量导入"""
    if not OPENPYXL_AVAILABLE:
        raise ValueError('openpyxl未安装，无法处理Excel文件')
    
    wb = openpyxl.load_workbook(excel_file, data_only=True)
    ws = wb.active
    
    # 读取表头
    headers_row = [str(cell.value or '').strip() for cell in ws[1]]
    headers_lower = {h.lower(): h for h in headers_row if h}
    
    # 列名映射（忽略大小写）
    col_map = {}
    for target_key, possible_names in {
        'bpo_id': ['合同号', 'bpo_id', '合同ID', '合同编号'],
        'creator_account': ['创建人账号', 'creator_account', 'bpo_edit_person', 'create_user'],
        'orig_creator_account': ['原创建人账号', 'orig_creator_account', '原账号'],
        'orig_creator_name': ['原创建人姓名', 'orig_creator_name', '原姓名'],
    }.items():
        for name in possible_names:
            if name.lower() in headers_lower:
                col_map[target_key] = headers_lower[name.lower()]
                break
    
    records = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        row_dict = dict(zip(headers_row, row))
        
        bpo_id = str(row_dict.get(col_map.get('bpo_id', ''), '')).strip()
        creator_account = str(row_dict.get(col_map.get('creator_account', ''), '')).strip()
        
        # 跳过空行
        if not bpo_id and not creator_account:
            continue
        
        records.append({
            'bpo_id': bpo_id,
            'creator_account': creator_account,
            'orig_creator_account': str(row_dict.get(col_map.get('orig_creator_account', ''), '')).strip(),
            'orig_creator_name': str(row_dict.get(col_map.get('orig_creator_name', ''), '')).strip(),
        })
    
    if not records:
        raise ValueError('Excel中没有有效数据行')
    
    # 执行数据校验
    validation_result = validate_contract_creator_records(records)
    
    if not validation_result['valid']:
        # 校验失败，生成失败文件
        temp_file = generate_validation_failure_excel(
            excel_file,
            validation_result,
            'contract_creator'
        )
        
        if temp_file:
            return {
                'validation_failure': {
                    'total': validation_result['total'],
                    'passed': validation_result['passed'],
                    'failed': validation_result['failed'],
                    'filename': get_temp_filename_from_path(temp_file)
                }
            }
    
    # 校验通过，生成SQL
    sqls, rollback_sqls = _generate_sqls(records, dynamic_id, ops_remark)
    
    if sqls or rollback_sqls:
        saved_file = _save_and_download(sqls, rollback_sqls, dynamic_id, ops_remark)
        return {'saved_file': saved_file}
    else:
        raise ValueError('未生成任何SQL语句，请检查数据')


def _generate_sqls(data, dynamic_id, ops_remark):
    """生成SQL语句"""
    from ..config import get_config
    
    sqls = []
    rollback_sqls = []
    
    cfg = get_config()
    
    for item in data:
        bpo_id = item.get('bpo_id')
        creator_account = item.get('creator_account')
        orig_creator_account = item.get('orig_creator_account', '')
        orig_creator_name = item.get('orig_creator_name', '')
        
        # 使用校验阶段查询到的用户信息
        creator_name = item.get('creator_name', '')
        
        # 执行SQL - 更新合同主表
        sql = (f"update TPHCT01 set "
               f"BPO_EDIT_PERSON='{creator_account}',"
               f"CREATE_USER='{creator_account}',"
               f"CREATE_USER_NAME='{creator_name}',"
               f"ops_remark='{ops_remark}' "
               f"WHERE BPO_ID='{bpo_id}';")
        sqls.append(sql)
        
        # 生成回退SQL（支持补充创建人场景：原本为空则回退时清空）
        if not orig_creator_account:
            # 补充创建人场景：原本为空，回退时清空
            rollback_sql = (f"update tphct01 set "
                           f"create_user='',"
                           f"CREATE_USER_NAME='',"
                           f"BPO_EDIT_PERSON='',"
                           f"ops_remark='' "
                           f"where bpo_id='{bpo_id}';")
            rollback_sqls.append(rollback_sql)
        else:
            # 修改创建人场景：原本有值，回退到原值
            # 查询原创建人信息（如果账号无效，使用手动补充的姓名或空值）
            orig_user_org = UserOrgDetail.objects.filter(login_name=orig_creator_account).first()
            
            if orig_user_org:
                # 账号有效，使用查询的姓名
                final_orig_name = orig_user_org.user_name or ''
            else:
                # 账号无效，使用手动补充的姓名（如果有），否则为空
                final_orig_name = orig_creator_name or ''
            
            rollback_sql = (f"update tphct01 set "
                           f"create_user='{orig_creator_account}',"
                           f"CREATE_USER_NAME='{final_orig_name}',"
                           f"BPO_EDIT_PERSON='{orig_creator_account}',"
                           f"ops_remark='' "
                           f"where bpo_id='{bpo_id}';")
            rollback_sqls.append(rollback_sql)
    
    return sqls, rollback_sqls


def _save_and_download(sqls, rollback_sqls, dynamic_id, ops_remark_raw):
    """保存SQL文件到固定目录（静默输出）"""
    
    # 生成SQL内容（执行语句 + 回退语句合并在一个文件中）
    sql_content = "1、执行语句\n\n"
    
    for sql in sqls:
        sql_content += sql + '\n'
    
    sql_content += "\n\n\n\n\n\n2、回退语句\n\n"
    
    for sql in rollback_sqls:
        sql_content += sql + '\n'
    
    # 添加数据库信息
    sql_content += "\n\n3.数据库\n\nip：192.168.11.71\n\n库名：cnnc_ph"
    
    # 保存文件到固定目录
    filepath = save_sql_file(sql_content, '合同创建人修改', dynamic_id)
    
    return filepath


def contract_creator_update_view(request):
    """合同创建人修改视图"""
    saved_file = None
    validation_failure = None
    
    # 处理清除请求
    if request.method == 'GET' and request.GET.get('clear') == '1':
        form = ContractCreatorUpdateForm()
    elif request.method == 'POST':
        form = ContractCreatorUpdateForm(request.POST, request.FILES)
        if form.is_valid():
            try:
                cd = form.cleaned_data
                dynamic_id = cd.get('dynamic_id')
                ops_remark_raw = cd.get('ops_remark', '')
                ops_remark = parse_ops_remark(ops_remark_raw)
                
                # Excel批量导入
                if cd.get('excel_file'):
                    result = _process_excel_import(
                        cd['excel_file'], dynamic_id, ops_remark)
                    
                    if 'validation_failure' in result:
                        validation_failure = result['validation_failure']
                    elif 'saved_file' in result:
                        saved_file = result['saved_file']
                else:
                    # 单条修改
                    data = [{
                        'bpo_id': cd.get('bpo_id'),
                        'creator_account': cd.get('creator_account'),
                        'orig_creator_account': cd.get('orig_creator_account', ''),
                        'orig_creator_name': cd.get('orig_creator_name', ''),
                    }]
                    
                    # 执行数据校验
                    validation_result = validate_contract_creator_records(data)
                    
                    if not validation_result['valid']:
                        # 单条模式，直接显示错误
                        errors = []
                        for res in validation_result['results']:
                            if not res['valid']:
                                errors.extend(res['errors'])
                        raise ValueError('\n'.join(errors))
                    
                    # 校验通过，生成SQL
                    sqls, rollback_sqls = _generate_sqls(
                        data, dynamic_id, ops_remark)
                    
                    if sqls or rollback_sqls:
                        saved_file = _save_and_download(
                            sqls, rollback_sqls, dynamic_id, ops_remark_raw)
                    
            except Exception as e:
                logger.error(f"合同创建人修改错误: {str(e)}", exc_info=True)
                form.add_error(None, str(e))
    else:
        form = ContractCreatorUpdateForm()
    
    from ..navigation import get_sidebar_groups
    return render(request, 'contract_creator_form.html', {
        'form': form,
        'saved_file': saved_file,
        'validation_failure': validation_failure,
        'sidebar_groups': get_sidebar_groups(),
        'active_menu': 'contract_creator_update',
    })


def download_contract_creator_template(request):
    """下载合同创建人修改Excel模板"""
    if not OPENPYXL_AVAILABLE:
        return FileResponse(
            io.BytesIO(b"openpyxl not installed"),
            as_attachment=True,
            filename="error.txt"
        )
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "合同创建人修改"
    
    # 设置表头
    headers = ["合同号", "创建人账号", "原创建人账号", "原创建人姓名"]
    header_fill = PatternFill(start_color="4472C4",
                             end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # 示例数据
    ws.cell(row=2, column=1, value="CNEC03023002000-CLHT-25-0001-B001")
    ws.cell(row=2, column=2, value="cni23480626")
    ws.cell(row=2, column=3, value="")
    ws.cell(row=2, column=4, value="")
    
    # 调整列宽
    ws.column_dimensions['A'].width = 35
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 20
    
    # 保存到内存
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    response = FileResponse(
        output,
        as_attachment=True,
        filename='合同创建人修改模板.xlsx'
    )
    response['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    
    return response


def validate_contract_creator_api(request):
    """验证创建人账号是否存在"""
    login_name = request.GET.get('login_name', '').strip()
    
    if not login_name:
        return JsonResponse({
            'valid': False,
            'message': '请输入创建人账号'
        })
    
    # 查询user_org_detail表
    user_org = UserOrgDetail.objects.filter(login_name=login_name).first()
    
    if not user_org:
        return JsonResponse({
            'valid': False,
            'message': f'账号 {login_name} 不存在，请检查后重试'
        })
    
    return JsonResponse({
        'valid': True,
        'message': f'账号有效：{user_org.user_name or login_name}',
        'user_name': user_org.user_name or '',
    })


__all__ = ['contract_creator_update_view', 'download_contract_creator_template', 'validate_contract_creator_api']
