"""合同失效日期修改模块"""
import os
import io
import logging
import openpyxl
from datetime import datetime
from django.shortcuts import render
from django.http import FileResponse, HttpResponse
from django.conf import settings

from ..forms import EndDateUpdateForm
from ..navigation import get_sidebar_groups
from ..config import get_config
from ..sql_merge import chunk_list, format_in, merge_by_key
from .base import save_sql_file, parse_ops_remark
from ..validation_utils import (
    validate_required,
    validate_date_format,
    generate_validation_failure_excel,
    get_temp_filename_from_path,
)

logger = logging.getLogger('work_tools.view')


def validate_enddate_records(records):
    """
    校验失效日期修改记录
    """
    results = []
    passed_count = 0
    failed_count = 0
    
    for idx, record in enumerate(records):
        row_number = idx + 2
        errors = []
        
        # 必填项校验
        error = validate_required(record.get('bpo_id'), '合同编号')
        if error:
            errors.append(error)
        
        error = validate_required(record.get('end_date'), '新失效日期')
        if error:
            errors.append(error)
        else:
            # 日期格式校验
            error = validate_date_format(record.get('end_date'), '新失效日期')
            if error:
                errors.append(error)
        
        if errors:
            results.append({
                'row_number': row_number,
                'valid': False,
                'errors': errors
            })
            failed_count += 1
        else:
            results.append({
                'row_number': row_number,
                'valid': True,
                'errors': []
            })
            passed_count += 1
    
    return {
        'valid': failed_count == 0,
        'total': len(records),
        'passed': passed_count,
        'failed': failed_count,
        'results': results
    }


def parse_enddate_excel(file):
    """解析合同失效日期Excel文件"""
    wb = openpyxl.load_workbook(file)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    idx = {str(h): i for i, h in enumerate(headers)}

    def pick(alts):
        for a in alts:
            if a in idx:
                return idx[a]
        return None

    b_idx = pick(['bpo_id', 'BPO_ID', '合同号'])
    e_idx = pick(['end_date', 'END_DATE', '新失效日期'])
    o_idx = pick(['orig_end_date', '原失效日期'])

    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        bpo = row[b_idx] if b_idx is not None else None
        endd = row[e_idx] if e_idx is not None else None
        orig = row[o_idx] if o_idx is not None else None
        if bpo is None or endd is None:
            continue
        rec = {
            'bpo_id': str(bpo).strip(),
            'end_date': str(endd).strip(),
        }
        if orig is not None:
            rec['orig_end_date'] = str(orig).strip()
        records.append(rec)

    if not records:
        raise ValueError("Excel 中没有有效的行数据")
    return records


def generate_enddate_sql_bulk(records, ops_remark=None):
    """生成失效日期修改SQL"""
    cfg = get_config()
    sql = []
    sql.append("1、执行语句")

    if cfg.get('MERGE_MODULES', {}).get('enddate', True):
        def key_fn(r):
            return r.get('end_date')

        groups = merge_by_key(records, key_fn)
        for endd, recs in groups.items():
            ids = [r['bpo_id'] for r in recs if r.get('bpo_id')]
            for chunk in chunk_list(ids, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
                inlist = format_in(chunk)
                sql.append(
                    f"UPDATE tphct01 SET END_DATE='{endd}', APPR_STATE='ACTIVE', OPS_REMARK='{ops_remark or ''}' WHERE BPO_ID IN ({inlist}) AND ALIVE_FLAG='1';")
                sql.append(
                    f"UPDATE tphct02 SET APPR_STATE='ACTIVE', OPS_REMARK='{ops_remark or ''}' WHERE BPO_ID IN ({inlist}) AND ALIVE_FLAG='1';")
    else:
        for r in records:
            sql.append(
                f"UPDATE tphct01 SET END_DATE='{r['end_date']}', APPR_STATE='ACTIVE', OPS_REMARK='{ops_remark or ''}' WHERE BPO_ID='{r['bpo_id']}' AND ALIVE_FLAG='1';")
            sql.append(
                f"UPDATE tphct02 SET APPR_STATE='ACTIVE', OPS_REMARK='{ops_remark or ''}' WHERE BPO_ID='{r['bpo_id']}' AND ALIVE_FLAG='1';")

    sql.append("")
    sql.append("2、回退语句")

    if cfg.get('MERGE_MODULES', {}).get('enddate', True):
        def rb_key(r):
            return r.get('orig_end_date') or ''

        groups = merge_by_key(records, rb_key)
        for orig, recs in groups.items():
            ids = [r['bpo_id'] for r in recs if r.get('bpo_id')]
            for chunk in chunk_list(ids, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
                inlist = format_in(chunk)
                if orig:
                    sql.append(
                        f"UPDATE tphct01 SET END_DATE='{orig}', APPR_STATE='CLOSED', OPS_REMARK='' WHERE BPO_ID IN ({inlist}) AND ALIVE_FLAG='1';")
                else:
                    sql.append(
                        f"UPDATE tphct01 SET APPR_STATE='CLOSED', OPS_REMARK='' WHERE BPO_ID IN ({inlist}) AND ALIVE_FLAG='1';")
                sql.append(
                    f"UPDATE tphct02 SET APPR_STATE='CLOSED', OPS_REMARK='' WHERE BPO_ID IN ({inlist}) AND ALIVE_FLAG='1';")
    else:
        for r in records:
            if r.get('orig_end_date'):
                sql.append(
                    f"UPDATE tphct01 SET END_DATE='{r['orig_end_date']}', APPR_STATE='CLOSED', OPS_REMARK='' WHERE BPO_ID='{r['bpo_id']}' AND ALIVE_FLAG='1';")
            else:
                sql.append(
                    f"UPDATE tphct01 SET APPR_STATE='CLOSED', OPS_REMARK='' WHERE BPO_ID='{r['bpo_id']}' AND ALIVE_FLAG='1';")
            sql.append(
                f"UPDATE tphct02 SET APPR_STATE='CLOSED', OPS_REMARK='' WHERE BPO_ID='{r['bpo_id']}' AND ALIVE_FLAG='1';")

    sql.append("")
    sql.append("3、数据库")
    sql.append("ip：192.168.11.71")
    sql.append("库名：cnnc_ph")
    return "\n".join(sql)


def enddate_update_view(request):
    """合同失效日期修改视图"""
    saved_file = None
    validation_failure = None
    
    if request.method == 'POST':
        form = EndDateUpdateForm(request.POST, request.FILES)
        if form.is_valid():
            try:
                cd = form.cleaned_data
                ops_remark = parse_ops_remark(cd.get('ops_remark', ''))

                if cd.get('excel_file'):
                    records = parse_enddate_excel(cd['excel_file'])
                    
                    # 执行数据校验
                    validation_result = validate_enddate_records(records)
                    
                    if not validation_result['valid']:
                        temp_file = generate_validation_failure_excel(
                            cd['excel_file'],
                            validation_result,
                            'enddate'
                        )
                        
                        if temp_file:
                            validation_failure = {
                                'total': validation_result['total'],
                                'passed': validation_result['passed'],
                                'failed': validation_result['failed'],
                                'filename': get_temp_filename_from_path(temp_file)
                            }
                            logger.info(f"校验失败: 总行数={validation_failure['total']}, 失败行数={validation_failure['failed']}")
                        
                        return render(request, 'end_date_form.html', {
                            'form': form,
                            'validation_failure': validation_failure,
                            'active_menu': 'enddate',
                            'sidebar_groups': get_sidebar_groups(),
                        })
                    
                    sql_content = generate_enddate_sql_bulk(records, ops_remark)
                else:
                    records = [{
                        'bpo_id': cd['bpo_id'],
                        'end_date': cd['end_date'],
                        'orig_end_date': cd.get('orig_end_date')
                    }]
                    sql_content = generate_enddate_sql_bulk(records, ops_remark)

                # 保存SQL到固定目录
                saved_file = save_sql_file(
                    sql_content, '合同失效日期修改', cd.get('dynamic_id'))

                request.session['enddate_last'] = {
                    k: v for k, v in cd.items() if k != 'excel_file'
                }
            except Exception as e:
                logger.error(f"[失效日期修改] 处理失败: {e}", exc_info=True)
                form.add_error(None, f"生成SQL失败: {str(e)}")
    else:
        if request.GET.get('clear'):
            request.session.pop('enddate_last', None)
            form = EndDateUpdateForm()
        else:
            initial = request.session.get('enddate_last')
            form = EndDateUpdateForm(initial=initial)

    return render(request, 'end_date_form.html', {
        'form': form,
        'saved_file': saved_file,
        'validation_failure': validation_failure,
        'active_menu': 'enddate',
        'sidebar_groups': get_sidebar_groups(),
    })


def download_enddate_template(request):
    """下载失效日期模板"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '模板'
    ws.append(['合同号', '新失效日期', '原失效日期'])
    ws.append(['203S-25-00039', '20251125', '20251105'])
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    response = FileResponse(bio, as_attachment=True,
                            filename='enddate_template.xlsx')
    response['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    return response


__all__ = ['enddate_update_view', 'download_enddate_template']
