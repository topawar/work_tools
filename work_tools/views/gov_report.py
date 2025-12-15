"""政采云报告模块"""
import os
import io
import logging
import openpyxl
from datetime import datetime
from django.shortcuts import render
from django.http import FileResponse, HttpResponse
from django.conf import settings

from ..forms import GovReportForm
from ..navigation import get_sidebar_groups
from ..config import get_config
from ..sql_merge import chunk_list, format_in, merge_by_key
from .base import save_sql_file, parse_ops_remark
from ..validation_utils import (
    validate_required,
    validate_at_least_one,
    generate_validation_failure_excel,
    get_temp_filename_from_path
)

logger = logging.getLogger('work_tools.view')


def validate_gov_report_records(records):
    """校验政采云报送记录"""
    results = []
    passed_count = 0
    failed_count = 0
    
    for idx, record in enumerate(records):
        row_number = idx + 2  # 跳过表头
        errors = []
        
        # 至少有一项必填:采购方案编号,询价单编号,合同编号
        error = validate_at_least_one(
            [record.get('scheme_no'), record.get('inq_id'), record.get('bpo_id')],
            ['采购方案编号', '询价单编号', '合同编号']
        )
        if error:
            errors.append(error)
        
        # 必填项校验:是否报送必选
        if record.get('report_bool') is None:
            errors.append("是否报送必选")
        
        # 记录结果
        if errors:
            results.append({
                'row_number': row_number,
                'valid': False,
                'errors': errors
            })
            failed_count += 1
        else:
            results.append({
                'row_number': row_number,
                'valid': True,
                'errors': []
            })
            passed_count += 1
    
    return {
        'valid': failed_count == 0,
        'total': len(records),
        'passed': passed_count,
        'failed': failed_count,
        'results': results
    }


def parse_gov_excel(file):
    """解析政采云报告Excel文件"""
    wb = openpyxl.load_workbook(file)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    idx = {str(h).strip() if h else '': i for i, h in enumerate(headers)}
    
    logger.info(f"[政采云报送] Excel表头: {headers}")
    logger.info(f"[政采云报送] 列索引映射: {idx}")

    def pick(alts):
        for a in alts:
            if a in idx:
                return idx[a]
        return None

    s_idx = pick(['scheme_no', 'PURCHASE_SCHEME_NO', '采购方案编号'])
    i_idx = pick(['inq_id', 'INQ_ID', '询价单编号'])
    b_idx = pick(['bpo_id', 'BPO_ID', '合同号'])
    r_idx = pick(['report', '是否报送'])
    or_idx = pick(['orig_report', '原是否报送'])
    
    logger.info(f"[政采云报送] 列索引: s_idx={s_idx}, i_idx={i_idx}, b_idx={b_idx}, r_idx={r_idx}, or_idx={or_idx}")

    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        rec = {}
        if s_idx is not None and row[s_idx] is not None:
            rec['scheme_no'] = str(row[s_idx]).strip()
        if i_idx is not None and row[i_idx] is not None:
            rec['inq_id'] = str(row[i_idx]).strip()
        if b_idx is not None and row[b_idx] is not None:
            rec['bpo_id'] = str(row[b_idx]).strip()
        if r_idx is not None:
            v = row[r_idx]
            rec['report_bool'] = str(v).strip() in [
                '是', '1', 'yes', 'Yes', 'YES', 'true', 'True'] if v else None
        if or_idx is not None:
            v = row[or_idx]
            rec['orig_report_bool'] = str(v).strip() in [
                '是', '1', 'yes', 'Yes', 'YES', 'true', 'True'] if v else None
        if rec.get('scheme_no') or rec.get('inq_id') or rec.get('bpo_id'):
            records.append(rec)

    logger.info(f"[政采云报送] 解析到 {len(records)} 条记录")
    if records:
        logger.info(f"[政采云报送] 第一条记录示例: {records[0]}")
    
    if not records:
        raise ValueError("Excel 中没有有效的行数据")
    return records


def generate_gov_sql_bulk(records, ops_remark=None, report_choice=None, orig_report_choice=None):
    """生成政采云报送修改SQL"""
    cfg = get_config()
    sql = []
    sql.append("1、执行语句")

    if cfg.get('MERGE_MODULES', {}).get('gov', True):
        def key_fn(r):
            rb = r.get('report_bool')
            if rb is None:
                rb = (report_choice == 'yes') if report_choice is not None else True
            return rb

        groups = merge_by_key(records, key_fn)
        for rb, recs in groups.items():
            report_val = 1 if rb else 0
            submit_val = 1 if rb else 0
            schemes = [r.get('scheme_no') for r in recs if r.get('scheme_no')]
            inqs = [r.get('inq_id') for r in recs if r.get('inq_id')]
            bpos = [r.get('bpo_id') for r in recs if r.get('bpo_id')]

            for chunk in chunk_list(schemes, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
                sql.append(
                    f"UPDATE tprfa02 SET IS_REPORT_TO_SASAC='{report_val}', OPS_REMARK='{ops_remark or ''}' WHERE PURCHASE_SCHEME_NO IN ({format_in(chunk)});")
            for chunk in chunk_list(inqs, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
                sql.append(
                    f"UPDATE tprxj02 SET IS_REPORT_TO_SASAC='{report_val}', OPS_REMARK='{ops_remark or ''}' WHERE INQ_ID IN ({format_in(chunk)});")
            for chunk in chunk_list(bpos, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
                sql.append(
                    f"UPDATE tphct01 SET IS_SUBMIT_SASAC='{submit_val}', OPS_REMARK='{ops_remark or ''}' WHERE BPO_ID IN ({format_in(chunk)}) AND ALIVE_FLAG='1';")
    else:
        for r in records:
            rb = r.get('report_bool')
            if rb is None:
                rb = (report_choice == 'yes') if report_choice is not None else True
            report_val = 1 if rb else 0
            submit_val = 1 if rb else 0
            if r.get('scheme_no'):
                sql.append(
                    f"UPDATE tprfa02 SET IS_REPORT_TO_SASAC='{report_val}', OPS_REMARK='{ops_remark or ''}' WHERE PURCHASE_SCHEME_NO='{r['scheme_no']}';")
            if r.get('inq_id'):
                sql.append(
                    f"UPDATE tprxj02 SET IS_REPORT_TO_SASAC='{report_val}', OPS_REMARK='{ops_remark or ''}' WHERE INQ_ID='{r['inq_id']}';")
            if r.get('bpo_id'):
                sql.append(
                    f"UPDATE tphct01 SET IS_SUBMIT_SASAC='{submit_val}', OPS_REMARK='{ops_remark or ''}' WHERE BPO_ID='{r['bpo_id']}' AND ALIVE_FLAG='1';")

    sql.append("")
    sql.append("2、回退语句")

    if cfg.get('MERGE_MODULES', {}).get('gov', True):
        def rb_key(r):
            ob = r.get('orig_report_bool')
            if ob is None:
                if orig_report_choice is not None:
                    ob = (orig_report_choice == 'yes')
                else:
                    rb = r.get('report_bool')
                    if rb is None:
                        rb = (report_choice ==
                              'yes') if report_choice is not None else True
                    ob = not rb
            return ob

        groups = merge_by_key(records, rb_key)
        for ob, recs in groups.items():
            report_val = 1 if ob else 0
            submit_val = 1 if ob else 0
            schemes = [r.get('scheme_no') for r in recs if r.get('scheme_no')]
            inqs = [r.get('inq_id') for r in recs if r.get('inq_id')]
            bpos = [r.get('bpo_id') for r in recs if r.get('bpo_id')]

            for chunk in chunk_list(schemes, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
                sql.append(
                    f"UPDATE tprfa02 SET IS_REPORT_TO_SASAC='{report_val}', OPS_REMARK='' WHERE PURCHASE_SCHEME_NO IN ({format_in(chunk)});")
            for chunk in chunk_list(inqs, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
                sql.append(
                    f"UPDATE tprxj02 SET IS_REPORT_TO_SASAC='{report_val}', OPS_REMARK='' WHERE INQ_ID IN ({format_in(chunk)});")
            for chunk in chunk_list(bpos, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
                sql.append(
                    f"UPDATE tphct01 SET IS_SUBMIT_SASAC='{submit_val}', OPS_REMARK='' WHERE BPO_ID IN ({format_in(chunk)}) AND ALIVE_FLAG='1';")
    else:
        for r in records:
            ob = r.get('orig_report_bool')
            if ob is None:
                if orig_report_choice is not None:
                    ob = (orig_report_choice == 'yes')
                else:
                    rb = r.get('report_bool')
                    if rb is None:
                        rb = (report_choice ==
                              'yes') if report_choice is not None else True
                    ob = not rb
            report_val = 1 if ob else 0
            submit_val = 1 if ob else 0
            if r.get('scheme_no'):
                sql.append(
                    f"UPDATE tprfa02 SET IS_REPORT_TO_SASAC='{report_val}', OPS_REMARK='' WHERE PURCHASE_SCHEME_NO='{r['scheme_no']}';")
            if r.get('inq_id'):
                sql.append(
                    f"UPDATE tprxj02 SET IS_REPORT_TO_SASAC='{report_val}', OPS_REMARK='' WHERE INQ_ID='{r['inq_id']}';")
            if r.get('bpo_id'):
                sql.append(
                    f"UPDATE tphct01 SET IS_SUBMIT_SASAC='{submit_val}', OPS_REMARK='' WHERE BPO_ID='{r['bpo_id']}' AND ALIVE_FLAG='1';")

    sql.append("")
    sql.append("3、数据库")
    sql.append("ip:192.168.11.69")
    sql.append("库名：cnnc_pr")
    sql.append("")
    sql.append("ip:192.168.11.71")
    sql.append("库名：cnnc_ph")
    return "\n".join(sql)


def gov_report_view(request):
    """政采云报送修改视图"""
    saved_file = None
    validation_failure = None  # 新增
    if request.method == 'POST':
        form = GovReportForm(request.POST, request.FILES)
        if form.is_valid():
            cd = form.cleaned_data
            ops_remark = parse_ops_remark(cd.get('ops_remark', ''))

            if cd.get('excel_file'):
                records = parse_gov_excel(cd['excel_file'])
                
                # 执行数据校验
                validation_result = validate_gov_report_records(records)
                
                if not validation_result['valid']:
                    # 校验失败,生成失败文件
                    temp_file = generate_validation_failure_excel(
                        cd['excel_file'],
                        validation_result,
                        'gov_report'
                    )
                    
                    if temp_file:
                        validation_failure = {
                            'total': validation_result['total'],
                            'passed': validation_result['passed'],
                            'failed': validation_result['failed'],
                            'filename': get_temp_filename_from_path(temp_file)
                        }
                    
                    return render(request, 'gov_report_form.html', {
                        'form': form,
                        'validation_failure': validation_failure,
                        'active_menu': 'gov_report',
                        'sidebar_groups': get_sidebar_groups(),
                    })
                
                sql_content = generate_gov_sql_bulk(records, ops_remark)
            else:
                rec = {
                    'scheme_no': cd.get('scheme_no'),
                    'inq_id': cd.get('inq_id'),
                    'bpo_id': cd.get('bpo_id'),
                }
                records = [rec]
                sql_content = generate_gov_sql_bulk(
                    records, ops_remark, cd.get(
                        'report_choice'), cd.get('orig_report_choice')
                )

            # 保存SQL到固定目录
            saved_file = save_sql_file(
                sql_content, '是否报送国资委', cd.get('dynamic_id'))

            request.session['gov_report_last'] = {
                k: v for k, v in cd.items() if k != 'excel_file'
            }
    else:
        if request.GET.get('clear'):
            request.session.pop('gov_report_last', None)
            form = GovReportForm()
        else:
            initial = request.session.get('gov_report_last')
            form = GovReportForm(initial=initial)

    return render(request, 'gov_report_form.html', {
        'form': form,
        'saved_file': saved_file,
        'validation_failure': validation_failure,  # 新增
        'active_menu': 'gov_report',
        'sidebar_groups': get_sidebar_groups(),
    })


def download_gov_template(request):
    """下载政采云报送模板"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '模板'
    ws.append(['采购方案编号', '询价单编号', '合同号', '是否报送', '原是否报送'])
    ws.append(['CNEC-CGFA-25-20795', 'CNEC-XJD-25-20009',
              '23EC-25-02605', '是', '否'])
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    response = FileResponse(bio, as_attachment=True,
                            filename='gov_report_template.xlsx')
    response['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    return response


__all__ = ['gov_report_view', 'download_gov_template']
