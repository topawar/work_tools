"""ERP合同终止模块"""
import os
import io
import openpyxl
from datetime import datetime
from django.shortcuts import render
from django.http import FileResponse, HttpResponse
from django.conf import settings

from ..forms import ErpTerminateForm
from ..navigation import get_sidebar_groups
from ..config import get_config
from ..sql_merge import chunk_list, format_in
from .base import save_sql_file, parse_ops_remark
from ..validation_utils import (
    validate_required,
    generate_validation_failure_excel,
    get_temp_filename_from_path
)


def validate_erp_terminate_records(records):
    """校验核电ERP终止记录"""
    results = []
    passed_count = 0
    failed_count = 0
    
    for idx, record in enumerate(records):
        row_number = idx + 2  # 跳过表头
        errors = []
        
        # 三个字段都必填:询价单编号、采购方案编号、采购包编号
        error = validate_required(record.get('inq_id'), '询价单编号')
        if error:
            errors.append(error)
        
        error = validate_required(record.get('purchase_scheme_no'), '采购方案编号')
        if error:
            errors.append(error)
        
        error = validate_required(record.get('purchase_package_no'), '采购包编号')
        if error:
            errors.append(error)
        
        # 记录结果
        if errors:
            results.append({
                'row_number': row_number,
                'valid': False,
                'errors': errors
            })
            failed_count += 1
        else:
            results.append({
                'row_number': row_number,
                'valid': True,
                'errors': []
            })
            passed_count += 1
    
    return {
        'valid': failed_count == 0,
        'total': len(records),
        'passed': passed_count,
        'failed': failed_count,
        'results': results
    }


def parse_erp_excel(file):
    """解析ERP终止Excel文件"""
    wb = openpyxl.load_workbook(file)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    idx = {str(h): i for i, h in enumerate(headers)}

    def pick(alts):
        for a in alts:
            if a in idx:
                return idx[a]
        return None

    inq_idx = pick(['INQ_ID', 'inq_id'])
    sch_idx = pick(['PURCHASE_SCHEME_NO', 'purchase_scheme_no'])
    pkg_idx = pick(['PURCHASE_PACKAGE_NO', 'purchase_package_no'])

    if inq_idx is None and sch_idx is None and pkg_idx is None:
        raise ValueError(
            "Excel 缺少必要列：INQ_ID或PURCHASE_SCHEME_NO或PURCHASE_PACKAGE_NO")

    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        rec = {}
        if inq_idx is not None:
            v = row[inq_idx]
            rec['inq_id'] = v and str(v).strip() or None
        if sch_idx is not None:
            v = row[sch_idx]
            rec['purchase_scheme_no'] = v and str(v).strip() or None
        if pkg_idx is not None:
            v = row[pkg_idx]
            rec['purchase_package_no'] = v and str(v).strip() or None
        if rec.get('inq_id') or rec.get('purchase_scheme_no') or rec.get('purchase_package_no'):
            records.append(rec)

    if not records:
        raise ValueError("Excel 中没有有效的行数据")
    return records


def generate_erp_terminate_sql(records, ops_remark=None):
    """生成ERP终止SQL"""
    cfg = get_config()
    sql = []
    sql.append("1、执行语句")
    remark = ops_remark or ''

    if cfg.get('MERGE_MODULES', {}).get('erp', True):
        inqs = [r.get('inq_id') for r in records if r.get('inq_id')]
        schs = [r.get('purchase_scheme_no')
                for r in records if r.get('purchase_scheme_no')]
        pkgs = [r.get('purchase_package_no')
                for r in records if r.get('purchase_package_no')]

        for chunk in chunk_list(inqs, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
            inlist = format_in(chunk)
            sql.append(
                f"update tprxj01 set INQ_STATUS='70', OPS_REMARK='{remark}' where INQ_ID in ({inlist});")
            sql.append(
                f"update tprxj06 set ALIVE_FLAG='0', OPS_REMARK='{remark}' where INQ_ID in ({inlist});")
        for chunk in chunk_list(schs, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
            sql.append(
                f"update tprfa01 set PURCHASE_SCHEME_STATUS='95', OPS_REMARK='{remark}' where PURCHASE_SCHEME_NO IN ({format_in(chunk)});")
        for chunk in chunk_list(pkgs, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
            inlist = format_in(chunk)
            sql.append(
                f"update tprly01 set status='60', OPS_REMARK='{remark}' where PURCHASE_PACKAGE_NO IN ({inlist});")
            sql.append(
                f"update tprly02 set status='60', OPS_REMARK='{remark}' where PURCHASE_PACKAGE_NO IN ({inlist});")
    else:
        for r in records:
            inq = r.get('inq_id')
            sch = r.get('purchase_scheme_no')
            pkg = r.get('purchase_package_no')
            if inq:
                sql.append(
                    f"update tprxj01 set INQ_STATUS='70', OPS_REMARK='{remark}' where INQ_ID in ('{inq}');")
                sql.append(
                    f"update tprxj06 set ALIVE_FLAG='0', OPS_REMARK='{remark}' where INQ_ID  in ('{inq}');")
            if sch:
                sql.append(
                    f"update tprfa01 set PURCHASE_SCHEME_STATUS='95', OPS_REMARK='{remark}' where PURCHASE_SCHEME_NO IN ('{sch}');")
            if pkg:
                sql.append(
                    f"update tprly01 set status='60', OPS_REMARK='{remark}' where PURCHASE_PACKAGE_NO IN ('{pkg}');")
                sql.append(
                    f"update tprly02 set status='60', OPS_REMARK='{remark}' where PURCHASE_PACKAGE_NO IN ('{pkg}');")

    sql.append("")
    sql.append("2、回退语句")

    if cfg.get('MERGE_MODULES', {}).get('erp', True):
        inqs = [r.get('inq_id') for r in records if r.get('inq_id')]
        schs = [r.get('purchase_scheme_no')
                for r in records if r.get('purchase_scheme_no')]
        pkgs = [r.get('purchase_package_no')
                for r in records if r.get('purchase_package_no')]

        for chunk in chunk_list(inqs, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
            inlist = format_in(chunk)
            sql.append(
                f"update tprxj01 set INQ_STATUS='55', OPS_REMARK='' where INQ_ID in ({inlist});")
            sql.append(
                f"update tprxj06 set ALIVE_FLAG='1', OPS_REMARK='' where INQ_ID in ({inlist});")
        for chunk in chunk_list(schs, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
            sql.append(
                f"update tprfa01 set PURCHASE_SCHEME_STATUS='55', OPS_REMARK='' where PURCHASE_SCHEME_NO IN ({format_in(chunk)});")
        for chunk in chunk_list(pkgs, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
            inlist = format_in(chunk)
            sql.append(
                f"update tprly01 set status='75', OPS_REMARK='' where PURCHASE_PACKAGE_NO IN ({inlist});")
            sql.append(
                f"update tprly02 set status='75', OPS_REMARK='' where PURCHASE_PACKAGE_NO IN ({inlist});")
    else:
        for r in records:
            inq = r.get('inq_id')
            sch = r.get('purchase_scheme_no')
            pkg = r.get('purchase_package_no')
            if inq:
                sql.append(
                    f"update tprxj01 set INQ_STATUS='55', OPS_REMARK='' where INQ_ID in ('{inq}');")
                sql.append(
                    f"update tprxj06 set ALIVE_FLAG='1', OPS_REMARK='' where INQ_ID in ('{inq}');")
            if sch:
                sql.append(
                    f"update tprfa01 set PURCHASE_SCHEME_STATUS='55', OPS_REMARK='' where PURCHASE_SCHEME_NO IN ('{sch}');")
            if pkg:
                sql.append(
                    f"update tprly01 set status='75', OPS_REMARK='' where PURCHASE_PACKAGE_NO IN ('{pkg}');")
                sql.append(
                    f"update tprly02 set status='75', OPS_REMARK='' where PURCHASE_PACKAGE_NO IN ('{pkg}');")

    sql.append("")
    sql.append("ip：192.168.11.69")
    sql.append("库名：cnnc_pr")
    return "\n".join(sql)


def erp_terminate_view(request):
    """ERP合同终止视图"""
    saved_file = None
    validation_failure = None  # 新增
    if request.method == 'POST':
        form = ErpTerminateForm(request.POST, request.FILES)
        if form.is_valid():
            cd = form.cleaned_data
            ops_remark = parse_ops_remark(cd.get('ops_remark', ''))

            if cd.get('excel_file'):
                records = parse_erp_excel(cd['excel_file'])
                
                # 执行数据校验
                validation_result = validate_erp_terminate_records(records)
                
                if not validation_result['valid']:
                    # 校验失败,生成失败文件
                    temp_file = generate_validation_failure_excel(
                        cd['excel_file'],
                        validation_result,
                        'erp_terminate'
                    )
                    
                    if temp_file:
                        validation_failure = {
                            'total': validation_result['total'],
                            'passed': validation_result['passed'],
                            'failed': validation_result['failed'],
                            'filename': get_temp_filename_from_path(temp_file)
                        }
                    
                    return render(request, 'erp_terminate_form.html', {
                        'form': form,
                        'validation_failure': validation_failure,
                        'active_menu': 'erp_terminate',
                        'sidebar_groups': get_sidebar_groups(),
                    })
            else:
                rec = {
                    'inq_id': (cd.get('inq_id') or '').strip() or None,
                    'purchase_scheme_no': (cd.get('purchase_scheme_no') or '').strip() or None,
                    'purchase_package_no': (cd.get('purchase_package_no') or '').strip() or None,
                }
                records = [rec]

            sql_content = generate_erp_terminate_sql(records, ops_remark)

            # 保存SQL到固定目录
            saved_file = save_sql_file(
                sql_content, 'ERP终止', cd.get('dynamic_id'))

            # 保存会话数据
            session_data = {}
            for k, v in cd.items():
                if k != 'excel_file':
                    if hasattr(v, 'to_eng_string'):
                        session_data[k] = v.to_eng_string()
                    else:
                        session_data[k] = v
            request.session['erp_terminate_last'] = session_data
    else:
        if request.GET.get('clear'):
            request.session.pop('erp_terminate_last', None)
            form = ErpTerminateForm()
        else:
            initial = request.session.get('erp_terminate_last')
            form = ErpTerminateForm(initial=initial)

    return render(request, 'erp_terminate_form.html', {
        'form': form,
        'saved_file': saved_file,
        'validation_failure': validation_failure,  # 新增
        'active_menu': 'erp_terminate',
        'sidebar_groups': get_sidebar_groups(),
    })


def download_erp_terminate_template(request):
    """下载ERP终止模板"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '模板'
    ws.append(['INQ_ID', 'PURCHASE_SCHEME_NO', 'PURCHASE_PACKAGE_NO'])
    ws.append(['CNSC-XJD-25-03203', 'YW01-CGFA-25-02641', 'YW01-CGB-25-02749'])
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    response = FileResponse(bio, as_attachment=True,
                            filename='erp_terminate_template.xlsx')
    response['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    return response


__all__ = ['erp_terminate_view', 'download_erp_terminate_template']
