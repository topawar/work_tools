"""浮动单价类型修改模块"""
import os
import io
import openpyxl
import logging
from datetime import datetime
from django.shortcuts import render
from django.http import FileResponse, HttpResponse
from django.conf import settings

from ..forms import FloatingPriceTypeForm
from ..navigation import SIDEBAR_GROUPS
from ..config import get_config
from ..sql_merge import chunk_list, format_in, merge_by_key
from .base import save_sql_file, parse_ops_remark

logger = logging.getLogger('work_tools.view')


def parse_price_type_excel(file):
    """解析浮动单价类型Excel文件"""
    wb = openpyxl.load_workbook(file)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    idx = {str(h): i for i, h in enumerate(headers)}

    def pick(alts):
        for a in alts:
            if a in idx:
                return idx[a]
        return None

    bpo_idx = pick(['BPO_ID', 'bpo_id', '合同编号'])
    if bpo_idx is None:
        raise ValueError('Excel 缺少必要列：合同编号')

    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        bpo = row[bpo_idx]
        if not bpo:
            continue
        rec = {
            'bpo_id': str(bpo).strip(),
        }
        records.append(rec)

    if not records:
        raise ValueError('Excel 中没有有效的行数据')
    return records


def generate_price_type_sql(records, ops_remark=''):
    """生成浮动单价类型修改SQL"""
    cfg = get_config()
    ops_remark = parse_ops_remark(ops_remark)
    sql = []
    sql.append('1、执行语句')
    
    logger.info(f"[SQL合并] 开始生成浮动单价类型SQL, 记录数={len(records)}, 配置启用={cfg.get('MERGE_MODULES', {}).get('price_type', True)}")
    
    newt = '20'
    if cfg.get('MERGE_MODULES', {}).get('price_type', True):
        # SQL合并模式
        bpo_ids = [r['bpo_id'] for r in records if r.get('bpo_id')]
        logger.info(f"[SQL合并] 浮动单价类型合并, 总BPO数={len(bpo_ids)}")
        
        for chunk in chunk_list(bpo_ids, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
            inlist = format_in(chunk)
            sql.append(
                f"update TPHCT02 set PUR_PRICE_TYPE='{newt}', OPS_REMARK='{ops_remark}' WHERE BPO_ID IN ({inlist}) and ALIVE_FLAG='1';")
            sql.append(
                f"update tprly04 set PUR_PRICE_TYPE='{newt}', OPS_REMARK='{ops_remark}' WHERE BPO_ID IN ({inlist}) and ALIVE_FLAG='1';")
    else:
        # 非合并模式
        for r in records:
            bpo = r['bpo_id']
            sql.append(
                f"update TPHCT02 set PUR_PRICE_TYPE='{newt}', OPS_REMARK='{ops_remark}' WHERE BPO_ID='{bpo}' and ALIVE_FLAG='1';")
            sql.append(
                f"update tprly04 set PUR_PRICE_TYPE='{newt}', OPS_REMARK='{ops_remark}' WHERE BPO_ID='{bpo}' and ALIVE_FLAG='1';")

    sql.append('')
    sql.append('2、回退语句')
    
    orig = '10'
    if cfg.get('MERGE_MODULES', {}).get('price_type', True):
        # SQL合并模式
        bpo_ids = [r['bpo_id'] for r in records if r.get('bpo_id')]
        for chunk in chunk_list(bpo_ids, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
            inlist = format_in(chunk)
            sql.append(
                f"update TPHCT02 set PUR_PRICE_TYPE='{orig}', OPS_REMARK='' WHERE BPO_ID IN ({inlist}) and ALIVE_FLAG='1';")
            sql.append(
                f"update tprly04 set PUR_PRICE_TYPE='{orig}', OPS_REMARK='' WHERE BPO_ID IN ({inlist}) and ALIVE_FLAG='1';")
    else:
        # 非合并模式
        for r in records:
            bpo = r['bpo_id']
            sql.append(
                f"update TPHCT02 set PUR_PRICE_TYPE='{orig}', OPS_REMARK='' WHERE BPO_ID='{bpo}' and ALIVE_FLAG='1';")
            sql.append(
                f"update tprly04 set PUR_PRICE_TYPE='{orig}', OPS_REMARK='' WHERE BPO_ID='{bpo}' and ALIVE_FLAG='1';")

    sql.append('')
    sql.append('3.数据库')
    sql.append('ip：192.168.11.71')
    sql.append('库名：cnnc_ph')
    return "\n".join(sql)


def floating_price_type_view(request):
    """浮动单价类型修改视图"""
    saved_file = None
    if request.method == 'POST':
        form = FloatingPriceTypeForm(request.POST, request.FILES)
        if form.is_valid():
            cd = form.cleaned_data
            ops_remark = cd.get('ops_remark', '')

            if cd.get('excel_file'):
                records = parse_price_type_excel(cd['excel_file'])
                valid = [r for r in records if r.get('bpo_id')]
                sql_content = generate_price_type_sql(valid, ops_remark)
            else:
                rec = {'bpo_id': cd['bpo_id']}
                sql_content = generate_price_type_sql([rec], ops_remark)

            # 保存SQL到固定目录
            saved_file = save_sql_file(
                sql_content, '浮动单价类型修改', cd.get('dynamic_id'))

            # 保存会话数据
            session_data = {}
            for k, v in cd.items():
                if k != 'excel_file':
                    if hasattr(v, 'to_eng_string'):
                        session_data[k] = v.to_eng_string()
                    else:
                        session_data[k] = v
            request.session['floating_price_type_last'] = session_data
    else:
        if request.GET.get('clear'):
            request.session.pop('floating_price_type_last', None)
            form = FloatingPriceTypeForm()
        else:
            initial = request.session.get('floating_price_type_last')
            form = FloatingPriceTypeForm(initial=initial)

    return render(request, 'floating_price_type_form.html', {
        'form': form,
        'saved_file': saved_file,
        'active_menu': 'price_type_update',
        'sidebar_groups': SIDEBAR_GROUPS,
    })


def download_price_type_template(request):
    """下载浮动单价类型模板"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '模板'
    ws.append(['合同编号'])
    ws.append(['22EC-25-02800'])
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    response = FileResponse(bio, as_attachment=True,
                            filename='price_type_template.xlsx')
    response['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    return response


__all__ = ['floating_price_type_view', 'download_price_type_template']
