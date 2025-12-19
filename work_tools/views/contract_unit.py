"""
合同单位修改模块
包含起草单位和签约主体的修改功能
"""
import io
import os
import logging
from datetime import datetime
from django.shortcuts import render
from django.http import FileResponse
from ..forms import UnitChangeForm
from ..navigation import get_sidebar_groups
from ..config import get_config
from ..models import OrgDetail
from ..sql_merge import chunk_list, format_in, merge_by_key, compose_or
from ..logger_utils import log_view_input
from .base import parse_ops_remark, extract_company_code, extract_company_name, _unique_code_by_name, save_sql_file
from ..validation_utils import (
    validate_required,
    validate_at_least_one,
    generate_validation_failure_excel,
    get_temp_filename_from_path
)

try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

logger = logging.getLogger('work_tools.view')


def parse_unit_name_with_code(input_value):
    """
    解析单位名称输入，支持两种格式：
    1. 纯名称：中核（上海）供应链管理有限公司
    2. 名称-编码：中核（上海）供应链管理有限公司-99280563919161110148
    
    返回: (company_name, company_code)
    """
    if not input_value:
        return None, None
    
    input_value = input_value.strip()
    if not input_value:  # 处理空字符串或只有空格的情况
        return None, None
    
    # 检查是否包含编码（以-分隔，且最后部分看起来像编码）
    if '-' in input_value:
        parts = input_value.rsplit('-', 1)  # 从右边分割，只分割一次
        if len(parts) == 2:
            name_part = parts[0].strip()
            code_part = parts[1].strip()
            
            # 简单判断code_part是否像编码（长度>5且包含数字）
            if len(code_part) > 5 and any(c.isdigit() for c in code_part):
                return name_part, code_part
    
    # 如果不包含编码或格式不对，返回纯名称
    return input_value, None


def validate_original_unit_names(old_drafting_name=None, old_party_name=None):
    """校验原单位名称是否在数据库中存在"""
    errors = []
    
    if old_drafting_name:
        name, code = parse_unit_name_with_code(old_drafting_name)
        
        if code:
            # 如果有编码，进行精确匹配（名称+编码）
            exists = OrgDetail.objects.filter(
                company_name=name, 
                company_code=code
            ).exists()
            if not exists:
                errors.append(f'原起草单位名称"{name}"（编码：{code}）在数据库中不存在')
        else:
            # 如果没有编码，检查名称是否存在
            records = OrgDetail.objects.filter(company_name=name)
            if not records.exists():
                errors.append(f'原起草单位名称"{name}"在数据库中不存在')
            elif records.count() > 1:
                # 如果存在多条记录，提示用户选择具体的编码
                codes = [r.company_code for r in records if r.company_code]
                if codes:
                    codes_str = '、'.join(codes)
                    errors.append(f'原起草单位名称"{name}"存在多条记录，请选择具体编码：{codes_str}')
    
    if old_party_name:
        name, code = parse_unit_name_with_code(old_party_name)
        
        if code:
            # 如果有编码，进行精确匹配（名称+编码）
            exists = OrgDetail.objects.filter(
                company_name=name, 
                company_code=code
            ).exists()
            if not exists:
                errors.append(f'原签约主体名称"{name}"（编码：{code}）在数据库中不存在')
        else:
            # 如果没有编码，检查名称是否存在
            records = OrgDetail.objects.filter(company_name=name)
            if not records.exists():
                errors.append(f'原签约主体名称"{name}"在数据库中不存在')
            elif records.count() > 1:
                # 如果存在多条记录，提示用户选择具体的编码
                codes = [r.company_code for r in records if r.company_code]
                if codes:
                    codes_str = '、'.join(codes)
                    errors.append(f'原签约主体名称"{name}"存在多条记录，请选择具体编码：{codes_str}')
    
    return errors


def validate_unit_change_records(records):
    """校验合同起草签约单位修改记录"""
    results = []
    passed_count = 0
    failed_count = 0
    
    for idx, record in enumerate(records):
        row_number = idx + 2  # 跳过表头
        errors = []
        
        # 至少有一项必填:新起草单位名称、新签约主体名称
        error = validate_at_least_one(
            [record.get('new_drafting_name'), record.get('new_party_name')],
            ['新起草单位名称', '新签约主体名称']
        )
        if error:
            errors.append(error)
        
        # 至少有一项必填:采购方案编号、询价单编号、定标结果编号
        error = validate_at_least_one(
            [record.get('scheme'), record.get('inquiry'), record.get('result')],
            ['采购方案编号', '询价单编号', '定标结果编号']
        )
        if error:
            errors.append(error)
        
        # 记录结果
        if errors:
            results.append({
                'row_number': row_number,
                'valid': False,
                'errors': errors
            })
            failed_count += 1
        else:
            results.append({
                'row_number': row_number,
                'valid': True,
                'errors': []
            })
            passed_count += 1
    
    return {
        'valid': failed_count == 0,
        'total': len(records),
        'passed': passed_count,
        'failed': failed_count,
        'results': results
    }


def generate_unit_sql(records, **kwargs):
    """生成单位修改SQL"""
    cfg = get_config()
    update_drafting = kwargs.get('update_drafting', False)
    update_contract_party = kwargs.get('update_contract_party', False)
    old_drafting_id = kwargs.get('old_drafting_id')
    old_drafting_name = kwargs.get('old_drafting_name')
    old_party_id = kwargs.get('old_party_id')
    old_party_name = kwargs.get('old_party_name')
    ops_remark = parse_ops_remark(kwargs.get('ops_remark', ''))

    sql = []
    sql.append("1、执行语句")

    if cfg.get('MERGE_MODULES', {}).get('unit', True):
        def set_parts_from(r):
            parts = []
            if update_drafting:
                ndid = r.get('new_drafting_id', kwargs.get('new_drafting_id'))
                ndnm = r.get('new_drafting_name',
                             kwargs.get('new_drafting_name'))
                if not ndid and ndnm:
                    code = extract_company_code(ndnm)
                    if code:
                        ndid = code
                if ndid:
                    parts.append(f"PRIMARY_CONTRACT_DRAFTING_UNIT='{ndid}'")
                if ndnm:
                    parts.append(
                        f"PRIMARY_CONTRACT_DRAFTING_UNIT_NAME='{extract_company_name(ndnm)}'")
            if update_contract_party:
                npid = r.get('new_party_id', kwargs.get('new_party_id'))
                npnm = r.get('new_party_name', kwargs.get('new_party_name'))
                if not npid and npnm:
                    code = extract_company_code(npnm)
                    if code:
                        npid = code
                if npid:
                    parts.append(f"CONTRACT_PARTY_ID='{npid}'")
                if npnm:
                    parts.append(
                        f"CONTRACT_PARTY_NAME='{extract_company_name(npnm)}'")
            return ', '.join(parts)

        groups = merge_by_key(records, set_parts_from)
        for set_clause, recs in groups.items():
            if not set_clause:
                continue
            schemes = [r['scheme'] for r in recs if r.get('scheme')]
            for chunk in chunk_list(schemes, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
                sql.append(
                    f"update tprfa03 set {set_clause}, OPS_REMARK='{ops_remark}' where PURCHASE_SCHEME_NO IN ({format_in(chunk)});")
            conds_xj = [{'PURCHASE_SCHEME_NO': r.get(
                'scheme'), 'INQ_ID': r.get('inquiry')} for r in recs]
            conds_xj = [
                c for c in conds_xj if c['PURCHASE_SCHEME_NO'] and c['INQ_ID']]
            if conds_xj:
                where_or = compose_or(
                    conds_xj, ['PURCHASE_SCHEME_NO', 'INQ_ID'])
                sql.append(
                    f"update tprxj05 set {set_clause}, OPS_REMARK='{ops_remark}' where {where_or};")
            conds_nq = [{'PURCHASE_SCHEME_NO': r.get('scheme'), 'INQ_ID': r.get(
                'inquiry'), 'sign_id': r.get('result')} for r in recs]
            conds_nq = [c for c in conds_nq if c['PURCHASE_SCHEME_NO']
                        and c['INQ_ID'] and c['sign_id']]
            if conds_nq:
                where_or = compose_or(
                    conds_nq, ['PURCHASE_SCHEME_NO', 'INQ_ID', 'sign_id'])
                sql.append(
                    f"update tprnq03 set {set_clause}, OPS_REMARK='{ops_remark}' where {where_or};")
    else:
        for r in records:
            parts = []
            if update_drafting:
                ndid = r.get('new_drafting_id', kwargs.get('new_drafting_id'))
                ndnm = r.get('new_drafting_name',
                             kwargs.get('new_drafting_name'))
                if not ndid and ndnm:
                    code = extract_company_code(ndnm)
                    if code:
                        ndid = code
                if ndid:
                    parts.append(f"PRIMARY_CONTRACT_DRAFTING_UNIT='{ndid}'")
                if ndnm:
                    parts.append(
                        f"PRIMARY_CONTRACT_DRAFTING_UNIT_NAME='{extract_company_name(ndnm)}'")
            if update_contract_party:
                npid = r.get('new_party_id', kwargs.get('new_party_id'))
                npnm = r.get('new_party_name', kwargs.get('new_party_name'))
                if not npid and npnm:
                    code = extract_company_code(npnm)
                    if code:
                        npid = code
                if npid:
                    parts.append(f"CONTRACT_PARTY_ID='{npid}'")
                if npnm:
                    parts.append(
                        f"CONTRACT_PARTY_NAME='{extract_company_name(npnm)}'")
            if parts:
                set_clause = ', '.join(parts)
                sql.append(
                    f"update tprfa03 set {set_clause}, OPS_REMARK='{ops_remark}' where PURCHASE_SCHEME_NO='{r['scheme']}';")
                sql.append(
                    f"update tprxj05 set {set_clause}, OPS_REMARK='{ops_remark}' where PURCHASE_SCHEME_NO='{r['scheme']}' AND INQ_ID='{r['inquiry']}';")
                sql.append(
                    f"update tprnq03 set {set_clause}, OPS_REMARK='{ops_remark}' where PURCHASE_SCHEME_NO='{r['scheme']}' AND INQ_ID='{r['inquiry']}' AND sign_id='{r['result']}';")

    sql.append("2.回退语句")
    if cfg.get('MERGE_MODULES', {}).get('unit', True):
        def rb_parts_from(r):
            rb = []
            if update_drafting:
                odid = r.get('orig_drafting_id', old_drafting_id)
                odnm = r.get('orig_drafting_name', old_drafting_name)
                if not odid and odnm:
                    code = extract_company_code(odnm)
                    if code:
                        odid = code
                if odid:
                    rb.append(f"PRIMARY_CONTRACT_DRAFTING_UNIT='{odid}'")
                if odnm:
                    rb.append(
                        f"PRIMARY_CONTRACT_DRAFTING_UNIT_NAME='{extract_company_name(odnm)}'")
            if update_contract_party:
                opid = r.get('orig_party_id', old_party_id)
                opnm = r.get('orig_party_name', old_party_name)
                if not opid and opnm:
                    code = extract_company_code(opnm)
                    if code:
                        opid = code
                if opid:
                    rb.append(f"CONTRACT_PARTY_ID='{opid}'")
                if opnm:
                    rb.append(
                        f"CONTRACT_PARTY_NAME='{extract_company_name(opnm)}'")
            return ', '.join(rb)
        groups_rb = merge_by_key(records, rb_parts_from)
        for rb_clause, recs in groups_rb.items():
            if not rb_clause:
                continue
            schemes = [r['scheme'] for r in recs if r.get('scheme')]
            for chunk in chunk_list(schemes, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
                sql.append(
                    f"update tprfa03 set {rb_clause}, OPS_REMARK='' where PURCHASE_SCHEME_NO IN ({format_in(chunk)});")
            conds_xj = [{'PURCHASE_SCHEME_NO': r.get(
                'scheme'), 'INQ_ID': r.get('inquiry')} for r in recs]
            conds_xj = [
                c for c in conds_xj if c['PURCHASE_SCHEME_NO'] and c['INQ_ID']]
            if conds_xj:
                where_or = compose_or(
                    conds_xj, ['PURCHASE_SCHEME_NO', 'INQ_ID'])
                sql.append(
                    f"update tprxj05 set {rb_clause}, OPS_REMARK='' where {where_or};")
            conds_nq = [{'PURCHASE_SCHEME_NO': r.get('scheme'), 'INQ_ID': r.get(
                'inquiry'), 'sign_id': r.get('result')} for r in recs]
            conds_nq = [c for c in conds_nq if c['PURCHASE_SCHEME_NO']
                        and c['INQ_ID'] and c['sign_id']]
            if conds_nq:
                where_or = compose_or(
                    conds_nq, ['PURCHASE_SCHEME_NO', 'INQ_ID', 'sign_id'])
                sql.append(
                    f"update tprnq03 set {rb_clause}, OPS_REMARK='' where {where_or};")
    else:
        for r in records:
            rb_parts = []
            if update_drafting:
                odid = r.get('orig_drafting_id', old_drafting_id)
                odnm = r.get('orig_drafting_name', old_drafting_name)
                if not odid and odnm:
                    code = extract_company_code(odnm)
                    if code:
                        odid = code
                if odid:
                    rb_parts.append(f"PRIMARY_CONTRACT_DRAFTING_UNIT='{odid}'")
                if odnm:
                    rb_parts.append(
                        f"PRIMARY_CONTRACT_DRAFTING_UNIT_NAME='{extract_company_name(odnm)}'")
            if update_contract_party:
                opid = r.get('orig_party_id', old_party_id)
                opnm = r.get('orig_party_name', old_party_name)
                if not opid and opnm:
                    code = extract_company_code(opnm)
                    if code:
                        opid = code
                if opid:
                    rb_parts.append(f"CONTRACT_PARTY_ID='{opid}'")
                if opnm:
                    rb_parts.append(
                        f"CONTRACT_PARTY_NAME='{extract_company_name(opnm)}'")
            if rb_parts:
                rb_clause = ', '.join(rb_parts)
                sql.append(
                    f"update tprfa03 set {rb_clause}, OPS_REMARK='' where PURCHASE_SCHEME_NO='{r['scheme']}';")
                sql.append(
                    f"update tprxj05 set {rb_clause}, OPS_REMARK='' where PURCHASE_SCHEME_NO='{r['scheme']}' AND INQ_ID='{r['inquiry']}';")
                sql.append(
                    f"update tprnq03 set {rb_clause}, OPS_REMARK='' where PURCHASE_SCHEME_NO='{r['scheme']}' AND INQ_ID='{r['inquiry']}' AND sign_id='{r['result']}';")

    sql.append("3.数据库")
    sql.append("ip：192.168.11.76")
    sql.append("库名：cnnc_ph")
    return "\n".join(sql)


def parse_unit_excel(file):
    """解析单位修改Excel"""
    wb = openpyxl.load_workbook(file)
    ws = wb.active
    records = []
    headers = [cell.value for cell in ws[1]]
    idx = {str(h): i for i, h in enumerate(headers)}

    scheme_alts = ['scheme_no', 'PURCHASE_SCHEME_NO',
                   'purchase_scheme_no', '采购方案编号']
    inquiry_alts = ['inquiry_no', 'inq_id', 'INQ_ID', '询价单编号']
    result_alts = ['result_no', 'sign_id', 'SIGN_ID', '定标结果编号']

    def pick(alts):
        for a in alts:
            if a in idx:
                return idx[a]
        return None

    s_idx = pick(scheme_alts)
    i_idx = pick(inquiry_alts)
    r_idx = pick(result_alts)
    if s_idx is None or i_idx is None or r_idx is None:
        raise ValueError("Excel 缺少必要列：方案、询价、定标")

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        rec = {
            'scheme': row[s_idx],
            'inquiry': row[i_idx],
            'result': row[r_idx]
        }
        if not all(rec.values()):
            continue

        # 提取新起草单位信息
        for k in ['new_drafting_id', '新起草单位ID']:
            if k in idx:
                rec['new_drafting_id'] = row[idx[k]]
                break
        for k in ['new_drafting_name', '新起草单位名称', '组织机构名称']:
            if k in idx:
                rec['new_drafting_name'] = row[idx[k]]
                break
        # 提取新签约主体信息
        for k in ['new_party_id', '新签约主体ID']:
            if k in idx:
                rec['new_party_id'] = row[idx[k]]
                break
        for k in ['new_party_name', '新签约主体名称']:
            if k in idx:
                rec['new_party_name'] = row[idx[k]]
                break
        # 提取原起草单位信息
        for k in ['orig_drafting_id', '原起草单位ID']:
            if k in idx:
                rec['orig_drafting_id'] = row[idx[k]]
                break
        for k in ['orig_drafting_name', '原起草单位名称']:
            if k in idx:
                rec['orig_drafting_name'] = row[idx[k]]
                break
        # 提取原签约主体信息
        for k in ['orig_party_id', '原签约主体ID']:
            if k in idx:
                rec['orig_party_id'] = row[idx[k]]
                break
        for k in ['orig_party_name', '原签约主体名称']:
            if k in idx:
                rec['orig_party_name'] = row[idx[k]]
                break

        records.append(rec)
    return records


def unit_change_view(request):
    """合同单位修改视图"""
    saved_file = None
    validation_failure = None  # 新增

    if request.method == 'POST':
        form = UnitChangeForm(request.POST, request.FILES)
        if form.is_valid():
            cd = form.cleaned_data
            log_view_input("合同起草、签约单位修改", cd)

            if cd.get('excel_file'):
                records = parse_unit_excel(cd['excel_file'])
                
                # 执行数据校验
                validation_result = validate_unit_change_records(records)
                
                if not validation_result['valid']:
                    # 校验失败,生成失败文件
                    temp_file = generate_validation_failure_excel(
                        cd['excel_file'],
                        validation_result,
                        'contract_unit'
                    )
                    
                    if temp_file:
                        validation_failure = {
                            'total': validation_result['total'],
                            'passed': validation_result['passed'],
                            'failed': validation_result['failed'],
                            'filename': get_temp_filename_from_path(temp_file)
                        }
                    
                    return render(request, 'unit_change_form.html', {
                        'form': form,
                        'validation_failure': validation_failure,
                        'active_menu': 'unit_change',
                        'sidebar_groups': get_sidebar_groups(),
                    })
                
                # 校验原单位名称是否存在
                unit_name_errors = []
                for idx, record in enumerate(records):
                    row_number = idx + 2  # 跳过表头
                    errors = validate_original_unit_names(
                        old_drafting_name=record.get('orig_drafting_name'),
                        old_party_name=record.get('orig_party_name')
                    )
                    if errors:
                        for error in errors:
                            unit_name_errors.append(f'第{row_number}行: {error}')
                
                if unit_name_errors:
                    for error in unit_name_errors:
                        form.add_error(None, error)
                    return render(request, 'unit_change_form.html', {
                        'form': form,
                        'active_menu': 'unit_change',
                        'sidebar_groups': get_sidebar_groups(),
                    })
                
                # 从记录中判断是否更新
                will_update_drafting = any(r.get('new_drafting_id') or r.get(
                    'new_drafting_name') for r in records)
                will_update_party = any(r.get('new_party_id') or r.get(
                    'new_party_name') for r in records)
                # 统一处理名称到编号的转换
                for r in records:
                    for k_name, k_id in [
                        ('new_drafting_name', 'new_drafting_id'),
                        ('orig_drafting_name', 'orig_drafting_id'),
                        ('new_party_name', 'new_party_id'),
                        ('orig_party_name', 'orig_party_id')
                    ]:
                        if r.get(k_name) and not r.get(k_id):
                            code = extract_company_code(str(r[k_name]))
                            if code:
                                r[k_id] = code
                            else:
                                uc = _unique_code_by_name(str(r[k_name]))
                                if uc:
                                    r[k_id] = uc
            else:
                # 单条记录模式 - 添加简单直接的校验
                
                # 校验1: 至少填写一个新单位信息
                has_new_unit = bool(cd.get('new_drafting_name') or cd.get('new_party_name') or 
                                   cd.get('new_drafting_id') or cd.get('new_party_id'))
                
                if not has_new_unit:
                    form.add_error(None, '新起草单位名称，新签约主体名称至少有一项必填')
                    return render(request, 'unit_change_form.html', {
                        'form': form,
                        'active_menu': 'unit_change',
                        'sidebar_groups': get_sidebar_groups(),
                    })
                
                # 校验2: 至少填写一个标识字段
                has_identifier = bool(cd.get('scheme_no') or cd.get('inquiry_no') or cd.get('result_no'))
                
                if not has_identifier:
                    form.add_error(None, '采购方案编号，询价单编号，定标结果编号至少有一项必填')
                    return render(request, 'unit_change_form.html', {
                        'form': form,
                        'active_menu': 'unit_change',
                        'sidebar_groups': get_sidebar_groups(),
                    })
                
                # 校验3: 检查原单位名称是否在数据库中存在
                unit_name_errors = validate_original_unit_names(
                    old_drafting_name=cd.get('old_drafting_name'),
                    old_party_name=cd.get('old_party_name')
                )
                
                if unit_name_errors:
                    for error in unit_name_errors:
                        form.add_error(None, error)
                    return render(request, 'unit_change_form.html', {
                        'form': form,
                        'active_menu': 'unit_change',
                        'sidebar_groups': get_sidebar_groups(),
                    })
                
                records = [{
                    'scheme': cd['scheme_no'],
                    'inquiry': cd['inquiry_no'],
                    'result': cd['result_no']
                }]
                will_update_drafting = bool(
                    cd.get('new_drafting_id') or cd.get('new_drafting_name'))
                will_update_party = bool(
                    cd.get('new_party_id') or cd.get('new_party_name'))
                # 单条记录的名称到编号转换
                for k_name, k_id in [
                    ('new_drafting_name', 'new_drafting_id'),
                    ('old_drafting_name', 'old_drafting_id'),
                    ('new_party_name', 'new_party_id'),
                    ('old_party_name', 'old_party_id')
                ]:
                    if cd.get(k_name) and not cd.get(k_id):
                        code = extract_company_code(cd[k_name])
                        if code:
                            cd[k_id] = code
                        else:
                            uc = _unique_code_by_name(cd[k_name])
                            if uc:
                                cd[k_id] = uc

            ops_remark = parse_ops_remark(cd.get('ops_remark', ''))
            sql_content = generate_unit_sql(
                records,
                update_drafting=will_update_drafting,
                update_contract_party=will_update_party,
                old_drafting_id=cd.get('old_drafting_id'),
                old_drafting_name=cd.get('old_drafting_name'),
                new_drafting_id=cd.get('new_drafting_id'),
                new_drafting_name=cd.get('new_drafting_name'),
                old_party_id=cd.get('old_party_id'),
                old_party_name=cd.get('old_party_name'),
                new_party_id=cd.get('new_party_id'),
                new_party_name=cd.get('new_party_name'),
                ops_remark=ops_remark
            )

            # 确定文件名前缀
            if will_update_drafting and will_update_party:
                suffix = "修改起草单位和签约主体"
            elif will_update_drafting:
                suffix = "修改起草单位"
            elif will_update_party:
                suffix = "修改签约主体"
            else:
                suffix = "修改单位信息"

            saved_file = save_sql_file(
                sql_content, suffix, cd.get('dynamic_id'))
            logger.info(f"SQL文件已保存: {saved_file}")

            # 保存会话数据
            session_data = {}
            for k, v in cd.items():
                if k != 'excel_file':
                    if hasattr(v, 'to_eng_string'):
                        session_data[k] = v.to_eng_string()
                    else:
                        session_data[k] = v
            request.session['unit_change_last'] = session_data
    else:
        if request.GET.get('clear'):
            request.session.pop('unit_change_last', None)
            form = UnitChangeForm()
        else:
            initial = request.session.get('unit_change_last')
            form = UnitChangeForm(initial=initial)

    return render(request, 'unit_change_form.html', {
        'form': form,
        'saved_file': saved_file,
        'validation_failure': validation_failure,  # 新增
        'active_menu': 'unit_change',
        'sidebar_groups': get_sidebar_groups(),
    })


def download_unit_template(request):
    """下载单位修改Excel模板"""
    if not OPENPYXL_AVAILABLE:
        from django.http import HttpResponse
        return HttpResponse("缺少 openpyxl，请安装后使用模板下载：pip install openpyxl", status=500)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '模板'
    ws.append(['采购方案编号', '询价单编号', '定标结果编号', '新起草单位ID', '新起草单位名称',
              '新签约主体ID', '新签约主体名称', '原起草单位ID', '原起草单位名称', '原签约主体ID', '原签约主体名称'])
    ws.append(['CNEC-EXAMPLE-001', 'XJD-EXAMPLE-001', '1',
              None, None, None, None, None, None, None, None])
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    response = FileResponse(bio, as_attachment=True,
                            filename='unit_change_template.xlsx')
    response['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    return response


__all__ = ['unit_change_view', 'download_unit_template',
           'generate_unit_sql', 'parse_unit_excel']
