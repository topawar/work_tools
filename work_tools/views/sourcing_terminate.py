""" 
终止简化寻源合同模块
包含终止简化寻源合同的所有视图函数和SQL生成逻辑
"""
import io
import os
import logging
from django.shortcuts import render
from django.http import FileResponse
from ..forms import SourcingTerminateForm, BID_STATUS_CHOICES
from ..navigation import get_sidebar_groups
from ..config import get_config
from ..sql_merge import chunk_list, format_in, merge_by_key
from ..logger_utils import log_view_input
from .base import parse_ops_remark, save_sql_file
from ..validation_utils import (
    validate_required,
    generate_validation_failure_excel,
    get_temp_filename_from_path
)

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

logger = logging.getLogger('work_tools.view')
sql_logger = logging.getLogger('work_tools.sql')


def validate_sourcing_terminate_records(records):
    """校验终止简化寻源合同记录"""
    results = []
    passed_count = 0
    failed_count = 0
    
    for idx, record in enumerate(records):
        row_number = idx + 2  # 跳过表头
        errors = []
        
        # 必填项校验:合同ID必填
        error = validate_required(record.get('bpo_id'), '合同ID')
        if error:
            errors.append(error)
        
        # 记录结果
        if errors:
            results.append({
                'row_number': row_number,
                'valid': False,
                'errors': errors
            })
            failed_count += 1
        else:
            results.append({
                'row_number': row_number,
                'valid': True,
                'errors': []
            })
            passed_count += 1
    
    return {
        'valid': failed_count == 0,
        'total': len(records),
        'passed': passed_count,
        'failed': failed_count,
        'results': results
    }


def parse_sourcing_terminate_excel(file):
    """解析Excel文件,提取终止简化寻源合同数据"""
    if not OPENPYXL_AVAILABLE:
        raise ValueError("缺少 openpyxl，无法解析 Excel")

    # 从数据库加载中标状态配置构建映射字典
    try:
        from ..dropdown_utils import get_dropdown_options
        options = get_dropdown_options('bid_status', include_empty=False)
        SOURCING_STATUS_MAP = {label: code for code, label in options if code}
    except Exception:
        # 回退到硬编码映射
        SOURCING_STATUS_MAP = {v: k for k, v in BID_STATUS_CHOICES if k}

    wb = openpyxl.load_workbook(file)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    idx = {str(h): i for i, h in enumerate(headers)}

    def pick(alts):
        for a in alts:
            if a in idx:
                return idx[a]
        return None

    bpo_id_idx = pick(['合同ID', 'BPO_ID', 'bpo_id'])
    package_idx = pick(['采购包编号', 'purchase_package_no', 'PURCHASE_PACKAGE_NO'])
    orig_status_idx = pick(['原简化寻源状态', '原状态', 'orig_sourcing_status'])

    if bpo_id_idx is None or package_idx is None:
        raise ValueError("Excel 缺少必要列：合同ID/采购包编号")

    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        bpo_id = row[bpo_id_idx]
        package_no = row[package_idx]
        if bpo_id is None or package_no is None:
            continue

        rec = {
            'bpo_id': str(bpo_id).strip(),
            'purchase_package_no': str(package_no).strip(),
        }

        if orig_status_idx is not None:
            orig_status = row[orig_status_idx]
            if orig_status:
                orig_status_val = str(orig_status).strip()
                if orig_status_val in SOURCING_STATUS_MAP:
                    orig_status_val = SOURCING_STATUS_MAP[orig_status_val]
                rec['orig_sourcing_status'] = orig_status_val

        records.append(rec)

    if not records:
        raise ValueError("Excel 中没有有效的行数据")

    logger.info(f"[终止简化寻源] Excel解析完成, 记录数={len(records)}")
    return records


def generate_sourcing_terminate_sql_bulk(records, ops_remark=None):
    """
    生成批量终止简化寻源合同SQL - 支持SQL合并策略
    包括删除合同和回退简化寻源状态两部分
    """
    cfg = get_config()
    sql = []
    sql.append("1、执行语句")

    sql_logger.info(
        f"[SQL合并] 开始生成终止简化寻源SQL, 记录数={len(records)}, 配置启用={cfg.get('MERGE_MODULES', {}).get('sourcing_terminate', True)}")

    if cfg.get('MERGE_MODULES', {}).get('sourcing_terminate', True):
        # 启用SQL合并策略: 合并删除合同语句
        bpo_ids = [r['bpo_id'] for r in records if r.get('bpo_id')]
        sql_logger.info(f"[SQL合并] 合并删除合同, 合同数={len(bpo_ids)}")

        for chunk in chunk_list(bpo_ids, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
            inlist = format_in(chunk)
            sql.append("--删除合同")
            # 删除主表
            stmt_main = (
                f"UPDATE tphct01 SET alive_flag='0', APPR_STATE='DELETE', "
                f"OPS_REMARK='{ops_remark or ''}' "
                f"WHERE BPO_ID IN ({inlist});"
            )
            sql.append(stmt_main)

            # 删除明细表
            stmt_detail = (
                f"UPDATE tphct02 SET alive_flag='0', APPR_STATE='DELETE', "
                f"OPS_REMARK='{ops_remark or ''}' "
                f"WHERE BPO_ID IN ({inlist});"
            )
            sql.append(stmt_detail)
            sql.append("")

        # 合并回退简化寻源语句
        package_nos = [r['purchase_package_no']
                       for r in records if r.get('purchase_package_no')]
        sql_logger.info(f"[SQL合并] 合并回退简化寻源, 采购包数={len(package_nos)}")

        for chunk in chunk_list(package_nos, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
            inlist = format_in(chunk)
            sql.append(
                f"UPDATE tprlyjy01 SET status='40', OPS_REMARK='{ops_remark or ''}' WHERE purchase_package_no IN ({inlist});")
            sql.append(
                f"UPDATE tprlyjy02 SET status='40', OPS_REMARK='{ops_remark or ''}' WHERE purchase_package_no IN ({inlist});")
            sql.append("")

    else:
        # 未启用合并策略: 逐条生成SQL
        sql_logger.info("[SQL合并] 合并策略已禁用, 逐条生成SQL")
        for r in records:
            bpo_id = r.get('bpo_id')
            package_no = r.get('purchase_package_no')

            if bpo_id:
                sql.append("--删除合同")
                sql.append(
                    f"UPDATE tphct01 SET alive_flag='0', APPR_STATE='DELETE', OPS_REMARK='{ops_remark or ''}' WHERE BPO_ID='{bpo_id}';")
                sql.append(
                    f"UPDATE tphct02 SET alive_flag='0', APPR_STATE='DELETE', OPS_REMARK='{ops_remark or ''}' WHERE BPO_ID='{bpo_id}';")
                sql.append("")

            if package_no:
                sql.append(
                    f"UPDATE tprlyjy01 SET status='40', OPS_REMARK='{ops_remark or ''}' WHERE purchase_package_no='{package_no}';")
                sql.append(
                    f"UPDATE tprlyjy02 SET status='40', OPS_REMARK='{ops_remark or ''}' WHERE purchase_package_no='{package_no}';")
                sql.append("")

    # 生成回退语句
    sql.append("2、回退语句")

    if cfg.get('MERGE_MODULES', {}).get('sourcing_terminate', True):
        # 按原状态分组回退简化寻源
        def rb_key(r):
            return r.get('orig_sourcing_status', '50')  # 默认回退到已签约

        groups = merge_by_key(records, rb_key)
        sql_logger.info(f"[SQL合并] 回退语句分组完成, 组数={len(groups)}")

        for orig_status, recs in groups.items():
            package_nos = [r['purchase_package_no']
                           for r in recs if r.get('purchase_package_no')]

            for chunk in chunk_list(package_nos, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
                inlist = format_in(chunk)
                sql.append(
                    f"UPDATE tprlyjy01 SET status='{orig_status}', OPS_REMARK='' WHERE purchase_package_no IN ({inlist});")
                sql.append(
                    f"UPDATE tprlyjy02 SET status='{orig_status}', OPS_REMARK='' WHERE purchase_package_no IN ({inlist});")

        sql.append("")

        # 合并恢复合同语句
        bpo_ids = [r['bpo_id'] for r in records if r.get('bpo_id')]
        for chunk in chunk_list(bpo_ids, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
            inlist = format_in(chunk)
            sql.append(
                f"UPDATE tphct01 SET alive_flag='1', APPR_STATE='ACTIVE', OPS_REMARK='' WHERE BPO_ID IN ({inlist});")
            sql.append(
                f"UPDATE tphct02 SET alive_flag='1', APPR_STATE='ACTIVE', OPS_REMARK='' WHERE BPO_ID IN ({inlist});")

    else:
        # 逐条生成回退语句
        for r in records:
            package_no = r.get('purchase_package_no')
            bpo_id = r.get('bpo_id')
            orig_status = r.get('orig_sourcing_status', '50')

            if package_no:
                sql.append(
                    f"UPDATE tprlyjy01 SET status='{orig_status}', OPS_REMARK='' WHERE purchase_package_no='{package_no}';")
                sql.append(
                    f"UPDATE tprlyjy02 SET status='{orig_status}', OPS_REMARK='' WHERE purchase_package_no='{package_no}';")
                sql.append("")

            if bpo_id:
                sql.append(
                    f"UPDATE tphct01 SET alive_flag='1', APPR_STATE='ACTIVE', OPS_REMARK='' WHERE BPO_ID='{bpo_id}';")
                sql.append(
                    f"UPDATE tphct02 SET alive_flag='1', APPR_STATE='ACTIVE', OPS_REMARK='' WHERE BPO_ID='{bpo_id}';")
                sql.append("")

    return sql


def sourcing_terminate_view(request):
    """终止简化寻源合同视图"""
    saved_file = None
    validation_failure = None  # 新增

    if request.method == 'POST':
        form = SourcingTerminateForm(request.POST, request.FILES)
        if form.is_valid():
            try:
                cd = form.cleaned_data
                ops_remark = parse_ops_remark(cd.get('ops_remark', ''))
                dynamic_id = cd.get('dynamic_id', '')

                # 记录视图输入
                log_view_input('sourcing_terminate', cd)

                logger.info(
                    f"[终止简化寻源] 开始处理, dynamic_id={dynamic_id}, ops_remark={ops_remark}")

                # 解析数据
                excel_file = cd.get('excel_file')
                if excel_file:
                    records = parse_sourcing_terminate_excel(excel_file)
                    logger.info(f"[终止简化寻源] Excel模式, 记录数={len(records)}")
                    
                    # 执行数据校验
                    validation_result = validate_sourcing_terminate_records(records)
                    
                    if not validation_result['valid']:
                        # 校验失败,生成失败文件
                        temp_file = generate_validation_failure_excel(
                            excel_file,
                            validation_result,
                            'sourcing_terminate'
                        )
                        
                        if temp_file:
                            validation_failure = {
                                'total': validation_result['total'],
                                'passed': validation_result['passed'],
                                'failed': validation_result['failed'],
                                'filename': get_temp_filename_from_path(temp_file)
                            }
                        
                        return render(request, 'sourcing_terminate.html', {
                            'form': form,
                            'validation_failure': validation_failure,
                            'active_menu': 'sourcing_terminate',
                            'sidebar_groups': get_sidebar_groups(),
                        })
                else:
                    # 单条记录
                    rec = {
                        'bpo_id': cd.get('bpo_id'),
                        'purchase_package_no': cd.get('purchase_package_no'),
                    }
                    if cd.get('orig_sourcing_status'):
                        rec['orig_sourcing_status'] = cd.get(
                            'orig_sourcing_status')
                    records = [rec]
                    logger.info(
                        f"[终止简化寻源] 单条模式, bpo_id={rec['bpo_id']}, package_no={rec['purchase_package_no']}")

                # 生成SQL
                sql_lines = generate_sourcing_terminate_sql_bulk(
                    records, ops_remark)
                logger.info(f"[终止简化寻源] SQL生成完成, 行数={len(sql_lines)}")

                # 保存SQL文件
                sql_content = '\n'.join(sql_lines)
                saved_file = save_sql_file(
                    sql_content, prefix='终止简化寻源合同', dynamic_id=dynamic_id)
                logger.info(f"[终止简化寻源] SQL文件已保存: {saved_file}")

            except Exception as e:
                logger.error(f"[终止简化寻源] 处理失败: {e}", exc_info=True)
                form.add_error(None, f"生成SQL失败: {str(e)}")
    else:
        form = SourcingTerminateForm()

    return render(request, 'sourcing_terminate.html', {
        'form': form,
        'saved_file': saved_file,
        'validation_failure': validation_failure,  # 新增
        'active_menu': 'sourcing_terminate',
        'sidebar_groups': get_sidebar_groups(),
    })


def download_sourcing_terminate_template(request):
    """下载终止简化寻源合同Excel模板"""
    if not OPENPYXL_AVAILABLE:
        return FileResponse(
            io.BytesIO(b"openpyxl not installed"),
            as_attachment=True,
            filename="error.txt"
        )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "终止简化寻源合同"

    # 设置表头
    headers = ["合同ID", "采购包编号", "原简化寻源状态"]
    header_fill = PatternFill(start_color="4472C4",
                              end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')

    # 设置列宽
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 20

    # 添加示例数据
    example_data = [
        ["HYY0-25-00144", "JYGS-CGB-25-00807", "50"],
        ["HYY0-25-00145", "JYGS-CGB-25-00808", "50"],
    ]

    for row_num, row_data in enumerate(example_data, 2):
        for col_num, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col_num)
            cell.value = value
            cell.alignment = Alignment(horizontal='left', vertical='center')

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)

    response = FileResponse(bio, as_attachment=True,
                            filename='sourcing_terminate_template.xlsx')
    response['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    return response


__all__ = [
    'sourcing_terminate_view',
    'download_sourcing_terminate_template',
]
