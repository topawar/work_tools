"""
合同物资编码修改模块
"""
import io
import logging
from django.shortcuts import render
from django.http import FileResponse
from ..forms import ContractItemUpdateForm
from ..navigation import SIDEBAR_GROUPS
from ..models import ItemDetail
from ..config import get_config
from ..sql_merge import chunk_list, format_in, merge_by_key
from .base import parse_ops_remark, normalize_item_id, save_sql_file
from ..validation_utils import (
    validate_required,
    generate_validation_failure_excel,
    get_temp_filename_from_path,
)

try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

logger = logging.getLogger('work_tools.view')


def validate_contract_item_records(records):
    """
    校验物资编码修改记录
    
    Args:
        records: 解析后的记录列表
        
    Returns:
        校验结果字典
    """
    results = []
    passed_count = 0
    failed_count = 0
    
    for idx, record in enumerate(records):
        row_number = idx + 2  # 从第2行开始（第1行是表头）
        errors = []
        
        # 必填项校验：明细行ID
        error = validate_required(record.get('line_id'), '明细行ID')
        if error:
            errors.append(error)
        
        # 记录校验结果
        if errors:
            results.append({
                'row_number': row_number,
                'valid': False,
                'errors': errors
            })
            failed_count += 1
        else:
            results.append({
                'row_number': row_number,
                'valid': True,
                'errors': []
            })
            passed_count += 1
    
    return {
        'valid': failed_count == 0,
        'total': len(records),
        'passed': passed_count,
        'failed': failed_count,
        'results': results
    }


def parse_item_excel(file):
    """解析物资编码Excel"""
    wb = openpyxl.load_workbook(file)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    idx = {str(h): i for i, h in enumerate(headers)}

    def pick(alts):
        for a in alts:
            if a in idx:
                return idx[a]
        return None

    lid_idx = pick(['line_id', 'BPO_LINE_ID', '明细行ID'])
    new_id_idx = pick(['ITEM_ID', 'item_id', '物资编码'])
    new_name_idx = pick(['ITEM_NAME', 'item_name', '物资名称'])
    new_uom_idx = pick(['ITEM_UOM', 'item_uom', '计量单位'])
    new_cat_idx = pick(['CATEGORY', 'category', '物资分类编码'])
    orig_id_idx = pick(['orig_item_id', '原物资编码'])
    orig_name_idx = pick(['orig_item_name', '原物资名称'])
    orig_uom_idx = pick(['orig_item_uom', '原计量单位'])
    orig_cat_idx = pick(['orig_category', '原物资分类编码'])

    if lid_idx is None:
        raise ValueError("Excel 缺少必要列：明细行ID")

    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        lid = row[lid_idx]
        # 不过滤空行，留给校验函数处理
        rec = {'line_id': str(lid).strip() if lid is not None else None}
        if new_id_idx is not None:
            rec['new_item_id'] = normalize_item_id(row[new_id_idx])
        if new_name_idx is not None:
            rec['new_item_name'] = row[new_name_idx]
        if new_uom_idx is not None:
            rec['new_item_uom'] = row[new_uom_idx]
        if new_cat_idx is not None:
            rec['new_category'] = row[new_cat_idx]
        if orig_id_idx is not None:
            rec['orig_item_id'] = normalize_item_id(row[orig_id_idx])
        if orig_name_idx is not None:
            rec['orig_item_name'] = row[orig_name_idx]
        if orig_uom_idx is not None:
            rec['orig_item_uom'] = row[orig_uom_idx]
        if orig_cat_idx is not None:
            rec['orig_category'] = row[orig_cat_idx]
        records.append(rec)

    if not records:
        raise ValueError("Excel 中没有有效的行数据")
    return records


def generate_item_sql_bulk(records, ops_remark=None):
    """生成物资编码修改SQL"""
    cfg = get_config()
    sql = []
    sql.append("1、执行语句")
    
    logger.info(f"[SQL合并] 开始生成物资编码SQL, 记录数={len(records)}, 配置启用={cfg.get('MERGE_MODULES', {}).get('item', True)}")
    
    if cfg.get('MERGE_MODULES', {}).get('item', True):
        # SQL合并模式
        def key_fn(r):
            return (
                str(r.get('new_item_id') or '').strip(),
                str(r.get('new_item_name') or '').strip(),
                str(r.get('new_item_uom') or '').strip(),
                str(r.get('new_category') or '').strip(),
            )
        
        groups = merge_by_key(records, key_fn)
        logger.info(f"[SQL合并] 物资编码分组完成, 组数={len(groups)}")
        
        for k, recs in groups.items():
            if not any(k):
                continue
            set_parts = []
            if k[0]:
                set_parts.append(f"ITEM_ID = '{k[0]}'")
            if k[1]:
                set_parts.append(f"ITEM_NAME = '{k[1]}'")
            if k[2]:
                set_parts.append(f"ITEM_UOM = '{k[2]}'")
            if k[3]:
                set_parts.append(f"CATEGORY = '{k[3]}'")
            
            ids = [r['line_id'] for r in recs if r.get('line_id')]
            for chunk in chunk_list(ids, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
                inlist = format_in(chunk)
                stmt = (
                    "UPDATE TPHCT02 SET " + ", ".join(set_parts) + ", "
                    f"OPS_REMARK = '{ops_remark or ''}' WHERE BPO_LINE_ID IN (" +
                    inlist + ") AND ALIVE_FLAG = '1';"
                )
                sql.append(stmt)
    else:
        # 非合并模式
        for r in records:
            parts = []
            if r.get('new_item_id'):
                parts.append(f"ITEM_ID = '{str(r['new_item_id']).strip()}'")
            if r.get('new_item_name'):
                parts.append(f"ITEM_NAME = '{str(r['new_item_name']).strip()}'")
            if r.get('new_item_uom'):
                parts.append(f"ITEM_UOM = '{str(r['new_item_uom']).strip()}'")
            if r.get('new_category'):
                parts.append(f"CATEGORY = '{str(r['new_category']).strip()}'")
            if parts:
                stmt = (
                    "UPDATE TPHCT02 "
                    "SET " + ", ".join(parts) + ", "
                    f"OPS_REMARK = '{ops_remark or ''}' "
                    f"WHERE BPO_LINE_ID = '{r['line_id']}' "
                    "AND ALIVE_FLAG = '1';"
                )
                sql.append(stmt)

    sql.append("")
    sql.append("2、回退语句")
    
    if cfg.get('MERGE_MODULES', {}).get('item', True):
        # SQL合并模式
        def rb_key(r):
            return (
                str(r.get('orig_item_id') or '').strip(),
                str(r.get('orig_item_name') or '').strip(),
                str(r.get('orig_item_uom') or '').strip(),
                str(r.get('orig_category') or '').strip(),
            )
        
        groups = merge_by_key(records, rb_key)
        for k, recs in groups.items():
            if not any(k):
                continue
            rb = []
            if k[0]:
                rb.append(f"ITEM_ID = '{k[0]}'")
            if k[1]:
                rb.append(f"ITEM_NAME = '{k[1]}'")
            if k[2]:
                rb.append(f"ITEM_UOM = '{k[2]}'")
            if k[3]:
                rb.append(f"CATEGORY = '{k[3]}'")
            
            ids = [r['line_id'] for r in recs if r.get('line_id')]
            for chunk in chunk_list(ids, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
                inlist = format_in(chunk)
                stmt = (
                    "UPDATE TPHCT02 SET " + ", ".join(rb) + ", "
                    f"OPS_REMARK = '' WHERE BPO_LINE_ID IN (" +
                    inlist + ") AND ALIVE_FLAG = '1';"
                )
                sql.append(stmt)
    else:
        # 非合并模式
        for r in records:
            rb = []
            if r.get('orig_item_id'):
                rb.append(f"ITEM_ID = '{str(r['orig_item_id']).strip()}'")
            if r.get('orig_item_name'):
                rb.append(f"ITEM_NAME = '{str(r['orig_item_name']).strip()}'")
            if r.get('orig_item_uom'):
                rb.append(f"ITEM_UOM = '{str(r['orig_item_uom']).strip()}'")
            if r.get('orig_category'):
                rb.append(f"CATEGORY = '{str(r['orig_category']).strip()}'")
            if rb:
                stmt = (
                    "UPDATE TPHCT02 "
                    "SET " + ", ".join(rb) + ", "
                    f"OPS_REMARK = '' "
                    f"WHERE BPO_LINE_ID = '{r['line_id']}' "
                    "AND ALIVE_FLAG = '1';"
                )
                sql.append(stmt)

    sql.append("")
    sql.append("3、数据库")
    sql.append("ip：192.168.11.71")
    sql.append("库名：cnnc_ph")
    return "\n".join(sql)


def contract_item_update_view(request):
    """合同物资编码修改视图"""
    saved_file = None
    validation_failure = None

    if request.method == 'POST':
        form = ContractItemUpdateForm(request.POST, request.FILES)
        if form.is_valid():
            cd = form.cleaned_data
            ops_remark = parse_ops_remark(cd.get('ops_remark', ''))

            if cd.get('excel_file'):
                records = parse_item_excel(cd['excel_file'])
                
                # 执行数据校验
                validation_result = validate_contract_item_records(records)
                
                if not validation_result['valid']:
                    # 校验失败，生成包含错误信息的Excel文件
                    temp_file = generate_validation_failure_excel(
                        cd['excel_file'], 
                        validation_result, 
                        'contract_item'
                    )
                    
                    if temp_file:
                        validation_failure = {
                            'total': validation_result['total'],
                            'passed': validation_result['passed'],
                            'failed': validation_result['failed'],
                            'filename': get_temp_filename_from_path(temp_file)
                        }
                        logger.info(f"物资编码校验失败: 总行数={validation_failure['total']}, 失败行数={validation_failure['failed']}")
                    
                    # 不生成SQL，直接返回页面显示错误
                    return render(request, 'contract_item_form.html', {
                        'form': form,
                        'validation_failure': validation_failure,
                        'active_menu': 'contract_item',
                        'sidebar_groups': SIDEBAR_GROUPS,
                    })
            else:
                rec = {
                    'line_id': cd['single_line_id'],
                    'new_item_id': normalize_item_id(cd.get('new_item_id')),
                    'new_item_name': cd.get('new_item_name'),
                    'new_item_uom': cd.get('new_item_uom'),
                    'new_category': cd.get('new_category'),
                    'orig_item_id': normalize_item_id(cd.get('orig_item_id')),
                    'orig_item_name': cd.get('orig_item_name'),
                    'orig_item_uom': cd.get('orig_item_uom'),
                    'orig_category': cd.get('orig_category'),
                }
                records = [rec]

            # 从数据库补充物资信息
            for r in records:
                for prefix in ['new', 'orig']:
                    iid = normalize_item_id(r.get(f'{prefix}_item_id'))
                    if iid and (not r.get(f'{prefix}_item_name') or not r.get(f'{prefix}_item_uom') or not r.get(f'{prefix}_category')):
                        obj = ItemDetail.objects.filter(item_id=iid).first()
                        if obj:
                            if not r.get(f'{prefix}_item_name'):
                                r[f'{prefix}_item_name'] = obj.item_name
                            if not r.get(f'{prefix}_item_uom'):
                                r[f'{prefix}_item_uom'] = obj.item_uom
                            if not r.get(f'{prefix}_category'):
                                r[f'{prefix}_category'] = obj.category
                            r[f'{prefix}_item_id'] = iid

            sql_content = generate_item_sql_bulk(records, ops_remark)
            saved_file = save_sql_file(
                sql_content, '修改合同物资编码', cd.get('dynamic_id'))
            logger.info(f"SQL文件已保存: {saved_file}")

            # 保存会话数据
            session_data = {}
            for k, v in cd.items():
                if k != 'excel_file':
                    if hasattr(v, 'to_eng_string'):
                        session_data[k] = v.to_eng_string()
                    else:
                        session_data[k] = v
            request.session['contract_item_last'] = session_data
    else:
        if request.GET.get('clear'):
            request.session.pop('contract_item_last', None)
            form = ContractItemUpdateForm()
        else:
            initial = request.session.get('contract_item_last')
            form = ContractItemUpdateForm(initial=initial)

    return render(request, 'contract_item_form.html', {
        'form': form,
        'saved_file': saved_file,
        'validation_failure': validation_failure,
        'active_menu': 'contract_item',
        'sidebar_groups': SIDEBAR_GROUPS,
    })


def download_item_template(request):
    """下载物资编码Excel模板"""
    if not OPENPYXL_AVAILABLE:
        from django.http import HttpResponse
        return HttpResponse("缺少 openpyxl，请安装后使用模板下载：pip install openpyxl", status=500)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '模板'
    ws.append(['明细行ID', '物资编码', '物资名称', '计量单位', '物资分类编码',
              '原物资编码', '原物资名称', '原计量单位', '原物资分类编码', '描述'])
    ws.append(['BPO-EXAMPLE-000001', '01607733', '鞋套',
              '双', '430798', None, None, None, None, None])
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    response = FileResponse(bio, as_attachment=True,
                            filename='contract_item_template.xlsx')
    response['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    return response


# 别名兼容
contract_item_view = contract_item_update_view
download_contract_item_template = download_item_template

__all__ = [
    'contract_item_view',
    'contract_item_update_view',
    'download_contract_item_template',
    'download_item_template',
    'parse_item_excel',
    'generate_item_sql_bulk',
]
