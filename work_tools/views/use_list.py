"""
适用清单修改模块
修改tphct04表的组织机构信息
"""
import io
import logging
from django.shortcuts import render
from django.http import FileResponse
from ..forms import UseListUpdateForm
from ..navigation import SIDEBAR_GROUPS
from ..config import get_config
from ..sql_merge import chunk_list, format_in, merge_by_key
from ..models import OrgDetail
from .base import parse_ops_remark, extract_company_code, extract_company_name, save_sql_file

try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

logger = logging.getLogger('work_tools.view')


def _unique_code_by_name(name: str):
    """根据名称查询组织机构编码"""
    if not name:
        return None
    qs = OrgDetail.objects.filter(company_name__iexact=str(name).strip())
    if qs.count() == 1:
        obj = qs.first()
        return obj.company_code if obj and obj.company_code else None
    return None


def parse_use_list_excel(file):
    """解析Excel文件"""
    if not OPENPYXL_AVAILABLE:
        raise ValueError("缺少 openpyxl，无法解析 Excel")

    wb = openpyxl.load_workbook(file)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    idx = {str(h): i for i, h in enumerate(headers)}

    def pick(alts):
        for a in alts:
            if a in idx:
                return idx[a]
        return None

    bid_idx = pick(['BUSINESS_ID', 'business_id', '合同ID'])
    new_name_idx = pick(['新组织机构名称', 'new_org_name'])
    new_code_idx = pick(['新组织机构编码', 'new_org_code'])
    orig_name_idx = pick(['原组织机构名称', 'orig_org_name'])
    orig_code_idx = pick(['原组织机构编码', 'orig_org_code'])

    if bid_idx is None:
        raise ValueError("Excel 缺少必要列：BUSINESS_ID")

    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        bid = row[bid_idx]
        if bid is None:
            continue

        rec = {'business_id': str(bid).strip()}

        # 新值
        if new_name_idx is not None and row[new_name_idx]:
            new_name = str(row[new_name_idx]).strip()
            rec['new_org_name'] = extract_company_name(new_name) or new_name
            code = extract_company_code(new_name)
            if code:
                rec['new_org_code'] = code
            elif new_code_idx is not None and row[new_code_idx]:
                rec['new_org_code'] = str(row[new_code_idx]).strip()
            else:
                rec['new_org_code'] = _unique_code_by_name(rec['new_org_name'])

        # 原值
        if orig_name_idx is not None and row[orig_name_idx]:
            orig_name = str(row[orig_name_idx]).strip()
            rec['orig_org_name'] = extract_company_name(orig_name) or orig_name
            code = extract_company_code(orig_name)
            if code:
                rec['orig_org_code'] = code
            elif orig_code_idx is not None and row[orig_code_idx]:
                rec['orig_org_code'] = str(row[orig_code_idx]).strip()
            else:
                rec['orig_org_code'] = _unique_code_by_name(
                    rec['orig_org_name'])

        records.append(rec)

    if not records:
        raise ValueError("Excel 中没有有效的行数据")
    return records


def generate_use_list_sql(records, ops_remark=''):
    """
    生成适用清单修改SQL - 支持SQL合并策略
    """
    ops_remark = parse_ops_remark(ops_remark)
    cfg = get_config()
    sql = []
    sql.append("1、执行语句")

    if cfg.get('MERGE_MODULES', {}).get('use_list', True):
        # 启用合并策略:按相同ORG_CODE和org_name分组
        def key_fn(r):
            new_code = str(r.get('new_org_code') or '').strip()
            new_name = str(r.get('new_org_name') or '').strip()
            return (new_code, new_name)

        groups = merge_by_key(records, key_fn)

        for k, recs in groups.items():
            new_code, new_name = k
            if not new_code and not new_name:
                continue

            set_parts = []
            if new_code:
                set_parts.append(f"ORG_CODE='{new_code}'")
            if new_name:
                set_parts.append(f"org_name='{new_name}'")

            if set_parts:
                set_clause = ', '.join(set_parts)
                business_ids = [r['business_id']
                                for r in recs if r.get('business_id')]

                for chunk in chunk_list(business_ids, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
                    inlist = format_in(chunk)
                    sql.append(
                        f"UPDATE tphct04 SET {set_clause}, OPS_REMARK='{ops_remark}' WHERE BUSINESS_ID IN ({inlist}) AND alive_flag='1';")
    else:
        # 未启用合并策略:逐条生成SQL
        for rec in records:
            set_parts = []
            if rec.get('new_org_code'):
                set_parts.append(f"ORG_CODE='{rec['new_org_code']}'")
            if rec.get('new_org_name'):
                set_parts.append(f"org_name='{rec['new_org_name']}'")

            if set_parts:
                set_clause = ', '.join(set_parts)
                sql.append(
                    f"UPDATE tphct04 SET {set_clause}, OPS_REMARK='{ops_remark}' WHERE BUSINESS_ID='{rec['business_id']}' AND alive_flag='1';")

    sql.append("")
    sql.append("2、回退语句")

    if cfg.get('MERGE_MODULES', {}).get('use_list', True):
        def rb_key(r):
            orig_code = str(r.get('orig_org_code') or '').strip()
            orig_name = str(r.get('orig_org_name') or '').strip()
            return (orig_code, orig_name)

        groups = merge_by_key(records, rb_key)

        for k, recs in groups.items():
            orig_code, orig_name = k
            if not orig_code and not orig_name:
                continue

            rb_parts = []
            if orig_code:
                rb_parts.append(f"ORG_CODE='{orig_code}'")
            if orig_name:
                rb_parts.append(f"org_name='{orig_name}'")

            if rb_parts:
                rb_clause = ', '.join(rb_parts)
                business_ids = [r['business_id']
                                for r in recs if r.get('business_id')]

                for chunk in chunk_list(business_ids, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
                    inlist = format_in(chunk)
                    sql.append(
                        f"UPDATE tphct04 SET {rb_clause}, OPS_REMARK='' WHERE BUSINESS_ID IN ({inlist}) AND alive_flag='1';")
    else:
        for rec in records:
            rb_parts = []
            if rec.get('orig_org_code'):
                rb_parts.append(f"ORG_CODE='{rec['orig_org_code']}'")
            if rec.get('orig_org_name'):
                rb_parts.append(f"org_name='{rec['orig_org_name']}'")

            if rb_parts:
                rb_clause = ', '.join(rb_parts)
                sql.append(
                    f"UPDATE tphct04 SET {rb_clause}, OPS_REMARK='' WHERE BUSINESS_ID='{rec['business_id']}' AND alive_flag='1';")

    sql.append("")
    sql.append("3.数据库")
    sql.append("ip：192.168.11.71")
    sql.append("库名：cnnc_ph")

    return "\n".join(sql)


def download_use_list_template(request):
    """下载适用清单修改Excel模板"""
    if not OPENPYXL_AVAILABLE:
        return render(request, 'error.html', {'error': '服务器缺少 openpyxl 库'})

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '适用清单修改'
    ws.append(['BUSINESS_ID', '新组织机构名称', '新组织机构编码', '原组织机构名称', '原组织机构编码'])
    ws.append(['CON20241014000182', '中核集团', 'ZHXN', '原单位名称', 'ORIG_CODE'])

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)

    response = FileResponse(bio, as_attachment=True,
                            filename='适用清单修改模板.xlsx')
    response['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    return response


def use_list_update_view(request):
    """适用清单修改视图"""
    saved_file = None

    if request.method == 'POST':
        form = UseListUpdateForm(request.POST, request.FILES)
        if form.is_valid():
            cd = form.cleaned_data
            ops_remark = cd.get('ops_remark', '')

            if cd.get('excel_file'):
                # Excel批量导入
                records = parse_use_list_excel(cd['excel_file'])
            else:
                # 单条记录
                new_name = cd.get('new_org_name', '')
                new_code = extract_company_code(new_name)
                if not new_code:
                    new_code = _unique_code_by_name(
                        extract_company_name(new_name) or new_name)

                orig_name = cd.get('orig_org_name', '')
                orig_code = extract_company_code(orig_name)
                if not orig_code:
                    orig_code = _unique_code_by_name(
                        extract_company_name(orig_name) or orig_name)

                records = [{
                    'business_id': cd['business_id'],
                    'new_org_name': extract_company_name(new_name) or new_name,
                    'new_org_code': new_code,
                    'orig_org_name': extract_company_name(orig_name) or orig_name,
                    'orig_org_code': orig_code,
                }]

            sql_content = generate_use_list_sql(records, ops_remark)
            # 保存SQL到固定目录
            saved_file = save_sql_file(
                sql_content, '适用清单修改', cd.get('dynamic_id'))
    else:
        form = UseListUpdateForm()

    return render(request, 'use_list_update_form.html', {
        'form': form,
        'saved_file': saved_file,
        'active_menu': 'use_list_update',
        'sidebar_groups': SIDEBAR_GROUPS,
    })


__all__ = ['use_list_update_view', 'download_use_list_template']
