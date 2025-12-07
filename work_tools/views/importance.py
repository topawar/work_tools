"""物项重要性修改模块"""
import os
import io
import openpyxl
from datetime import datetime
from django.shortcuts import render
from django.http import FileResponse, HttpResponse
from django.conf import settings

from ..forms import ImportanceForm
from ..navigation import SIDEBAR_GROUPS
from ..config import get_config
from ..sql_merge import chunk_list, format_in, merge_by_key
from .base import save_sql_file, parse_ops_remark
from ..validation_utils import (
    validate_required,
    validate_enum,
    validate_at_least_one,
    generate_validation_failure_excel,
    get_temp_filename_from_path,
)


def validate_importance_records(records):
    """
    校验物项重要性修改记录
    
    Args:
        records: 解析后的记录列表
        
    Returns:
        校验结果字典
    """
    results = []
    passed_count = 0
    failed_count = 0
    
    # 有效的重要性值
    valid_importance_values = ['0', '1', '2', '3', '4', '一般准入备案类', '一般自行管理类', '一般', '核心', '重要']
    importance_display = '一般准入备案类/一般自行管理类/一般/核心/重要'
    
    for idx, record in enumerate(records):
        row_number = idx + 2  # 从第2行开始（第1行是表头）
        errors = []
        
        # 必须填写采购方案编号/询价单编号/合同号至少一个
        at_least_one_error = validate_at_least_one(
            [record.get('scheme_no'), record.get('inq_id'), record.get('bpo_id')],
            ['采购方案编号', '询价单编号', '合同号']
        )
        if at_least_one_error:
            errors.append(at_least_one_error)
        
        # 校验物项重要性（如果填写了）
        importance_code = record.get('importance_code')
        if importance_code is not None:
            # 先转换为字符串并去除空格
            if isinstance(importance_code, str):
                importance_code = importance_code.strip()
                record['importance_code'] = importance_code
            
            if importance_code:  # 如果不为空，校验枚举值
                enum_error = validate_enum(importance_code, '物项重要性', valid_importance_values, importance_display)
                if enum_error:
                    errors.append(enum_error)
        
        # 记录校验结果
        if errors:
            results.append({
                'row_number': row_number,
                'valid': False,
                'errors': errors
            })
            failed_count += 1
        else:
            results.append({
                'row_number': row_number,
                'valid': True,
                'errors': []
            })
            passed_count += 1
    
    return {
        'valid': failed_count == 0,
        'total': len(records),
        'passed': passed_count,
        'failed': failed_count,
        'results': results
    }


def parse_importance_excel(file):
    """解析物项重要性Excel文件"""
    wb = openpyxl.load_workbook(file)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    idx = {str(h): i for i, h in enumerate(headers)}

    def pick(alts):
        for a in alts:
            if a in idx:
                return idx[a]
        return None

    s_idx = pick(['scheme_no', 'PURCHASE_SCHEME_NO', '采购方案编号'])
    i_idx = pick(['inq_id', 'INQ_ID', '询价单编号'])
    b_idx = pick(['bpo_id', 'BPO_ID', '合同号'])
    imp_idx = pick(['importance', 'PROJECT_IMPORTANCE', '物项重要性'])
    oimp_idx = pick(['orig_importance', '原重要性'])

    def map_imp(v):
        """映射重要性值，如果无法映射则返回原值（用于校验）"""
        if v is None:
            return None
        val = str(v).strip()
        if not val:
            return None
        m = {
            '一般准入备案类': '0',
            '一般自行管理类': '1',
            '一般': '4',
            '核心': '3',
            '重要': '2',
        }
        if val in m:
            return m[val]
        if val in ['0', '1', '2', '3', '4']:
            return val
        # 返回原值以便校验函数检测错误
        return val

    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        rec = {}
        if s_idx is not None and row[s_idx] is not None:
            rec['scheme_no'] = str(row[s_idx]).strip()
        if i_idx is not None and row[i_idx] is not None:
            rec['inq_id'] = str(row[i_idx]).strip()
        if b_idx is not None and row[b_idx] is not None:
            rec['bpo_id'] = str(row[b_idx]).strip()
        if imp_idx is not None:
            rec['importance_code'] = map_imp(row[imp_idx])
        if oimp_idx is not None:
            rec['orig_importance_code'] = map_imp(row[oimp_idx])
        # 保留所有行，包括缺少业务编号的行，由校验函数处理
        records.append(rec)

    if not records:
        raise ValueError("Excel 中没有有效的行数据")
    return records


def generate_importance_sql_bulk(records, ops_remark=None, importance_code=None, orig_importance_code=None):
    """生成物项重要性修改SQL"""
    cfg = get_config()
    sql = []
    sql.append("1、执行语句")

    if cfg.get('MERGE_MODULES', {}).get('importance', True):
        def key_fn(r):
            return r.get('importance_code') if r.get('importance_code') is not None else importance_code

        groups = merge_by_key(records, key_fn)
        for code, recs in groups.items():
            if code is None:
                continue
            schemes = [r.get('scheme_no') for r in recs if r.get('scheme_no')]
            inqs = [r.get('inq_id') for r in recs if r.get('inq_id')]
            bpos = [r.get('bpo_id') for r in recs if r.get('bpo_id')]

            for chunk in chunk_list(schemes, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
                sql.append(
                    f"UPDATE tprfa01 SET IMPORTANCE='{code}', OPS_REMARK='{ops_remark or ''}' WHERE PURCHASE_SCHEME_NO IN ({format_in(chunk)});")
            for chunk in chunk_list(inqs, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
                sql.append(
                    f"UPDATE tprxj01 SET IMPORTANCE='{code}', OPS_REMARK='{ops_remark or ''}' WHERE INQ_ID IN ({format_in(chunk)});")
                sql.append(
                    f"UPDATE tprnq01 SET PROJECT_IMPORTANCE='{code}', OPS_REMARK='{ops_remark or ''}' WHERE INQ_ID IN ({format_in(chunk)});")
            for chunk in chunk_list(bpos, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
                sql.append(
                    f"UPDATE tphct01 SET PROJECT_IMPORTANCE='{code}', OPS_REMARK='{ops_remark or ''}' WHERE BPO_ID IN ({format_in(chunk)});")
    else:
        for r in records:
            code = r.get('importance_code') if r.get(
                'importance_code') is not None else importance_code
            if r.get('scheme_no') and code is not None:
                sql.append(
                    f"UPDATE tprfa01 SET IMPORTANCE='{code}', OPS_REMARK='{ops_remark or ''}' WHERE PURCHASE_SCHEME_NO='{r['scheme_no']}';")
            if r.get('inq_id') and code is not None:
                sql.append(
                    f"UPDATE tprxj01 SET IMPORTANCE='{code}', OPS_REMARK='{ops_remark or ''}' WHERE INQ_ID='{r['inq_id']}';")
                sql.append(
                    f"UPDATE tprnq01 SET PROJECT_IMPORTANCE='{code}', OPS_REMARK='{ops_remark or ''}' WHERE INQ_ID='{r['inq_id']}';")
            if r.get('bpo_id') and code is not None:
                sql.append(
                    f"UPDATE tphct01 SET PROJECT_IMPORTANCE='{code}', OPS_REMARK='{ops_remark or ''}' WHERE BPO_ID='{r['bpo_id']}';")

    sql.append("")
    sql.append("2、回退语句")

    if cfg.get('MERGE_MODULES', {}).get('importance', True):
        def rb_key(r):
            if r.get('orig_importance_code') is not None:
                return r.get('orig_importance_code')
            return orig_importance_code if orig_importance_code is not None else '2'

        groups = merge_by_key(records, rb_key)
        for rcode, recs in groups.items():
            schemes = [r.get('scheme_no') for r in recs if r.get('scheme_no')]
            inqs = [r.get('inq_id') for r in recs if r.get('inq_id')]
            bpos = [r.get('bpo_id') for r in recs if r.get('bpo_id')]

            for chunk in chunk_list(schemes, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
                sql.append(
                    f"UPDATE tprfa01 SET IMPORTANCE='{rcode}', OPS_REMARK='' WHERE PURCHASE_SCHEME_NO IN ({format_in(chunk)});")
            for chunk in chunk_list(inqs, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
                sql.append(
                    f"UPDATE tprxj01 SET IMPORTANCE='{rcode}', OPS_REMARK='' WHERE INQ_ID IN ({format_in(chunk)});")
                sql.append(
                    f"UPDATE tprnq01 SET PROJECT_IMPORTANCE='{rcode}', OPS_REMARK='' WHERE INQ_ID IN ({format_in(chunk)});")
            for chunk in chunk_list(bpos, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
                sql.append(
                    f"UPDATE tphct01 SET PROJECT_IMPORTANCE='{rcode}', OPS_REMARK='' WHERE BPO_ID IN ({format_in(chunk)});")
    else:
        for r in records:
            rcode = r.get('orig_importance_code') if r.get('orig_importance_code') is not None else (
                orig_importance_code if orig_importance_code is not None else '2')
            if r.get('scheme_no'):
                sql.append(
                    f"UPDATE tprfa01 SET IMPORTANCE='{rcode}', OPS_REMARK='' WHERE PURCHASE_SCHEME_NO='{r['scheme_no']}';")
            if r.get('inq_id'):
                sql.append(
                    f"UPDATE tprxj01 SET IMPORTANCE='{rcode}', OPS_REMARK='' WHERE INQ_ID='{r['inq_id']}';")
                sql.append(
                    f"UPDATE tprnq01 SET PROJECT_IMPORTANCE='{rcode}', OPS_REMARK='' WHERE INQ_ID='{r['inq_id']}';")
            if r.get('bpo_id'):
                sql.append(
                    f"UPDATE tphct01 SET PROJECT_IMPORTANCE='{rcode}', OPS_REMARK='' WHERE BPO_ID='{r['bpo_id']}';")

    sql.append("")
    sql.append("3、数据库")
    sql.append("ip:192.168.11.69")
    sql.append("库名：cnnc_pr")
    sql.append("")
    sql.append("ip：192.168.11.71")
    sql.append("库名：cnnc_ph")
    return "\n".join(sql)


def importance_update_view(request):
    """物项重要性修改视图"""
    saved_file = None
    validation_failure = None
    
    if request.method == 'POST':
        form = ImportanceForm(request.POST, request.FILES)
        if form.is_valid():
            cd = form.cleaned_data
            ops_remark = parse_ops_remark(cd.get('ops_remark', ''))

            if cd.get('excel_file'):
                records = parse_importance_excel(cd['excel_file'])
                
                # 执行数据校验
                validation_result = validate_importance_records(records)
                
                if not validation_result['valid']:
                    # 校验失败，生成包含错误信息的Excel文件
                    temp_file = generate_validation_failure_excel(
                        cd['excel_file'], 
                        validation_result, 
                        'importance'
                    )
                    
                    if temp_file:
                        validation_failure = {
                            'total': validation_result['total'],
                            'passed': validation_result['passed'],
                            'failed': validation_result['failed'],
                            'filename': get_temp_filename_from_path(temp_file)
                        }
                    
                    # 不生成SQL，直接返回页面显示错误
                    return render(request, 'importance_form.html', {
                        'form': form,
                        'validation_failure': validation_failure,
                        'active_menu': 'importance',
                        'sidebar_groups': SIDEBAR_GROUPS,
                    })
                
                sql_content = generate_importance_sql_bulk(records, ops_remark)
            else:
                rec = {
                    'scheme_no': cd.get('scheme_no'),
                    'inq_id': cd.get('inq_id'),
                    'bpo_id': cd.get('bpo_id'),
                }
                records = [rec]
                sql_content = generate_importance_sql_bulk(
                    records, ops_remark, cd.get('importance_choice'), cd.get(
                        'orig_importance_choice')
                )

            # 保存SQL到固定目录
            saved_file = save_sql_file(
                sql_content, '物项重要性修改', cd.get('dynamic_id'))

            request.session['importance_last'] = {
                k: v for k, v in cd.items() if k != 'excel_file'
            }
    else:
        if request.GET.get('clear'):
            request.session.pop('importance_last', None)
            form = ImportanceForm()
        else:
            initial = request.session.get('importance_last')
            form = ImportanceForm(initial=initial)

    return render(request, 'importance_form.html', {
        'form': form,
        'saved_file': saved_file,
        'validation_failure': validation_failure,
        'active_menu': 'importance',
        'sidebar_groups': SIDEBAR_GROUPS,
    })


def download_importance_template(request):
    """下载物项重要性模板"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '模板'
    ws.append(['采购方案编号', '询价单编号', '合同号', '物项重要性', '原重要性'])
    ws.append(['HNGS-CGFA-25-01658', 'CNSC-XJD-25-02704',
              'HNGS-25-00255', '一般', '重要'])
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    response = FileResponse(bio, as_attachment=True,
                            filename='importance_template.xlsx')
    response['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    return response


# 别名兼容
importance_view = importance_update_view

__all__ = [
    'importance_view',
    'importance_update_view',
    'download_importance_template',
]
