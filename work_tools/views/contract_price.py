"""
合同明细单价修改模块
包含单价修改的所有视图函数和SQL生成逻辑
"""
import io
import logging
from django.shortcuts import render
from django.http import FileResponse
from ..forms import ContractDetailPriceForm
from ..navigation import get_sidebar_groups
from ..config import get_config
from ..sql_merge import chunk_list, format_in, merge_by_key
from ..logger_utils import log_view_input
from .base import parse_ops_remark, save_sql_file
from ..validation_utils import (
    validate_required,
    generate_validation_failure_excel,
    get_temp_filename_from_path,
)

try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

logger = logging.getLogger('work_tools.view')
sql_logger = logging.getLogger('work_tools.sql')


def parse_price_excel(file):
    """解析Excel文件,提取单价修改数据"""
    if not OPENPYXL_AVAILABLE:
        raise ValueError("缺少 openpyxl，无法解析 Excel")

    wb = openpyxl.load_workbook(file)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    idx = {str(h): i for i, h in enumerate(headers)}

    def pick(alts):
        for a in alts:
            if a in idx:
                return idx[a]
        return None

    lid_idx = pick(['line_id', 'BPO_LINE_ID', '明细行ID'])
    price_idx = pick(['price', 'BPO_PRICE', '单价'])
    qty_idx = pick(['quantity', 'BPO_QTY', '数量'])
    tax_rate_idx = pick(['tax_rate', 'TAX_RATE', '税率'])
    orig_qty_idx = pick(['orig_quantity', '原数量'])
    orig_price_idx = pick(['orig_price', '原单价'])
    orig_tax_rate_idx = pick(['orig_tax_rate', '原税率'])

    if lid_idx is None or price_idx is None:
        raise ValueError("Excel 缺少必要列：明细行ID/单价")

    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        lid = row[lid_idx]
        prc = row[price_idx]
        qty = row[qty_idx] if qty_idx is not None else None
        tax_rate = row[tax_rate_idx] if tax_rate_idx is not None else None
        if lid is None or prc is None:
            continue
        rec = {
            'line_id': str(lid).strip(),
            'price': prc,
            'quantity': qty,
            'tax_rate': tax_rate,
        }
        if orig_qty_idx is not None:
            rec['orig_quantity'] = row[orig_qty_idx]
        if orig_price_idx is not None:
            rec['orig_price'] = row[orig_price_idx]
        if orig_tax_rate_idx is not None:
            rec['orig_tax_rate'] = row[orig_tax_rate_idx]
        records.append(rec)

    if not records:
        raise ValueError("Excel 中没有有效的行数据")
    return records


def generate_price_sql_bulk(records, ops_remark=None, global_tax_rate=None):
    """
    生成批量单价修改SQL - 支持SQL合并策略
    相同单价和数量的记录会自动合并为IN查询
    
    Args:
        records: 记录列表
        ops_remark: 操作备注
        global_tax_rate: 全局税率（单条模式使用），如果记录中有tax_rate则优先使用记录中的
    """
    cfg = get_config()
    sql = []
    sql.append("1、执行语句")

    # 日志: 记录SQL生成开始
    sql_logger.info(
        f"[SQL合并] 开始生成单价SQL, 记录数={len(records)}, 配置启用={cfg.get('MERGE_MODULES', {}).get('price', True)}, 全局税率={global_tax_rate}")

    if cfg.get('MERGE_MODULES', {}).get('price', True):
        # 启用SQL合并策略: 按单价、数量和税率分组
        def key_fn(r):
            # 优先使用记录中的税率，否则使用全局税率
            tax = r.get('tax_rate') if r.get('tax_rate') is not None else global_tax_rate
            return (r.get('price'), r.get('quantity'), tax)

        groups = merge_by_key(records, key_fn)
        sql_logger.info(f"[SQL合并] 分组结果: {len(groups)}个分组")

        for k, recs in groups.items():
            p, q, tax = k
            ids = [r['line_id'] for r in recs if r.get('line_id')]
            sql_logger.info(
                f"[SQL合并] 处理分组: price={p}, qty={q}, tax_rate={tax}, 记录数={len(ids)}")

            for chunk in chunk_list(ids, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
                inlist = format_in(chunk)
                # 税率表达式：如果有税率则使用固定值，否则使用表中的TAX_RATE
                tax_expr = str(tax) if tax is not None else 'TAX_RATE'
                
                # 情况1: 数量和单价都有值
                if q is not None and p is not None:
                    stmt = (
                        f"UPDATE tphct02 SET "
                        f"BPO_QTY={q}, "
                        f"BPO_PRICE={p}, "
                        f"BPO_AMT={p}*{q}, "
                        f"BPO_NOTAX_PRICE={p}/(1+{tax_expr}), "
                        f"BPO_NOTAX_AMT=({p}*{q})/(1+{tax_expr}), "
                        f"OPS_REMARK='{ops_remark or ''}' "
                        f"WHERE BPO_LINE_ID IN ({inlist}) AND ALIVE_FLAG='1';"
                    )
                    sql_logger.info(f"[SQL合并] 同时修改数量和单价: {len(chunk)}条记录")
                # 情况2: 仅修改数量，单价使用原表字段
                elif q is not None and p is None:
                    stmt = (
                        f"UPDATE tphct02 SET "
                        f"BPO_QTY={q}, "
                        f"BPO_AMT=BPO_PRICE*{q}, "
                        f"BPO_NOTAX_AMT=(BPO_PRICE*{q})/(1+{tax_expr}), "
                        f"OPS_REMARK='{ops_remark or ''}' "
                        f"WHERE BPO_LINE_ID IN ({inlist}) AND ALIVE_FLAG='1';"
                    )
                    sql_logger.info(f"[SQL合并] 仅修改数量，单价保持原值: {len(chunk)}条记录")
                # 情况3: 仅修改单价，数量使用原表字段
                elif q is None and p is not None:
                    stmt = (
                        f"UPDATE tphct02 SET "
                        f"BPO_PRICE={p}, "
                        f"BPO_AMT={p}*BPO_QTY, "
                        f"BPO_NOTAX_PRICE={p}/(1+{tax_expr}), "
                        f"BPO_NOTAX_AMT=({p}*BPO_QTY)/(1+{tax_expr}), "
                        f"OPS_REMARK='{ops_remark or ''}' "
                        f"WHERE BPO_LINE_ID IN ({inlist}) AND ALIVE_FLAG='1';"
                    )
                    sql_logger.info(f"[SQL合并] 仅修改单价，数量保持原值: {len(chunk)}条记录")
                # 情况4: 仅修改税率，单价和数量使用原表字段
                elif q is None and p is None and tax is not None:
                    stmt = (
                        f"UPDATE tphct02 SET "
                        f"TAX_RATE={tax}, "
                        f"BPO_NOTAX_PRICE=BPO_PRICE/(1+{tax}), "
                        f"BPO_NOTAX_AMT=BPO_AMT/(1+{tax}), "
                        f"OPS_REMARK='{ops_remark or ''}' "
                        f"WHERE BPO_LINE_ID IN ({inlist}) AND ALIVE_FLAG='1';"
                    )
                    sql_logger.info(f"[SQL合并] 仅修改税率: {len(chunk)}条记录")
                else:
                    # 数量、单价和税率都为空，跳过
                    sql_logger.warning(f"[SQL合并] 跳过无效记录: 数量、单价和税率都为空")
                    continue
                sql.append(stmt)
    else:
        # 未启用合并策略: 逐条生成SQL
        sql_logger.info("[SQL合并] 合并策略已禁用, 逐条生成SQL")
        for r in records:
            q = r.get('quantity')
            p = r.get('price')
            # 优先使用记录中的税率，否则使用全局税率
            tax = r.get('tax_rate') if r.get('tax_rate') is not None else global_tax_rate
            tax_expr = str(tax) if tax is not None else 'TAX_RATE'
            
            # 情况1: 数量和单价都有值
            if q is not None and p is not None:
                exec_stmt = (
                    f"UPDATE tphct02 SET "
                    f"BPO_QTY={q}, "
                    f"BPO_PRICE={p}, "
                    f"BPO_AMT={p}*{q}, "
                    f"BPO_NOTAX_PRICE={p}/(1+{tax_expr}), "
                    f"BPO_NOTAX_AMT=({p}*{q})/(1+{tax_expr}), "
                    f"OPS_REMARK='{ops_remark or ''}' "
                    f"WHERE BPO_LINE_ID='{r['line_id']}' AND ALIVE_FLAG='1';"
                )
                sql_logger.info(f"[SQL生成] 同时修改数量和单价: {r['line_id']}")
            # 情况2: 仅修改数量，单价使用原表字段
            elif q is not None and p is None:
                exec_stmt = (
                    f"UPDATE tphct02 SET "
                    f"BPO_QTY={q}, "
                    f"BPO_AMT=BPO_PRICE*{q}, "
                    f"BPO_NOTAX_AMT=(BPO_PRICE*{q})/(1+{tax_expr}), "
                    f"OPS_REMARK='{ops_remark or ''}' "
                    f"WHERE BPO_LINE_ID='{r['line_id']}' AND ALIVE_FLAG='1';"
                )
                sql_logger.info(f"[SQL生成] 仅修改数量，单价保持原值: {r['line_id']}")
            # 情况3: 仅修改单价，数量使用原表字段
            elif q is None and p is not None:
                exec_stmt = (
                    f"UPDATE tphct02 SET "
                    f"BPO_PRICE={p}, "
                    f"BPO_AMT={p}*BPO_QTY, "
                    f"BPO_NOTAX_PRICE={p}/(1+{tax_expr}), "
                    f"BPO_NOTAX_AMT=({p}*BPO_QTY)/(1+{tax_expr}), "
                    f"OPS_REMARK='{ops_remark or ''}' "
                    f"WHERE BPO_LINE_ID='{r['line_id']}' AND ALIVE_FLAG='1';"
                )
                sql_logger.info(f"[SQL生成] 仅修改单价，数量保持原值: {r['line_id']}")
            # 情况4: 仅修改税率，单价和数量使用原表字段
            elif q is None and p is None and tax is not None:
                exec_stmt = (
                    f"UPDATE tphct02 SET "
                    f"TAX_RATE={tax}, "
                    f"BPO_NOTAX_PRICE=BPO_PRICE/(1+{tax}), "
                    f"BPO_NOTAX_AMT=BPO_AMT/(1+{tax}), "
                    f"OPS_REMARK='{ops_remark or ''}' "
                    f"WHERE BPO_LINE_ID='{r['line_id']}' AND ALIVE_FLAG='1';"
                )
                sql_logger.info(f"[SQL生成] 仅修改税率: {r['line_id']}")
            else:
                # 数量、单价和税率都为空，跳过
                sql_logger.warning(f"[SQL生成] 跳过无效记录: {r['line_id']} - 数量、单价和税率都为空")
                continue
            sql.append(exec_stmt)

    # 更新合同主表汇总金额
    bpo_ids = []
    for r in records:
        parts = r['line_id'].split('-')
        if len(parts) > 4:
            bpo_id = '-'.join(parts[:-1])
        else:
            bpo_id = '-'.join(parts[:-1]) if len(parts) > 1 else r['line_id']
        if bpo_id not in bpo_ids:
            bpo_ids.append(bpo_id)

    sql.append("-- 批量更新合同的主表汇总金额")
    for bpo_id in bpo_ids:
        main_stmt = (
            f"UPDATE tphct01 SET "
            f"BPO_AMT=(SELECT SUM(BPO_AMT) FROM tphct02 WHERE BPO_ID='{bpo_id}' AND ALIVE_FLAG='1'), "
            f"TARGET_AMT=(SELECT SUM(BPO_AMT) FROM tphct02 WHERE BPO_ID='{bpo_id}' AND ALIVE_FLAG='1'), "
            f"BPO_NOTAX_AMT=(SELECT SUM(BPO_NOTAX_AMT) FROM tphct02 WHERE BPO_ID='{bpo_id}' AND ALIVE_FLAG='1'), "
            f"OPS_REMARK='{ops_remark or ''}' "
            f"WHERE BPO_ID='{bpo_id}' AND ALIVE_FLAG='1';"
        )
        sql.append(main_stmt)

    # 生成回退语句
    sql.append("2、回退语句")
    if cfg.get('MERGE_MODULES', {}).get('price', True):
        def rb_key(r):
            return (r.get('orig_price'), r.get('orig_quantity'), r.get('orig_tax_rate'))

        groups = merge_by_key(records, rb_key)
        for k, recs in groups.items():
            rp, rq, rt = k
            # 如果原单价、原数量、原税率都为空，跳过
            if rp is None and rq is None and rt is None:
                continue
            ids = [r['line_id'] for r in recs if r.get('line_id')]
            # 原税率表达式
            orig_tax_expr = str(rt) if rt is not None else 'TAX_RATE'

            for chunk in chunk_list(ids, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
                inlist = format_in(chunk)
                if rq is not None and rp is not None:
                    rollback_stmt = (
                        f"UPDATE tphct02 SET "
                        f"BPO_QTY={rq}, "
                        f"BPO_PRICE={rp}, "
                        f"BPO_AMT={rp}*{rq}, "
                        f"BPO_NOTAX_PRICE={rp}/(1+{orig_tax_expr}), "
                        f"BPO_NOTAX_AMT=({rp}*{rq})/(1+{orig_tax_expr}), "
                        f"OPS_REMARK='' "
                        f"WHERE BPO_LINE_ID IN ({inlist}) AND ALIVE_FLAG='1';"
                    )
                elif rp is not None and rq is None:
                    rollback_stmt = (
                        f"UPDATE tphct02 SET "
                        f"BPO_PRICE={rp}, "
                        f"BPO_AMT={rp}*BPO_QTY, "
                        f"BPO_NOTAX_PRICE={rp}/(1+{orig_tax_expr}), "
                        f"BPO_NOTAX_AMT=({rp}*BPO_QTY)/(1+{orig_tax_expr}), "
                        f"OPS_REMARK='' "
                        f"WHERE BPO_LINE_ID IN ({inlist}) AND ALIVE_FLAG='1';"
                    )
                elif rp is None and rq is not None:
                    rollback_stmt = (
                        f"UPDATE tphct02 SET "
                        f"BPO_QTY={rq}, "
                        f"BPO_AMT=BPO_PRICE*{rq}, "
                        f"BPO_NOTAX_AMT=(BPO_PRICE*{rq})/(1+{orig_tax_expr}), "
                        f"OPS_REMARK='' "
                        f"WHERE BPO_LINE_ID IN ({inlist}) AND ALIVE_FLAG='1';"
                    )
                elif rp is None and rq is None and rt is not None:
                    # 仅回退税率
                    rollback_stmt = (
                        f"UPDATE tphct02 SET "
                        f"TAX_RATE={rt}, "
                        f"BPO_NOTAX_PRICE=BPO_PRICE/(1+{rt}), "
                        f"BPO_NOTAX_AMT=BPO_AMT/(1+{rt}), "
                        f"OPS_REMARK='' "
                        f"WHERE BPO_LINE_ID IN ({inlist}) AND ALIVE_FLAG='1';"
                    )
                else:
                    rollback_stmt = None
                if rollback_stmt:
                    sql.append(rollback_stmt)
    else:
        for r in records:
            rq = r.get('orig_quantity')
            rp = r.get('orig_price')
            rt = r.get('orig_tax_rate')
            orig_tax_expr = str(rt) if rt is not None else 'TAX_RATE'
            
            if rq is not None and rp is not None:
                rollback_stmt = (
                    f"UPDATE tphct02 SET "
                    f"BPO_QTY={rq}, "
                    f"BPO_PRICE={rp}, "
                    f"BPO_AMT={rp}*{rq}, "
                    f"BPO_NOTAX_PRICE={rp}/(1+{orig_tax_expr}), "
                    f"BPO_NOTAX_AMT=({rp}*{rq})/(1+{orig_tax_expr}), "
                    f"OPS_REMARK='' "
                    f"WHERE BPO_LINE_ID='{r['line_id']}' AND ALIVE_FLAG='1';"
                )
                sql.append(rollback_stmt)
            elif rp is not None and rq is None:
                rollback_stmt = (
                    f"UPDATE tphct02 SET "
                    f"BPO_PRICE={rp}, "
                    f"BPO_AMT={rp}*BPO_QTY, "
                    f"BPO_NOTAX_PRICE={rp}/(1+{orig_tax_expr}), "
                    f"BPO_NOTAX_AMT=({rp}*BPO_QTY)/(1+{orig_tax_expr}), "
                    f"OPS_REMARK='' "
                    f"WHERE BPO_LINE_ID='{r['line_id']}' AND ALIVE_FLAG='1';"
                )
                sql.append(rollback_stmt)
            elif rp is None and rq is not None:
                rollback_stmt = (
                    f"UPDATE tphct02 SET "
                    f"BPO_QTY={rq}, "
                    f"BPO_AMT=BPO_PRICE*{rq}, "
                    f"BPO_NOTAX_AMT=(BPO_PRICE*{rq})/(1+{orig_tax_expr}), "
                    f"OPS_REMARK='' "
                    f"WHERE BPO_LINE_ID='{r['line_id']}' AND ALIVE_FLAG='1';"
                )
                sql.append(rollback_stmt)
            elif rp is None and rq is None and rt is not None:
                # 仅回退税率
                rollback_stmt = (
                    f"UPDATE tphct02 SET "
                    f"TAX_RATE={rt}, "
                    f"BPO_NOTAX_PRICE=BPO_PRICE/(1+{rt}), "
                    f"BPO_NOTAX_AMT=BPO_AMT/(1+{rt}), "
                    f"OPS_REMARK='' "
                    f"WHERE BPO_LINE_ID='{r['line_id']}' AND ALIVE_FLAG='1';"
                )
                sql.append(rollback_stmt)

    sql.append("-- 回退合同主表")
    for bpo_id in bpo_ids:
        main_rollback_stmt = (
            f"UPDATE tphct01 SET "
            f"BPO_AMT=(SELECT SUM(BPO_AMT) FROM tphct02 WHERE BPO_ID='{bpo_id}' AND ALIVE_FLAG='1'), "
            f"TARGET_AMT=(SELECT SUM(BPO_AMT) FROM tphct02 WHERE BPO_ID='{bpo_id}' AND ALIVE_FLAG='1'), "
            f"BPO_NOTAX_AMT=(SELECT SUM(BPO_NOTAX_AMT) FROM tphct02 WHERE BPO_ID='{bpo_id}' AND ALIVE_FLAG='1'), "
            f"OPS_REMARK='' "
            f"WHERE BPO_ID='{bpo_id}' AND ALIVE_FLAG='1';"
        )
        sql.append(main_rollback_stmt)

    sql.append("3.数据库")
    sql.append("ip：192.168.11.71")
    sql.append("库名：cnnc_ph")

    sql_logger.info(f"[SQL合并] SQL生成完成, 总行数={len(sql)}")
    return "\n".join(sql)


def download_contract_price_template(request):
    """下载合同单价修改Excel模板"""
    if not OPENPYXL_AVAILABLE:
        return render(request, 'error.html', {'error': '服务器缺少 openpyxl 库'})

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '合同明细单价修改'
    ws.append(['明细行ID', '单价', '数量', '税率', '原单价', '原数量', '原税率'])
    ws.append(['BPO-EXAMPLE-000001', 100, 1, 0.13, None, None, None])

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)

    response = FileResponse(bio, as_attachment=True,
                            filename='contract_price_template.xlsx')
    response['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    return response


def validate_price_records(records):
    """
    校验单价修改记录
    
    Args:
        records: 记录列表
        
    Returns:
        校验结果字典
    """
    results = []
    passed_count = 0
    failed_count = 0
    
    for idx, record in enumerate(records):
        row_number = idx + 2  # 跳过表头
        errors = []
        
        # 必填项校验：明细行ID必填
        error = validate_required(record.get('line_id'), '明细行ID')
        if error:
            errors.append(error)
        
        # 记录结果
        if errors:
            results.append({
                'row_number': row_number,
                'valid': False,
                'errors': errors
            })
            failed_count += 1
        else:
            results.append({
                'row_number': row_number,
                'valid': True,
                'errors': []
            })
            passed_count += 1
    
    return {
        'valid': failed_count == 0,
        'total': len(records),
        'passed': passed_count,
        'failed': failed_count,
        'results': results
    }


def contract_detail_price_view(request):
    """合同明细单价修改视图 - 支持单条和批量导入"""
    saved_file = None
    validation_failure = None

    if request.method == 'POST':
        form = ContractDetailPriceForm(request.POST, request.FILES)
        if form.is_valid():
            try:
                cd = form.cleaned_data
                log_view_input("合同明细单价修改", cd)

                ops_remark = parse_ops_remark(cd.get('ops_remark', ''))

                if cd.get('excel_file'):
                    # Excel批量导入模式
                    records = parse_price_excel(cd['excel_file'])
                    
                    # 执行数据校验
                    validation_result = validate_price_records(records)
                    
                    if not validation_result['valid']:
                        # 校验失败，生成失败文件
                        temp_file = generate_validation_failure_excel(
                            cd['excel_file'],
                            validation_result,
                            'contract_price'
                        )
                        
                        if temp_file:
                            validation_failure = {
                                'total': validation_result['total'],
                                'passed': validation_result['passed'],
                                'failed': validation_result['failed'],
                                'filename': get_temp_filename_from_path(temp_file)
                            }
                            logger.info(f"校验失败: 总行数={validation_failure['total']}, "
                                      f"失败行数={validation_failure['failed']}")
                        
                        return render(request, 'contract_price_form.html', {
                            'form': form,
                            'validation_failure': validation_failure,
                            'active_menu': 'contract_price',
                            'sidebar_groups': get_sidebar_groups(),
                        })
                    
                    # 校验通过，生成SQL
                    # Excel中每行可以有自己的税率，但如果没有则使用表单中的全局税率
                    global_tax_rate = cd.get('tax_rate')
                    sql_content = generate_price_sql_bulk(records, ops_remark, global_tax_rate)
                else:
                    # 单条记录模式 - 也使用SQL合并策略
                    new_quantity = cd.get('new_quantity')
                    new_price = cd.get('new_price')
                    tax_rate = cd.get('tax_rate')
                    orig_quantity = cd.get('orig_quantity')
                    orig_price = cd.get('orig_price')

                    single_record = {
                        'line_id': cd['single_line_id'],
                        'price': new_price,
                        'quantity': new_quantity,
                        'tax_rate': tax_rate,
                        'orig_price': orig_price,
                        'orig_quantity': orig_quantity,
                        'orig_tax_rate': cd.get('orig_tax_rate')
                    }
                    sql_content = generate_price_sql_bulk(
                        [single_record], ops_remark, tax_rate)

                # 保存SQL到固定目录
                saved_file = save_sql_file(
                    sql_content, '修改合同明细单价', cd.get('dynamic_id'))
            except Exception as e:
                logger.error(f"[合同明细单价修改] 处理失败: {e}", exc_info=True)
                form.add_error(None, f"生成SQL失败: {str(e)}")
    else:
        form = ContractDetailPriceForm()

    return render(request, 'contract_price_form.html', {
        'form': form,
        'saved_file': saved_file,
        'validation_failure': validation_failure,
        'active_menu': 'contract_price',
        'sidebar_groups': get_sidebar_groups(),
    })


# 为了兼容旧代码,保留contract_price_view作为别名
contract_price_view = contract_detail_price_view
download_price_template = download_contract_price_template


__all__ = [
    'parse_price_excel',
    'generate_price_sql_bulk',
    'download_contract_price_template',
    'download_price_template',
    'contract_detail_price_view',
    'contract_price_view',
]
