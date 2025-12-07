"""
数据校验工具模块
提供批量导入数据的校验功能，包括必填项、数值、日期、枚举等校验
"""
import os
import re
import logging
from datetime import datetime
from decimal import Decimal, InvalidOperation
from django.conf import settings

try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

logger = logging.getLogger('work_tools.validation')


def validate_required(value, field_name):
    """
    必填项校验
    
    Args:
        value: 待校验的值
        field_name: 字段名称
        
    Returns:
        错误信息字符串，如果校验通过返回None
    """
    if value is None or (isinstance(value, str) and not value.strip()):
        return f"{field_name}不能为空"
    return None


def validate_number(value, field_name):
    """
    数值校验
    
    Args:
        value: 待校验的值
        field_name: 字段名称
        
    Returns:
        错误信息字符串，如果校验通过返回None
    """
    if value is None:
        return None
    
    try:
        if isinstance(value, str):
            value = value.strip()
            if not value:
                return None
        Decimal(str(value))
        return None
    except (InvalidOperation, ValueError):
        return f"{field_name}必须为数值"


def validate_date_format(value, field_name, date_format='%Y%m%d'):
    """
    日期格式校验
    
    Args:
        value: 待校验的值
        field_name: 字段名称
        date_format: 日期格式，默认为YYYYMMDD
        
    Returns:
        错误信息字符串，如果校验通过返回None
    """
    if value is None:
        return None
    
    try:
        if isinstance(value, str):
            value = value.strip()
            if not value:
                return None
        datetime.strptime(str(value), date_format)
        return None
    except ValueError:
        if date_format == '%Y%m%d':
            return f"{field_name}格式错误，应为YYYYMMDD格式"
        else:
            return f"{field_name}格式错误"


def validate_enum(value, field_name, valid_values, value_display=None):
    """
    枚举值校验
    
    Args:
        value: 待校验的值
        field_name: 字段名称
        valid_values: 有效值列表
        value_display: 用于错误提示的值显示字符串，默认为None
        
    Returns:
        错误信息字符串，如果校验通过返回None
    """
    if value is None:
        return None
    
    if isinstance(value, str):
        value = value.strip()
        if not value:
            return None
    
    if value not in valid_values:
        if value_display:
            return f"{field_name}取值无效，应为：{value_display}"
        else:
            return f"{field_name}取值无效"
    return None


def validate_reference(value, field_name, model_class, query_field='id'):
    """
    引用完整性校验（检查值是否在数据库中存在）
    
    Args:
        value: 待校验的值
        field_name: 字段名称
        model_class: Django模型类
        query_field: 查询字段名称，默认为'id'
        
    Returns:
        错误信息字符串，如果校验通过返回None
    """
    if value is None:
        return None
    
    if isinstance(value, str):
        value = value.strip()
        if not value:
            return None
    
    try:
        exists = model_class.objects.filter(**{query_field: value}).exists()
        if not exists:
            return f"{field_name}不存在"
        return None
    except Exception as e:
        logger.error(f"引用完整性校验失败: {field_name}={value}, error={str(e)}")
        return f"{field_name}校验失败"


def validate_at_least_one(values, field_names):
    """
    至少填一个校验（多个字段中至少有一个不为空）
    
    Args:
        values: 字段值列表
        field_names: 字段名称列表（用于错误提示）
        
    Returns:
        错误信息字符串，如果校验通过返回None
    """
    has_value = False
    for v in values:
        if v is not None:
            if isinstance(v, str):
                if v.strip():
                    has_value = True
                    break
            else:
                has_value = True
                break
    
    if not has_value:
        field_list = '/'.join(field_names)
        return f"必须填写{field_list}至少一个"
    return None


def generate_validation_failure_excel(uploaded_file, validation_result, module_name):
    """
    生成校验失败Excel文件
    
    Args:
        uploaded_file: 用户上传的Excel文件对象（Django UploadedFile）
        validation_result: 校验结果字典
        module_name: 模块名称（用于文件命名）
        
    Returns:
        临时文件路径，如果生成失败返回None
    """
    if not OPENPYXL_AVAILABLE:
        logger.error("openpyxl库不可用，无法生成校验失败文件")
        return None
    
    try:
        # 重新读取上传的文件
        uploaded_file.seek(0)
        wb = openpyxl.load_workbook(uploaded_file, data_only=True)
        ws = wb.active
        
        # 在最后一列添加“校验结果”列
        max_col = ws.max_column
        ws.cell(row=1, column=max_col + 1, value='校验结果')
        
        # 填充校验结果
        for result in validation_result.get('results', []):
            row_num = result.get('row_number')
            if row_num is None:
                continue
                
            if result.get('valid'):
                # 校验通过，留空或显示“通过”
                ws.cell(row=row_num, column=max_col + 1, value='')
            else:
                # 校验失败，显示错误信息
                errors = result.get('errors', [])
                error_msg = '; '.join(errors)
                ws.cell(row=row_num, column=max_col + 1, value=error_msg)
        
        # 保存到临时文件
        temp_path = save_validation_temp_file(wb, module_name)
        return temp_path
        
    except Exception as e:
        logger.error(f"生成校验失败文件异常: module={module_name}, error={str(e)}")
        import traceback
        logger.error(traceback.format_exc())
        return None


def save_validation_temp_file(workbook, module_name):
    """
    保存校验失败文件到临时目录
    
    Args:
        workbook: openpyxl工作簿对象
        module_name: 模块名称
        
    Returns:
        文件完整路径
    """
    # 创建临时目录
    temp_dir = os.path.join(settings.BASE_DIR, 'temp_uploads', 'validation_failures')
    os.makedirs(temp_dir, exist_ok=True)
    
    # 生成文件名
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    filename = f"validation_failed_{module_name}_{timestamp}.xlsx"
    filepath = os.path.join(temp_dir, filename)
    
    # 保存文件
    workbook.save(filepath)
    logger.info(f"校验失败文件已保存: {filepath}")
    
    return filepath


def get_temp_filename_from_path(filepath):
    """
    从完整路径中提取文件名
    
    Args:
        filepath: 文件完整路径
        
    Returns:
        文件名
    """
    return os.path.basename(filepath)


def cleanup_temp_files(hours=24):
    """
    清理过期的临时文件
    
    Args:
        hours: 文件过期时间（小时），默认24小时
        
    Returns:
        清理的文件数量
    """
    temp_dir = os.path.join(settings.BASE_DIR, 'temp_uploads', 'validation_failures')
    if not os.path.exists(temp_dir):
        return 0
    
    count = 0
    now = datetime.now()
    
    try:
        for filename in os.listdir(temp_dir):
            filepath = os.path.join(temp_dir, filename)
            
            # 检查文件是否过期
            file_mtime = datetime.fromtimestamp(os.path.getmtime(filepath))
            age_hours = (now - file_mtime).total_seconds() / 3600
            
            if age_hours > hours:
                os.remove(filepath)
                count += 1
                logger.info(f"已清理过期文件: {filename}")
                
    except Exception as e:
        logger.error(f"清理临时文件异常: {str(e)}")
    
    return count


def validate_file_size(uploaded_file, max_size_mb=10):
    """
    校验文件大小
    
    Args:
        uploaded_file: 上传的文件对象
        max_size_mb: 最大文件大小（MB），默认10MB
        
    Returns:
        错误信息字符串，如果校验通过返回None
    """
    if uploaded_file.size > max_size_mb * 1024 * 1024:
        return f"文件过大，请确保文件不超过{max_size_mb}MB"
    return None


__all__ = [
    'validate_required',
    'validate_number',
    'validate_date_format',
    'validate_enum',
    'validate_reference',
    'validate_at_least_one',
    'generate_validation_failure_excel',
    'save_validation_temp_file',
    'get_temp_filename_from_path',
    'cleanup_temp_files',
    'validate_file_size',
]
