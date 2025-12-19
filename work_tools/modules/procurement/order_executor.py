"""
订单执行人修改视图
"""
import os
import json
import logging
from datetime import datetime
from django.shortcuts import render
from django.http import JsonResponse
from django.conf import settings
from work_tools.forms import OrderExecutorUpdateForm
from work_tools.models import UserOrgDetail
from work_tools.navigation import get_sidebar_groups
from work_tools.views.base import save_sql_file, parse_ops_remark
from work_tools.validation_utils import (
    validate_required,
    validate_at_least_one,
    validate_reference,
    generate_validation_failure_excel,
    get_temp_filename_from_path
)

view_logger = logging.getLogger('work_tools.view')


def validate_order_executor_records(records):
    """校验订单执行人修改记录"""
    results = []
    passed_count = 0
    failed_count = 0
    
    for idx, record in enumerate(records):
        row_number = idx + 2  # 跳过表头
        errors = []
        
        # 至少有一项必填:采购包编号、订单号
        error = validate_at_least_one(
            [record.get('purchase_package_no'), record.get('order_id')],
            ['采购包编号', '订单号']
        )
        if error:
            errors.append(error)
        
        # 必填项校验:订单执行人账号必填
        error = validate_required(record.get('order_executor'), '订单执行人账号')
        if error:
            errors.append(error)
        else:
            # 校验账号在UserOrgDetail表中存在
            error = validate_reference(
                UserOrgDetail,
                'login_name',
                record.get('order_executor'),
                '订单执行人账号'
            )
            if error:
                errors.append(error)
        
        # 原订单执行人账号和名称允许为空，允许查询不到（因为可能原本就是空）
        # 不对原订单执行人进行强制校验
        
        # 记录结果
        if errors:
            results.append({
                'row_number': row_number,
                'valid': False,
                'errors': errors
            })
            failed_count += 1
        else:
            results.append({
                'row_number': row_number,
                'valid': True,
                'errors': []
            })
            passed_count += 1
    
    return {
        'valid': failed_count == 0,
        'total': len(records),
        'passed': passed_count,
        'failed': failed_count,
        'results': results
    }


def order_executor_update_view(request):
    """订单执行人修改视图"""
    saved_file = None
    validation_failure = None  # 新增
    
    # 处理清除请求
    if request.method == 'GET' and request.GET.get('clear') == '1':
        form = OrderExecutorUpdateForm()
        return render(request, 'modules/procurement/order_executor_form.html', {
            'form': form,
            'active_menu': 'order_executor_update',
            'sidebar_groups': get_sidebar_groups(),
        })
    elif request.method == 'POST':
        form = OrderExecutorUpdateForm(request.POST, request.FILES)
        if form.is_valid():
            dynamic_id = form.cleaned_data['dynamic_id']
            ops_remark_raw = form.cleaned_data['ops_remark']
            ops_remark = parse_ops_remark(ops_remark_raw)
            excel_file = form.cleaned_data.get('excel_file')
            
            view_logger.info(f"[订单执行人修改] 开始处理 - 动态编号:{dynamic_id}")
            
            if excel_file:
                # 批量导入
                try:
                    result = _process_excel_import(
                        excel_file, dynamic_id, ops_remark, ops_remark_raw, form)
                    if isinstance(result, tuple):
                        saved_file, _, _ = result
                    elif isinstance(result, dict):
                        # 返回了validation_failure
                        validation_failure = result
                    else:
                        return result
                except Exception as e:
                    view_logger.error(f"[订单执行人修改] Excel处理失败: {str(e)}")
                    form.add_error('excel_file', f'处理失败：{str(e)}')
            else:
                # 单条记录处理
                purchase_package_no = form.cleaned_data.get('purchase_package_no')
                order_id = form.cleaned_data.get('order_id')
                order_executor = form.cleaned_data.get('order_executor')
                orig_order_executor = form.cleaned_data.get('orig_order_executor')
                orig_order_executor_name = form.cleaned_data.get('orig_order_executor_name')
                
                data = [{
                    'purchase_package_no': purchase_package_no,
                    'order_id': order_id,
                    'order_executor': order_executor,
                    'orig_order_executor': orig_order_executor,
                    'orig_order_executor_name': orig_order_executor_name,
                }]
                
                sqls, rollback_sqls, validation_failures = _generate_sqls(
                    data, dynamic_id, ops_remark)
                
                if validation_failures:
                    return _handle_validation_failures(
                        validation_failures, dynamic_id, 'single')
                
                if sqls or rollback_sqls:
                    saved_file, _, _ = _save_and_download(
                        sqls, rollback_sqls, dynamic_id, ops_remark_raw)
                else:
                    form.add_error(None, '未生成任何SQL语句，请检查数据')
        
        return render(request, 'modules/procurement/order_executor_form.html', {
            'form': form,
            'saved_file': saved_file,
            'validation_failure': validation_failure,
            'active_menu': 'order_executor_update',
            'sidebar_groups': get_sidebar_groups(),
        })
    else:
        form = OrderExecutorUpdateForm()
        return render(request, 'modules/procurement/order_executor_form.html', {
            'form': form,
            'active_menu': 'order_executor_update',
            'sidebar_groups': get_sidebar_groups(),
        })


def _process_excel_import(excel_file, dynamic_id, ops_remark, ops_remark_raw, form):
    """处理Excel批量导入"""
    import openpyxl
    
    wb = openpyxl.load_workbook(excel_file, data_only=True)
    ws = wb.active
    
    # 读取表头（支持忽略大小写）
    headers_row = list(ws.iter_rows(min_row=1, max_row=1, values_only=True))[0]
    headers_lower = {str(h).strip().lower(): str(h).strip() 
                     for h in headers_row if h}
    
    # 建立列名映射
    col_map = {}
    for target_key, possible_names in {
        'purchase_package_no': ['采购包编号', 'purchase_package_no'],
        'order_id': ['订单号', 'order_id'],
        'order_executor': ['订单执行人账号', 'order_executor'],
        'orig_order_executor': ['原订单执行人账号', 'orig_order_executor'],
        'orig_order_executor_name': ['原订单执行人名称', 'orig_order_executor_name'],
    }.items():
        for name in possible_names:
            if name.lower() in headers_lower:
                col_map[target_key] = headers_lower[name.lower()]
                break
    
    data = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        row_dict = dict(zip(headers_row, row))
        
        purchase_package_no = str(row_dict.get(col_map.get('purchase_package_no', ''), '')).strip()
        order_id = str(row_dict.get(col_map.get('order_id', ''), '')).strip()
        order_executor = str(row_dict.get(col_map.get('order_executor', ''), '')).strip()
        
        # 至少要有一个编号
        if not purchase_package_no and not order_id:
            continue
        
        # 必须有订单执行人
        if not order_executor:
            continue
        
        data.append({
            'purchase_package_no': purchase_package_no,
            'order_id': order_id,
            'order_executor': order_executor,
            'orig_order_executor': str(row_dict.get(col_map.get('orig_order_executor', ''), '')).strip(),
            'orig_order_executor_name': str(row_dict.get(col_map.get('orig_order_executor_name', ''), '')).strip(),
            'row_idx': row_idx,
        })
    
    if not data:
        raise ValueError('Excel中没有有效数据行')
    
    # 执行数据校验
    validation_result = validate_order_executor_records(data)
    
    if not validation_result['valid']:
        # 校验失败,生成失败文件
        temp_file = generate_validation_failure_excel(
            excel_file,
            validation_result,
            'order_executor'
        )
        
        if temp_file:
            return {
                'total': validation_result['total'],
                'passed': validation_result['passed'],
                'failed': validation_result['failed'],
                'filename': get_temp_filename_from_path(temp_file)
            }
    
    sqls, rollback_sqls, validation_failures = _generate_sqls(
        data, dynamic_id, ops_remark)
    
    if validation_failures:
        return _handle_validation_failures(
            validation_failures, dynamic_id, 'batch')
    
    if sqls or rollback_sqls:
        return _save_and_download(sqls, rollback_sqls, dynamic_id, ops_remark_raw)
    else:
        raise ValueError('未生成任何SQL语句，请检查数据')


def _generate_sqls(data, dynamic_id, ops_remark):
    """生成SQL语句"""
    from work_tools.config import get_config
    
    sqls = []
    rollback_sqls = []
    validation_failures = []
    
    cfg = get_config()
    
    for item in data:
        purchase_package_no = item.get('purchase_package_no')
        order_id = item.get('order_id')
        order_executor = item.get('order_executor')
        orig_order_executor = item.get('orig_order_executor', '')
        
        # 验证必填项
        if not order_executor:
            row_idx = item.get('row_idx', '')
            validation_failures.append({
                'row': row_idx,
                'error': '订单执行人账号为必填项'
            })
            continue
        
        if not purchase_package_no and not order_id:
            row_idx = item.get('row_idx', '')
            validation_failures.append({
                'row': row_idx,
                'error': '采购包编号和订单号至少需要一个'
            })
            continue
        
        # 根据order_executor从user_org_detail查询信息
        user_org = UserOrgDetail.objects.filter(login_name=order_executor).first()
        if not user_org:
            row_idx = item.get('row_idx', '')
            validation_failures.append({
                'row': row_idx,
                'order_executor': order_executor,
                'error': f'未找到订单执行人信息：{order_executor}'
            })
            continue
        
        # 执行SQL
        if purchase_package_no:
            # 更新采购包表
            sql = (f"UPDATE tprly02 SET ORDER_EXECUTOR = '{order_executor}',"
                   f"ORDER_EXECUTOR_NAME='{user_org.user_name or ''}',"
                   f"OPS_REMARK='{ops_remark}' "
                   f"WHERE PURCHASE_PACKAGE_NO IN ('{purchase_package_no}');")
            sqls.append(sql)
        
        if order_id:
            # 更新订单主表 tpodd01
            sql_dd01 = (f"update tpodd01 set "
                       f"PUR_ORG_ID='{user_org.company_code or ''}',"
                       f"PUR_ORG_NAME='{user_org.company_name or ''}',"
                       f"CREATE_DEPT_CODE='{user_org.dept_code or ''}',"
                       f"CREATE_DEPT_NAME='{user_org.dept_name or ''}',"
                       f"CREATE_ORG_CODE='{user_org.company_code or ''}',"
                       f"CREATE_ORG_NAME='{user_org.company_name or ''}',"
                       f"CREATE_USER='{order_executor}',"
                       f"CREATE_USER_ID='{order_executor}',"
                       f"CREATE_USER_NAME='{user_org.user_name or ''}',"
                       f"CREATE_PLATE_CODE='{user_org.plate_code or ''}',"
                       f"CREATE_PLATE_NAME='{user_org.plate_name or ''}',"
                       f"OPS_REMARK='{ops_remark}' "
                       f"where ORDER_ID='{order_id}';")
            sqls.append(sql_dd01)
            
            # 更新订单明细表 tpodd02
            sql_dd02 = (f"update tpodd02 set "
                       f"PUR_ORG_ID='{user_org.company_code or ''}',"
                       f"PUR_ORG_NAME='{user_org.company_name or ''}',"
                       f"CREATE_DEPT_CODE='{user_org.dept_code or ''}',"
                       f"CREATE_DEPT_NAME='{user_org.dept_name or ''}',"
                       f"CREATE_ORG_CODE='{user_org.company_code or ''}',"
                       f"CREATE_ORG_NAME='{user_org.company_name or ''}',"
                       f"CREATE_USER='{order_executor}',"
                       f"CREATE_USER_ID='{order_executor}',"
                       f"CREATE_USER_NAME='{user_org.user_name or ''}',"
                       f"CREATE_PLATE_CODE='{user_org.plate_code or ''}',"
                       f"CREATE_PLATE_NAME='{user_org.plate_name or ''}',"
                       f"OPS_REMARK='{ops_remark}' "
                       f"where ORDER_ID='{order_id}';")
            sqls.append(sql_dd02)
        
        # 生成回退SQL（支持补充订单执行人场景：原本为空则回退时清空）
        # 情况1：orig_order_executor为空，表示补充订单执行人（原本没有），回退时清空所有字段
        # 情况2：orig_order_executor不为空，需要验证账号存在，回退到原账号信息
        
        if not orig_order_executor:
            # 补充订单执行人场景：原本为空，回退时清空
            if purchase_package_no:
                rollback_sql = (f"UPDATE tprly02 SET ORDER_EXECUTOR = '', "
                               f"ORDER_EXECUTOR_NAME='',"
                               f"OPS_REMARK='' "
                               f"WHERE PURCHASE_PACKAGE_NO IN ('{purchase_package_no}');")
                rollback_sqls.append(rollback_sql)
            
            if order_id:
                # 回退订单主表（清空）
                rollback_dd01 = (f"update tpodd01 set "
                                f"PUR_ORG_ID='',"
                                f"PUR_ORG_NAME='',"
                                f"CREATE_DEPT_CODE='',"
                                f"CREATE_DEPT_NAME='',"
                                f"CREATE_ORG_CODE='',"
                                f"CREATE_ORG_NAME='',"
                                f"CREATE_USER='',"
                                f"CREATE_USER_ID='',"
                                f"CREATE_USER_NAME='',"
                                f"CREATE_PLATE_CODE='',"
                                f"CREATE_PLATE_NAME='',"
                                f"OPS_REMARK='' "
                                f"where ORDER_ID='{order_id}';")
                rollback_sqls.append(rollback_dd01)
                
                # 回退订单明细表（清空）
                rollback_dd02 = (f"update tpodd02 set "
                                f"PUR_ORG_ID='',"
                                f"PUR_ORG_NAME='',"
                                f"CREATE_DEPT_CODE='',"
                                f"CREATE_DEPT_NAME='',"
                                f"CREATE_ORG_CODE='',"
                                f"CREATE_ORG_NAME='',"
                                f"CREATE_USER='',"
                                f"CREATE_USER_ID='',"
                                f"CREATE_USER_NAME='',"
                                f"CREATE_PLATE_CODE='',"
                                f"CREATE_PLATE_NAME='',"
                                f"OPS_REMARK='' "
                                f"where ORDER_ID='{order_id}';")
                rollback_sqls.append(rollback_dd02)
        else:
            # 修改订单执行人场景：原本有值，回退到原值
            # 查询原执行人信息
            orig_user_org = UserOrgDetail.objects.filter(login_name=orig_order_executor).first()
            
            # 如果原订单执行人账号无效，使用提供的名称或空值，不报错
            if not orig_user_org:
                # 使用提供的原订单执行人名称，如果没有则为空
                orig_executor_name = item.get('orig_order_executor_name', '')
                
                if purchase_package_no:
                    rollback_sql = (f"UPDATE tprly02 SET ORDER_EXECUTOR = '{orig_order_executor}', "
                                   f"ORDER_EXECUTOR_NAME='{orig_executor_name}',"
                                   f"OPS_REMARK='' "
                                   f"WHERE PURCHASE_PACKAGE_NO IN ('{purchase_package_no}');")
                    rollback_sqls.append(rollback_sql)
                
                if order_id:
                    # 回退订单主表（使用提供的信息或空值）
                    rollback_dd01 = (f"update tpodd01 set "
                                    f"PUR_ORG_ID='',"
                                    f"PUR_ORG_NAME='',"
                                    f"CREATE_DEPT_CODE='',"
                                    f"CREATE_DEPT_NAME='',"
                                    f"CREATE_ORG_CODE='',"
                                    f"CREATE_ORG_NAME='',"
                                    f"CREATE_USER='{orig_order_executor}',"
                                    f"CREATE_USER_ID='{orig_order_executor}',"
                                    f"CREATE_USER_NAME='{orig_executor_name}',"
                                    f"CREATE_PLATE_CODE='',"
                                    f"CREATE_PLATE_NAME='',"
                                    f"OPS_REMARK='' "
                                    f"where ORDER_ID='{order_id}';")
                    rollback_sqls.append(rollback_dd01)
                    
                    # 回退订单明细表（使用提供的信息或空值）
                    rollback_dd02 = (f"update tpodd02 set "
                                    f"PUR_ORG_ID='',"
                                    f"PUR_ORG_NAME='',"
                                    f"CREATE_DEPT_CODE='',"
                                    f"CREATE_DEPT_NAME='',"
                                    f"CREATE_ORG_CODE='',"
                                    f"CREATE_ORG_NAME='',"
                                    f"CREATE_USER='{orig_order_executor}',"
                                    f"CREATE_USER_ID='{orig_order_executor}',"
                                    f"CREATE_USER_NAME='{orig_executor_name}',"
                                    f"CREATE_PLATE_CODE='',"
                                    f"CREATE_PLATE_NAME='',"
                                    f"OPS_REMARK='' "
                                    f"where ORDER_ID='{order_id}';")
                    rollback_sqls.append(rollback_dd02)
                continue
            
            if purchase_package_no:
                rollback_sql = (f"UPDATE tprly02 SET ORDER_EXECUTOR = '{orig_order_executor}', "
                               f"ORDER_EXECUTOR_NAME='{orig_user_org.user_name or ''}',"
                               f"OPS_REMARK='' "
                               f"WHERE PURCHASE_PACKAGE_NO IN ('{purchase_package_no}');")
                rollback_sqls.append(rollback_sql)
            
            if order_id:
                # 回退订单主表
                rollback_dd01 = (f"update tpodd01 set "
                                f"PUR_ORG_ID='{orig_user_org.company_code or ''}',"
                                f"PUR_ORG_NAME='{orig_user_org.company_name or ''}',"
                                f"CREATE_DEPT_CODE='{orig_user_org.dept_code or ''}',"
                                f"CREATE_DEPT_NAME='{orig_user_org.dept_name or ''}',"
                                f"CREATE_ORG_CODE='{orig_user_org.company_code or ''}',"
                                f"CREATE_ORG_NAME='{orig_user_org.company_name or ''}',"
                                f"CREATE_USER='{orig_order_executor}',"
                                f"CREATE_USER_ID='{orig_order_executor}',"
                                f"CREATE_USER_NAME='{orig_user_org.user_name or ''}',"
                                f"CREATE_PLATE_CODE='{orig_user_org.plate_code or ''}',"
                                f"CREATE_PLATE_NAME='{orig_user_org.plate_name or ''}',"
                                f"OPS_REMARK='' "
                                f"where ORDER_ID='{order_id}';")
                rollback_sqls.append(rollback_dd01)
                
                # 回退订单明细表
                rollback_dd02 = (f"update tpodd02 set "
                                f"PUR_ORG_ID='{orig_user_org.company_code or ''}',"
                                f"PUR_ORG_NAME='{orig_user_org.company_name or ''}',"
                                f"CREATE_DEPT_CODE='{orig_user_org.dept_code or ''}',"
                                f"CREATE_DEPT_NAME='{orig_user_org.dept_name or ''}',"
                                f"CREATE_ORG_CODE='{orig_user_org.company_code or ''}',"
                                f"CREATE_ORG_NAME='{orig_user_org.company_name or ''}',"
                                f"CREATE_USER='{orig_order_executor}',"
                                f"CREATE_USER_ID='{orig_order_executor}',"
                                f"CREATE_USER_NAME='{orig_user_org.user_name or ''}',"
                                f"CREATE_PLATE_CODE='{orig_user_org.plate_code or ''}',"
                                f"CREATE_PLATE_NAME='{orig_user_org.plate_name or ''}',"
                                f"OPS_REMARK='' "
                                f"where ORDER_ID='{order_id}';")
                rollback_sqls.append(rollback_dd02)
    
    return sqls, rollback_sqls, validation_failures


def _save_and_download(sqls, rollback_sqls, dynamic_id, ops_remark_raw):
    """保存SQL文件到固定目录（静默输出）"""
    
    # 生成SQL内容（执行语句 + 回退语句合并在一个文件中）
    sql_content = "1、执行语句\n\n"
    
    for sql in sqls:
        sql_content += sql + '\n'
    
    sql_content += "\n\n\n\n\n\n2.回退语句\n\n"
    
    for sql in rollback_sqls:
        sql_content += sql + '\n'
    
    # 添加数据库信息
    sql_content += "\n\n\n3.数据库\n\nip：192.168.11.69\n\n库名：cnnc_pr"
    
    # 保存文件到固定目录
    filepath = save_sql_file(sql_content, '订单执行人修改', dynamic_id)
    
    return (filepath, None, None)


def _handle_validation_failures(failures, dynamic_id, mode):
    """处理校验失败"""
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    failure_filename = f'order_executor_failures_{timestamp}.json'
    
    failure_path = os.path.join(
        settings.BASE_DIR, 'sql_output', failure_filename)
    os.makedirs(os.path.dirname(failure_path), exist_ok=True)
    
    with open(failure_path, 'w', encoding='utf-8') as f:
        json.dump(failures, f, ensure_ascii=False, indent=2)
    
    return JsonResponse({
        'status': 'validation_failed',
        'message': f'有 {len(failures)} 条数据校验失败',
        'failure_file': failure_filename,
        'failures': failures
    })


def download_order_executor_template(request):
    """下载订单执行人修改Excel模板"""
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    from django.http import HttpResponse
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "订单执行人修改"
    
    # 表头
    headers = ['采购包编号', '订单号', '订单执行人账号', '原订单执行人账号']
    ws.append(headers)
    
    # 样式设置
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", 
                              end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
    
    # 示例数据
    ws.append(['CNCH-CGB-25-00004', '', 'cnch_zmw01', '1891855700607'])
    ws.append(['', 'CABX-25-00027-DD-0001', 'cnch_zmw01', '1891855700607'])
    
    # 调整列宽
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[column].width = max_length + 2
    
    # 生成响应
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="order_executor_template.xlsx"'
    wb.save(response)
    return response


def validate_order_executor_api(request):
    """验证订单执行人账号是否存在"""
    login_name = request.GET.get('login_name', '').strip()
    
    if not login_name:
        return JsonResponse({
            'valid': False,
            'message': '请输入订单执行人账号'
        })
    
    # 查询user_org_detail表
    user_org = UserOrgDetail.objects.filter(login_name=login_name).first()
    
    if not user_org:
        return JsonResponse({
            'valid': False,
            'message': f'账号 {login_name} 不存在，请检查后重试'
        })
    
    return JsonResponse({
        'valid': True,
        'message': f'账号有效：{user_org.user_name or login_name}',
        'user_name': user_org.user_name or '',
        'company_name': user_org.company_name or '',
        'dept_name': user_org.dept_name or ''
    })


__all__ = ['order_executor_update_view', 'download_order_executor_template', 'validate_order_executor_api']
