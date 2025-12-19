"""需求计划明细日期修改模块"""
import os
import json
from datetime import datetime
from django.shortcuts import render
from django.http import HttpResponse, JsonResponse
from django.conf import settings

from work_tools.forms import PlanDateUpdateForm
from work_tools.navigation import get_sidebar_groups
from work_tools.views.base import parse_ops_remark, save_sql_file
from work_tools.validation_utils import (
    validate_required,
    validate_at_least_one,
    validate_date_format,
    generate_validation_failure_excel,
    get_temp_filename_from_path
)

# 日志记录器
import logging
view_logger = logging.getLogger('app.view')


def validate_plan_date_records(records):
    """校验需求计划明细日期修改记录"""
    results = []
    passed_count = 0
    failed_count = 0
    
    for idx, record in enumerate(records):
        row_number = idx + 2  # 跳过表头
        errors = []
        
        # 必填项校验:明细行号必填
        error = validate_required(record.get('plan_line_no'), '明细行号')
        if error:
            errors.append(error)
        
        # 至少有一项必填:需用日期、开始日期、结束日期
        error = validate_at_least_one(
            [record.get('required_date'), record.get('start_date'), record.get('end_date')],
            ['需用日期', '开始日期', '结束日期']
        )
        if error:
            errors.append(error)
        
        # 日期格式校验(YYYYMMDD)
        if record.get('required_date'):
            error = validate_date_format(record.get('required_date'), '需用日期')
            if error:
                errors.append(error)
        
        if record.get('start_date'):
            error = validate_date_format(record.get('start_date'), '开始日期')
            if error:
                errors.append(error)
        
        if record.get('end_date'):
            error = validate_date_format(record.get('end_date'), '结束日期')
            if error:
                errors.append(error)
        
        # 记录结果
        if errors:
            results.append({
                'row_number': row_number,
                'valid': False,
                'errors': errors
            })
            failed_count += 1
        else:
            results.append({
                'row_number': row_number,
                'valid': True,
                'errors': []
            })
            passed_count += 1
    
    return {
        'valid': failed_count == 0,
        'total': len(records),
        'passed': passed_count,
        'failed': failed_count,
        'results': results
    }


def plan_date_update_view(request):
    """需求计划明细日期修改视图"""
    saved_file = None
    validation_failure = None  # 新增
    
    # 处理清除请求
    if request.method == 'GET' and request.GET.get('clear') == '1':
        form = PlanDateUpdateForm()
        return render(request, 'modules/procurement/plan_date_form.html', {
            'form': form,
            'active_menu': 'plan_date_update',
            'sidebar_groups': get_sidebar_groups(),
        })
    elif request.method == 'POST':
        form = PlanDateUpdateForm(request.POST, request.FILES)
        if form.is_valid():
            dynamic_id = form.cleaned_data['dynamic_id']
            ops_remark_raw = form.cleaned_data['ops_remark']
            ops_remark = parse_ops_remark(ops_remark_raw)
            excel_file = form.cleaned_data.get('excel_file')
            
            view_logger.info(f"[需求计划日期修改] 开始处理 - 动态编号:{dynamic_id}")
            
            if excel_file:
                # 批量导入
                try:
                    result = _process_excel_import(
                        excel_file, dynamic_id, ops_remark, ops_remark_raw, form)
                    if isinstance(result, tuple):
                        saved_file, _, _ = result
                    elif isinstance(result, dict):
                        # 返回了validation_failure
                        validation_failure = result
                    else:
                        return result
                except Exception as e:
                    view_logger.error(f"[需求计划日期修改] Excel处理失败: {str(e)}")
                    form.add_error('excel_file', f'处理失败：{str(e)}')
            else:
                # 单条记录处理
                plan_line_no = form.cleaned_data['plan_line_no']
                required_date = form.cleaned_data.get('required_date')
                start_date = form.cleaned_data.get('start_date')
                end_date = form.cleaned_data.get('end_date')
                orig_required_date = form.cleaned_data.get('orig_required_date')
                orig_start_date = form.cleaned_data.get('orig_start_date')
                orig_end_date = form.cleaned_data.get('orig_end_date')
                
                data = [{
                    'plan_line_no': plan_line_no,
                    'required_date': required_date,
                    'start_date': start_date,
                    'end_date': end_date,
                    'orig_required_date': orig_required_date,
                    'orig_start_date': orig_start_date,
                    'orig_end_date': orig_end_date,
                }]
                
                sqls, rollback_sqls, validation_failures = _generate_sqls(
                    data, dynamic_id, ops_remark)
                
                if validation_failures:
                    return _handle_validation_failures(
                        validation_failures, dynamic_id, 'single')
                
                if sqls or rollback_sqls:
                    saved_file, _, _ = _save_and_download(
                        sqls, rollback_sqls, dynamic_id, ops_remark_raw)
                else:
                    form.add_error(None, '未生成任何SQL语句，请检查数据')
        
        return render(request, 'modules/procurement/plan_date_form.html', {
            'form': form,
            'saved_file': saved_file,
            'validation_failure': validation_failure,
            'active_menu': 'plan_date_update',
            'sidebar_groups': get_sidebar_groups(),
        })
    else:
        form = PlanDateUpdateForm()
        return render(request, 'modules/procurement/plan_date_form.html', {
            'form': form,
            'active_menu': 'plan_date_update',
            'sidebar_groups': get_sidebar_groups(),
        })


def _process_excel_import(excel_file, dynamic_id, ops_remark, ops_remark_raw, form):
    """处理Excel批量导入"""
    import openpyxl
    
    wb = openpyxl.load_workbook(excel_file, data_only=True)
    ws = wb.active
    
    # 读取表头（支持忽略大小写）
    headers_row = list(ws.iter_rows(min_row=1, max_row=1, values_only=True))[0]
    headers_lower = {str(h).strip().lower(): str(h).strip() 
                     for h in headers_row if h}
    
    # 必需列（忽略大小写）
    required_cols = {'明细行号', 'plan_line_no'}
    header_keys = set(headers_lower.keys())
    
    if not any(col in header_keys for col in required_cols):
        raise ValueError('Excel必须包含"明细行号"或"plan_line_no"列')
    
    # 建立列名映射
    col_map = {}
    for target_key, possible_names in {
        'plan_line_no': ['明细行号', 'plan_line_no'],
        'required_date': ['需用日期', 'required_date'],
        'start_date': ['开始日期', 'start_date'],
        'end_date': ['结束日期', 'end_date'],
        'orig_required_date': ['原需用日期', 'orig_required_date'],
        'orig_start_date': ['原开始日期', 'orig_start_date'],
        'orig_end_date': ['原结束日期', 'orig_end_date'],
    }.items():
        for name in possible_names:
            if name.lower() in headers_lower:
                col_map[target_key] = headers_lower[name.lower()]
                break
    
    data = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        row_dict = dict(zip(headers_row, row))
        
        plan_line_no = str(row_dict.get(col_map.get('plan_line_no', ''), '')).strip()
        if not plan_line_no:
            continue
        
        data.append({
            'plan_line_no': plan_line_no,
            'required_date': str(row_dict.get(col_map.get('required_date', ''), '')).strip(),
            'start_date': str(row_dict.get(col_map.get('start_date', ''), '')).strip(),
            'end_date': str(row_dict.get(col_map.get('end_date', ''), '')).strip(),
            'orig_required_date': str(row_dict.get(col_map.get('orig_required_date', ''), '')).strip(),
            'orig_start_date': str(row_dict.get(col_map.get('orig_start_date', ''), '')).strip(),
            'orig_end_date': str(row_dict.get(col_map.get('orig_end_date', ''), '')).strip(),
            'row_idx': row_idx,
        })
    
    if not data:
        raise ValueError('Excel中没有有效数据行')
    
    # 执行数据校验
    validation_result = validate_plan_date_records(data)
    
    if not validation_result['valid']:
        # 校验失败,生成失败文件
        temp_file = generate_validation_failure_excel(
            excel_file,
            validation_result,
            'plan_date'
        )
        
        if temp_file:
            return {
                'total': validation_result['total'],
                'passed': validation_result['passed'],
                'failed': validation_result['failed'],
                'filename': get_temp_filename_from_path(temp_file)
            }
    
    sqls, rollback_sqls, validation_failures = _generate_sqls(
        data, dynamic_id, ops_remark)
    
    if validation_failures:
        return _handle_validation_failures(
            validation_failures, dynamic_id, 'batch')
    
    if sqls or rollback_sqls:
        return _save_and_download(sqls, rollback_sqls, dynamic_id, ops_remark_raw)
    else:
        raise ValueError('未生成任何SQL语句，请检查数据')


def _generate_sqls(data, dynamic_id, ops_remark):
    """生成SQL语句（按需求计划-询价单分组）"""
    from work_tools.config import get_config
    
    sqls = []
    rollback_sqls = []
    validation_failures = []
    
    # 获取配置（暂时不使用SQL合并功能，因为这里是按明细行号单条生成）
    cfg = get_config()
    
    # 按明细行号分组
    grouped_data = {}
    for item in data:
        plan_line_no = item['plan_line_no']
        if plan_line_no not in grouped_data:
            grouped_data[plan_line_no] = []
        grouped_data[plan_line_no].append(item)
    
    # 为每个明细行号生成SQL
    for plan_line_no, items in grouped_data.items():
        for item in items:
            required_date = item.get('required_date')
            start_date = item.get('start_date')
            end_date = item.get('end_date')
            orig_required_date = item.get('orig_required_date')
            orig_start_date = item.get('orig_start_date')
            orig_end_date = item.get('orig_end_date')
            
            # 检查至少有一个日期字段
            if not any([required_date, start_date, end_date]):
                row_idx = item.get('row_idx', '')
                validation_failures.append({
                    'row': row_idx,
                    'plan_line_no': plan_line_no,
                    'error': '至少需要填写一个日期字段'
                })
                continue
            
            # 生成执行SQL（4个表）
            # 1. 需求计划表 tprnd02
            if required_date:
                sql_nd02 = f"update tprnd02 set required_date='{required_date}',ops_remark='{ops_remark}' where plan_line_no='{plan_line_no}';"
                sqls.append(('需求计划', sql_nd02))
            
            # 2. 采购包表 tprly02
            set_fields_ly02 = [f"OPS_REMARK='{ops_remark}'"]
            if required_date:
                set_fields_ly02.append(f"REQUIRED_DATE='{required_date}'")
            if start_date:
                set_fields_ly02.append(f"start_date='{start_date}'")
            if end_date:
                set_fields_ly02.append(f"end_date='{end_date}'")
            
            sql_ly02 = f"update tprly02 set {','.join(set_fields_ly02)} where plan_line_no='{plan_line_no}' and ALIVE_FLAG='1';"
            sqls.append(('采购包', sql_ly02))
            
            # 3. 采购方案表 tprfa03
            set_fields_fa03 = [f"OPS_REMARK='{ops_remark}'"]
            if required_date:
                set_fields_fa03.append(f"REQUIRED_DATE='{required_date}'")
            if start_date:
                set_fields_fa03.append(f"start_date='{start_date}'")
            if end_date:
                set_fields_fa03.append(f"end_date='{end_date}'")
            
            sql_fa03 = f"update tprfa03 set {','.join(set_fields_fa03)} where plan_line_no='{plan_line_no}' and ALIVE_FLAG='1';"
            sqls.append(('采购方案', sql_fa03))
            
            # 4. 询价单表 tprxj05
            set_fields_xj05 = [f"OPS_REMARK='{ops_remark}'"]
            if required_date:
                set_fields_xj05.append(f"REQUIRED_DATE='{required_date}'")
            if start_date:
                set_fields_xj05.append(f"start_date='{start_date}'")
            if end_date:
                set_fields_xj05.append(f"end_date='{end_date}'")
            
            sql_xj05 = f"update tprxj05 set {','.join(set_fields_xj05)} where plan_line_no='{plan_line_no}' and ALIVE_FLAG='1';"
            sqls.append(('询价单', sql_xj05))
            
            # 生成回退SQL（无论是否有原值，都生成回退语句）
            # 1. 需求计划回退（只有required_date）
            if required_date:  # 只有修改了required_date才生成回退
                rollback_nd02 = f"update tprnd02 set required_date='{orig_required_date or ''}',ops_remark='{ops_remark}' where plan_line_no='{plan_line_no}';"
                rollback_sqls.append(('需求计划', rollback_nd02))
            
            # 2. 采购包回退
            rollback_fields_ly02 = [f"OPS_REMARK='{ops_remark}'"]
            if required_date:  # 修改了就加回退字段
                rollback_fields_ly02.append(f"REQUIRED_DATE='{orig_required_date or ''}'")
            if start_date:
                rollback_fields_ly02.append(f"start_date='{orig_start_date or ''}'")
            if end_date:
                rollback_fields_ly02.append(f"end_date='{orig_end_date or ''}'")
            
            if len(rollback_fields_ly02) > 1:  # 有修改字段才生成
                rollback_ly02 = f"update tprly02 set {','.join(rollback_fields_ly02)} where plan_line_no='{plan_line_no}' and ALIVE_FLAG='1';"
                rollback_sqls.append(('采购包', rollback_ly02))
            
            # 3. 采购方案回退
            rollback_fields_fa03 = [f"OPS_REMARK='{ops_remark}'"]
            if required_date:
                rollback_fields_fa03.append(f"REQUIRED_DATE='{orig_required_date or ''}'")
            if start_date:
                rollback_fields_fa03.append(f"start_date='{orig_start_date or ''}'")
            if end_date:
                rollback_fields_fa03.append(f"end_date='{orig_end_date or ''}'")
            
            if len(rollback_fields_fa03) > 1:
                rollback_fa03 = f"update tprfa03 set {','.join(rollback_fields_fa03)} where plan_line_no='{plan_line_no}' and ALIVE_FLAG='1';"
                rollback_sqls.append(('采购方案', rollback_fa03))
            
            # 4. 询价单回退
            rollback_fields_xj05 = [f"OPS_REMARK='{ops_remark}'"]
            if required_date:
                rollback_fields_xj05.append(f"REQUIRED_DATE='{orig_required_date or ''}'")
            if start_date:
                rollback_fields_xj05.append(f"start_date='{orig_start_date or ''}'")
            if end_date:
                rollback_fields_xj05.append(f"end_date='{orig_end_date or ''}'")
            
            if len(rollback_fields_xj05) > 1:
                rollback_xj05 = f"update tprxj05 set {','.join(rollback_fields_xj05)} where plan_line_no='{plan_line_no}' and ALIVE_FLAG='1';"
                rollback_sqls.append(('询价单', rollback_xj05))
    
    # 按表分组排序（需求计划-采购包-采购方案-询价单）
    table_order = {'需求计划': 1, '采购包': 2, '采购方案': 3, '询价单': 4}
    sqls.sort(key=lambda x: table_order.get(x[0], 999))
    rollback_sqls.sort(key=lambda x: table_order.get(x[0], 999))
    
    return sqls, rollback_sqls, validation_failures


def _save_and_download(sqls, rollback_sqls, dynamic_id, ops_remark_raw):
    """保存SQL文件到固定目录（静默输出）"""
    
    # 生成SQL内容（执行语句 + 回退语句合并在一个文件中）
    sql_content = "1、执行语句\n\n"
    
    # 执行语句部分（按表分组）
    current_table = None
    for table_name, sql in sqls:
        if table_name != current_table:
            sql_content += f"--{table_name}\n"
            current_table = table_name
        sql_content += sql + '\n'
    
    # 分隔符
    sql_content += "\n\n\n\n\n\n2.回退语句\n\n"
    
    # 回退语句部分（带注释）
    current_table = None
    for table_name, sql in rollback_sqls:
        if table_name != current_table:
            sql_content += f"--{table_name}\n"
            current_table = table_name
        sql_content += sql + '\n'
    
    # 添加数据库信息
    sql_content += "\n\n\n3.数据库\n\nip：192.168.11.76\n\n库名：cnnc_ph"
    
    # 保存文件到固定目录
    filepath = save_sql_file(sql_content, '需求计划明细日期修改', dynamic_id)
    
    return (filepath, None, None)


def _handle_validation_failures(failures, dynamic_id, mode):
    """处理校验失败"""
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    failure_filename = f'plan_date_failures_{timestamp}.json'
    
    failure_path = os.path.join(
        settings.BASE_DIR, 'sql_output', failure_filename)
    os.makedirs(os.path.dirname(failure_path), exist_ok=True)
    
    with open(failure_path, 'w', encoding='utf-8') as f:
        json.dump(failures, f, ensure_ascii=False, indent=2)
    
    return JsonResponse({
        'status': 'validation_failed',
        'message': f'有 {len(failures)} 条数据校验失败',
        'failure_file': failure_filename,
        'failures': failures
    })


def download_plan_date_template(request):
    """下载需求计划日期修改Excel模板"""
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "需求计划日期修改"
    
    # 表头
    headers = ['明细行号', '需用日期', '开始日期', '结束日期', 
               '原需用日期', '原开始日期', '原结束日期']
    ws.append(headers)
    
    # 样式设置
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", 
                              end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
    
    # 示例数据
    ws.append(['ZH25-XQJH-25-00028-000001', '20250531', '20250424', 
               '20251231', '20250430', '20250320', '20251130'])
    ws.append(['ZJXC-XQJH-25-00758-000001', '20250902', '', '', 
               '20250801', '', ''])
    
    # 调整列宽
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[column].width = max_length + 2
    
    # 生成响应
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="plan_date_template.xlsx"'
    wb.save(response)
    return response


__all__ = ['plan_date_update_view', 'download_plan_date_template']
