"""" 
合同状态修改模块
包含合同状态修改的所有视图函数和SQL生成逻辑
"""
import io
import os
import logging
from django.shortcuts import render
from django.http import FileResponse
from work_tools.forms import ApprStateChangeForm, APPR_STATE_CHOICES
from work_tools.navigation import get_sidebar_groups
from work_tools.config import get_config
from work_tools.sql_merge import chunk_list, format_in, merge_by_key
from work_tools.logger_utils import log_view_input
from work_tools.views.base import parse_ops_remark, save_sql_file
from work_tools.validation_utils import (
    validate_required,
    generate_validation_failure_excel,
    get_temp_filename_from_path
)

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

logger = logging.getLogger('work_tools.view')
sql_logger = logging.getLogger('work_tools.sql')


def parse_appr_state_excel(file):
    """解析Excel文件,提取合同状态修改数据"""
    if not OPENPYXL_AVAILABLE:
        raise ValueError("缺少 openpyxl，无法解析 Excel")

    # 从数据库加载合同状态配置构建映射字典
    try:
        from work_tools.dropdown_utils import get_dropdown_options
        options = get_dropdown_options('contract_status', include_empty=False)
        STATE_MAP = {label: code for code, label in options if code}
    except Exception:
        # 回退到硬编码映射
        STATE_MAP = {v: k for k, v in APPR_STATE_CHOICES if k}

    wb = openpyxl.load_workbook(file)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    idx = {str(h): i for i, h in enumerate(headers)}

    def pick(alts):
        for a in alts:
            if a in idx:
                return idx[a]
        return None

    bpo_id_idx = pick(['合同ID', 'BPO_ID', 'bpo_id'])
    new_state_idx = pick(['新合同状态', '新状态', 'new_appr_state'])
    orig_state_idx = pick(['原合同状态', '原状态', 'orig_appr_state'])

    if bpo_id_idx is None or new_state_idx is None:
        raise ValueError("Excel 缺少必要列：合同ID/新合同状态")

    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        bpo_id = row[bpo_id_idx]
        new_state = row[new_state_idx]
        if bpo_id is None or new_state is None:
            continue

        # 转换状态值（支持中文和英文）
        new_state_val = str(new_state).strip()
        if new_state_val in STATE_MAP:
            new_state_val = STATE_MAP[new_state_val]

        rec = {
            'bpo_id': str(bpo_id).strip(),
            'new_appr_state': new_state_val,
        }

        if orig_state_idx is not None:
            orig_state = row[orig_state_idx]
            if orig_state:
                orig_state_val = str(orig_state).strip()
                if orig_state_val in STATE_MAP:
                    orig_state_val = STATE_MAP[orig_state_val]
                rec['orig_appr_state'] = orig_state_val

        records.append(rec)

    if not records:
        raise ValueError("Excel 中没有有效的行数据")

    logger.info(f"[合同状态] Excel解析完成, 记录数={len(records)}")
    return records


def validate_appr_state_records(records):
    """校验合同状态修改记录"""
    results = []
    passed_count = 0
    failed_count = 0
    
    for idx, record in enumerate(records):
        row_number = idx + 2  # 跳过表头
        errors = []
        
        # 必填项校验：合同ID必填
        error = validate_required(record.get('bpo_id'), '合同ID')
        if error:
            errors.append(error)
        
        # 必填项校验：新合同状态必填
        error = validate_required(record.get('new_appr_state'), '新合同状态')
        if error:
            errors.append(error)
        
        # 记录结果
        if errors:
            results.append({
                'row_number': row_number,
                'valid': False,
                'errors': errors
            })
            failed_count += 1
        else:
            results.append({
                'row_number': row_number,
                'valid': True,
                'errors': []
            })
            passed_count += 1
    
    return {
        'valid': failed_count == 0,
        'total': len(records),
        'passed': passed_count,
        'failed': failed_count,
        'results': results
    }


def generate_appr_state_sql_bulk(records, ops_remark=None):
    """
    生成批量合同状态修改SQL - 支持SQL合并策略
    相同状态的记录会自动合并为IN查询
    """
    cfg = get_config()
    sql = []
    sql.append("1、执行语句")

    sql_logger.info(
        f"[SQL合并] 开始生成合同状态SQL, 记录数={len(records)}, 配置启用={cfg.get('MERGE_MODULES', {}).get('appr_state', True)}")

    if cfg.get('MERGE_MODULES', {}).get('appr_state', True):
        # 启用SQL合并策略: 按新状态分组
        def key_fn(r):
            return r.get('new_appr_state')

        groups = merge_by_key(records, key_fn)
        sql_logger.info(f"[SQL合并] 合同状态分组完成, 组数={len(groups)}")

        for new_state, recs in groups.items():
            if not new_state:
                continue

            bpo_ids = [r['bpo_id'] for r in recs if r.get('bpo_id')]
            sql_logger.info(
                f"[SQL合并] 处理分组: new_state={new_state}, 合同数={len(bpo_ids)}")

            for chunk in chunk_list(bpo_ids, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
                inlist = format_in(chunk)

                # 更新主表 tphct01
                stmt_main = (
                    f"UPDATE tphct01 SET APPR_STATE='{new_state}', "
                    f"OPS_REMARK='{ops_remark or ''}' "
                    f"WHERE BPO_ID IN ({inlist}) AND ALIVE_FLAG='1';"
                )
                sql.append(stmt_main)

                # 更新明细表 tphct02
                stmt_detail = (
                    f"UPDATE tphct02 SET APPR_STATE='{new_state}', "
                    f"OPS_REMARK='{ops_remark or ''}' "
                    f"WHERE BPO_ID IN ({inlist}) AND ALIVE_FLAG='1';"
                )
                sql.append(stmt_detail)

                sql_logger.info(
                    f"[SQL合并] 生成IN查询: {len(chunk)}条合同, 状态={new_state}")
    else:
        # 未启用合并策略: 逐条生成SQL
        sql_logger.info("[SQL合并] 合并策略已禁用, 逐条生成SQL")
        for r in records:
            new_state = r.get('new_appr_state')
            if not new_state:
                continue

            # 更新主表
            stmt_main = (
                f"UPDATE tphct01 SET APPR_STATE='{new_state}', "
                f"OPS_REMARK='{ops_remark or ''}' "
                f"WHERE BPO_ID='{r['bpo_id']}' AND ALIVE_FLAG='1';"
            )
            sql.append(stmt_main)

            # 更新明细表
            stmt_detail = (
                f"UPDATE tphct02 SET APPR_STATE='{new_state}', "
                f"OPS_REMARK='{ops_remark or ''}' "
                f"WHERE BPO_ID='{r['bpo_id']}' AND ALIVE_FLAG='1';"
            )
            sql.append(stmt_detail)

    # 生成回退语句
    sql.append("")
    sql.append("2、回退语句")

    if cfg.get('MERGE_MODULES', {}).get('appr_state', True):
        # 按原状态分组
        def rb_key(r):
            return r.get('orig_appr_state')

        groups = merge_by_key(records, rb_key)
        sql_logger.info(f"[SQL合并] 回退语句分组完成, 组数={len(groups)}")

        for orig_state, recs in groups.items():
            if not orig_state:
                continue

            bpo_ids = [r['bpo_id'] for r in recs if r.get('bpo_id')]

            for chunk in chunk_list(bpo_ids, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
                inlist = format_in(chunk)

                # 回退主表
                rb_main = (
                    f"UPDATE tphct01 SET APPR_STATE='{orig_state}', "
                    f"OPS_REMARK='' "
                    f"WHERE BPO_ID IN ({inlist}) AND ALIVE_FLAG='1';"
                )
                sql.append(rb_main)

                # 回退明细表
                rb_detail = (
                    f"UPDATE tphct02 SET APPR_STATE='{orig_state}', "
                    f"OPS_REMARK='' "
                    f"WHERE BPO_ID IN ({inlist}) AND ALIVE_FLAG='1';"
                )
                sql.append(rb_detail)
    else:
        # 逐条生成回退语句
        for r in records:
            orig_state = r.get('orig_appr_state')
            if not orig_state:
                continue

            # 回退主表
            rb_main = (
                f"UPDATE tphct01 SET APPR_STATE='{orig_state}', "
                f"OPS_REMARK='' "
                f"WHERE BPO_ID='{r['bpo_id']}' AND ALIVE_FLAG='1';"
            )
            sql.append(rb_main)

            # 回退明细表
            rb_detail = (
                f"UPDATE tphct02 SET APPR_STATE='{orig_state}', "
                f"OPS_REMARK='' "
                f"WHERE BPO_ID='{r['bpo_id']}' AND ALIVE_FLAG='1';"
            )
            sql.append(rb_detail)

    # 数据库信息
    sql.append("")
    sql.append("3、数据库")
    sql.append("ip：192.168.11.71")
    sql.append("库名：cnnc_ph")

    return sql


def download_appr_state_template(request):
    """下载合同状态修改Excel模板"""
    if not OPENPYXL_AVAILABLE:
        return FileResponse(
            io.BytesIO(b"openpyxl not installed"),
            as_attachment=True,
            filename="error.txt"
        )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "合同状态修改"

    # 设置表头
    headers = ["合同ID", "新合同状态", "原合同状态"]
    header_fill = PatternFill(start_color="4472C4",
                              end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')

    # 设置列宽
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 25

    # 添加示例数据
    example_data = [
        ["CNSC-25-00187", "DRAFT", "ACTIVE"],
        ["CNSC-25-00186", "DRAFT", "ACTIVE"],
    ]

    for row_num, row_data in enumerate(example_data, 2):
        for col_num, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col_num)
            cell.value = value
            cell.alignment = Alignment(horizontal='left', vertical='center')

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)

    response = FileResponse(bio, as_attachment=True,
                            filename='appr_state_template.xlsx')
    response['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    return response


def appr_state_change_view(request):
    """合同状态修改视图"""
    saved_file = None
    validation_failure = None
    
    # 加载合同状态选项传递给模板
    try:
        from work_tools.dropdown_utils import get_dropdown_options
        contract_status_options = get_dropdown_options('contract_status', include_empty=True, empty_label='请选择合同状态')
    except Exception:
        # 回退到硬编码选项
        contract_status_options = APPR_STATE_CHOICES
    
    if request.method == 'POST':
        form = ApprStateChangeForm(request.POST, request.FILES)
        if form.is_valid():
            try:
                cd = form.cleaned_data
                ops_remark = parse_ops_remark(cd.get('ops_remark', ''))
                dynamic_id = cd.get('dynamic_id', '')

                # 记录视图输入
                log_view_input('appr_state_change', cd)

                logger.info(
                    f"[合同状态] 开始处理, dynamic_id={dynamic_id}, ops_remark={ops_remark}")

                # 解析数据
                excel_file = cd.get('excel_file')
                if excel_file:
                    records = parse_appr_state_excel(excel_file)
                    logger.info(f"[合同状态] Excel模式, 记录数={len(records)}")
                    
                    # 执行数据校验
                    validation_result = validate_appr_state_records(records)
                    
                    if not validation_result['valid']:
                        # 校验失败，生成失败文件
                        temp_file = generate_validation_failure_excel(
                            excel_file,
                            validation_result,
                            'appr_state'
                        )
                        
                        if temp_file:
                            validation_failure = {
                                'total': validation_result['total'],
                                'passed': validation_result['passed'],
                                'failed': validation_result['failed'],
                                'filename': get_temp_filename_from_path(temp_file)
                            }
                        
                        return render(request, 'modules/contract/appr_state_change.html', {
                            'form': form,
                            'validation_failure': validation_failure,
                            'active_menu': 'appr_state_change',
                            'sidebar_groups': get_sidebar_groups(),
                            'contract_status_options': contract_status_options,
                        })
                else:
                    # 单条记录
                    rec = {
                        'bpo_id': cd.get('bpo_id'),
                        'new_appr_state': cd.get('new_appr_state'),
                    }
                    if cd.get('orig_appr_state'):
                        rec['orig_appr_state'] = cd.get('orig_appr_state')
                    records = [rec]
                    logger.info(
                        f"[合同状态] 单条模式, bpo_id={rec['bpo_id']}, new_state={rec['new_appr_state']}")

                # 生成SQL
                sql_lines = generate_appr_state_sql_bulk(records, ops_remark)
                logger.info(f"[合同状态] SQL生成完成, 行数={len(sql_lines)}")

                # 保存SQL文件
                sql_content = '\n'.join(sql_lines)
                saved_file = save_sql_file(
                    sql_content, prefix='合同状态修改', dynamic_id=dynamic_id)
                logger.info(f"[合同状态] SQL文件已保存: {saved_file}")

            except Exception as e:
                logger.error(f"[合同状态] 处理失败: {e}", exc_info=True)
                form.add_error(None, f"生成SQL失败: {str(e)}")
    else:
        form = ApprStateChangeForm()

    return render(request, 'modules/contract/appr_state_change.html', {
        'form': form,
        'saved_file': saved_file,
        'validation_failure': validation_failure,
        'active_menu': 'appr_state_change',
        'sidebar_groups': get_sidebar_groups(),
        'contract_status_options': contract_status_options,
    })


__all__ = ['appr_state_change_view', 'download_appr_state_template']
