"""合同预算修改模块"""
import os
import io
import openpyxl
import logging
from datetime import datetime
from django.shortcuts import render
from django.http import FileResponse, HttpResponse
from django.conf import settings

from work_tools.forms import ContractBudgetUpdateForm
from work_tools.navigation import get_sidebar_groups
from work_tools.config import get_config
from work_tools.sql_merge import chunk_list, format_in, merge_by_key
from work_tools.views.base import save_sql_file, parse_ops_remark
from work_tools.validation_utils import (
    validate_required,
    generate_validation_failure_excel,
    get_temp_filename_from_path,
)

logger = logging.getLogger('work_tools.view')


def validate_budget_records(records):
    """
    校验预算修改记录
    
    Args:
        records: 记录列表
        
    Returns:
        校验结果字典
    """
    results = []
    passed_count = 0
    failed_count = 0
    
    for idx, record in enumerate(records):
        row_number = idx + 2  # 跳过表头
        errors = []
        
        # 必填项校验：询价单标段编号必填
        error = validate_required(record.get('section_no'), '询价单标段编号')
        if error:
            errors.append(error)
        
        # 记录结果
        if errors:
            results.append({
                'row_number': row_number,
                'valid': False,
                'errors': errors
            })
            failed_count += 1
        else:
            results.append({
                'row_number': row_number,
                'valid': True,
                'errors': []
            })
            passed_count += 1
    
    return {
        'valid': failed_count == 0,
        'total': len(records),
        'passed': passed_count,
        'failed': failed_count,
        'results': results
    }


def parse_budget_excel(file):
    """解析预算修改Excel文件"""
    wb = openpyxl.load_workbook(file)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    idx = {str(h): i for i, h in enumerate(headers)}

    def pick(alts):
        for a in alts:
            if a in idx:
                return idx[a]
        return None

    bpo_idx = pick(['BPO_ID', 'bpo_id', '合同号'])
    sec_idx = pick(['SECTION_NO', 'section_no', '标段编号'])
    new_idx = pick(['BUDGET_PRICE', 'new_budget', '预算单价', '新预算'])
    orig_idx = pick(['orig_budget', '原预算', '原预算单价'])

    if (sec_idx is None and bpo_idx is None) or new_idx is None:
        raise ValueError("Excel 缺少必要列：标段编号或合同号，以及新预算")

    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        sec = row[sec_idx] if sec_idx is not None else None
        newv = row[new_idx]
        if ((sec is None) and (bpo_idx is None or row[bpo_idx] is None)) or newv is None:
            continue
        rec = {
            'new_budget': str(newv).strip(),
        }
        if sec is not None:
            rec['section_no'] = str(sec).strip()
        if bpo_idx is not None and row[bpo_idx] is not None:
            rec['contract_bpo_id'] = str(row[bpo_idx]).strip()
        if orig_idx is not None and row[orig_idx] is not None:
            rec['orig_budget'] = str(row[orig_idx]).strip()
        records.append(rec)

    if not records:
        raise ValueError("Excel 中没有有效的行数据")
    return records


def generate_budget_sql_bulk(records, ops_remark=None):
    """生成预算修改SQL"""
    cfg = get_config()
    sql = []
    sql.append("1、执行语句")
    
    logger.info(f"[SQL合并] 开始生成预算SQL, 记录数={len(records)}, 配置启用={cfg.get('MERGE_MODULES', {}).get('budget', True)}")
    
    if cfg.get('MERGE_MODULES', {}).get('budget', True):
        # SQL合并模式 - PR明细
        def key_fn_part(r):
            return r.get('new_budget')
        
        part_records = [r for r in records if r.get('section_no')]
        if part_records:
            groups_part = merge_by_key(part_records, key_fn_part)
            logger.info(f"[SQL合并] PR预算分组完成, 组数={len(groups_part)}")
            
            sql.append("-- PR 明细：tprxj07 / tprxj10 / tprnq02")
            for budget, recs in groups_part.items():
                if budget is None:
                    continue
                sections = [r['section_no'] for r in recs if r.get('section_no')]
                for chunk in chunk_list(sections, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
                    inlist = format_in(chunk)
                    sql.append(
                        f"update tprxj07 set BUDGET_PRICE='{budget}', OPS_REMARK='{ops_remark or ''}' where SECTION_NO IN ({inlist}) and ALIVE_FLAG='1';"
                    )
                    sql.append(
                        f"update tprxj10 set BUDGET_PRICE='{budget}', OPS_REMARK='{ops_remark or ''}' where SECTION_NO IN ({inlist}) and ALIVE_FLAG='1';"
                    )
                    sql.append(
                        f"update tprnq02 set BUDGET_PRICE='{budget}', OPS_REMARK='{ops_remark or ''}' where SECTION_NO IN ({inlist}) and ALIVE_FLAG='1';"
                    )
        
        # SQL合并模式 - PH主表
        contract_records = [r for r in records if r.get('contract_bpo_id')]
        if contract_records:
            groups_contract = merge_by_key(contract_records, key_fn_part)
            logger.info(f"[SQL合并] PH预算分组完成, 组数={len(groups_contract)}")
            
            sql.append("-- PH 主表：修改合同预算 TPHCT01")
            for budget, recs in groups_contract.items():
                if budget is None:
                    continue
                bpo_ids = [r['contract_bpo_id'] for r in recs if r.get('contract_bpo_id')]
                for chunk in chunk_list(bpo_ids, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
                    inlist = format_in(chunk)
                    sql.append(
                        f"update TPHCT01 set BUDGETED_AMOUNT='{budget}', OPS_REMARK='{ops_remark or ''}' WHERE bpo_id IN ({inlist}) and ALIVE_FLAG='1';"
                    )
    else:
        # 非合并模式
        part_exec = []
        contract_exec = []

        for r in records:
            if r.get('section_no') and r.get('new_budget') is not None:
                cond = f"SECTION_NO='{r['section_no']}' and ALIVE_FLAG='1'"
                if r.get('supplier_id'):
                    cond += f" AND SUPPLIER_ID='{r['supplier_id']}'"
                part_exec.append(
                    f"update tprxj07 set BUDGET_PRICE='{r['new_budget']}', OPS_REMARK='{ops_remark or ''}' where {cond};"
                )
                part_exec.append(
                    f"update tprxj10 set BUDGET_PRICE='{r['new_budget']}', OPS_REMARK='{ops_remark or ''}' where {cond};"
                )
                part_exec.append(
                    f"update tprnq02 set BUDGET_PRICE='{r['new_budget']}', OPS_REMARK='{ops_remark or ''}' where {cond};"
                )
            if r.get('contract_bpo_id') and r.get('new_budget') is not None:
                contract_exec.append(
                    f"update TPHCT01 set BUDGETED_AMOUNT='{r['new_budget']}', OPS_REMARK='{ops_remark or ''}' WHERE bpo_id='{r['contract_bpo_id']}' and ALIVE_FLAG='1';"
                )

        if part_exec:
            sql.append("-- PR 明细：tprxj07 / tprxj10 / tprnq02")
            sql.extend(part_exec)
        if contract_exec:
            sql.append("-- PH 主表：修改合同预算 TPHCT01")
            sql.extend(contract_exec)

    sql.append("")
    sql.append("2、回退语句")
    
    if cfg.get('MERGE_MODULES', {}).get('budget', True):
        # SQL合并模式 - PR明细回退
        def rb_key_part(r):
            return r.get('orig_budget')
        
        part_records_rb = [r for r in records if r.get('section_no') and r.get('orig_budget')]
        if part_records_rb:
            groups_part_rb = merge_by_key(part_records_rb, rb_key_part)
            sql.append("-- PR 回退：tprxj07 / tprxj10 / tprnq02")
            for orig_budget, recs in groups_part_rb.items():
                if orig_budget is None:
                    continue
                sections = [r['section_no'] for r in recs if r.get('section_no')]
                for chunk in chunk_list(sections, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
                    inlist = format_in(chunk)
                    sql.append(
                        f"update tprxj07 set BUDGET_PRICE='{orig_budget}', OPS_REMARK='' where SECTION_NO IN ({inlist}) and ALIVE_FLAG='1';"
                    )
                    sql.append(
                        f"update tprxj10 set BUDGET_PRICE='{orig_budget}', OPS_REMARK='' where SECTION_NO IN ({inlist}) and ALIVE_FLAG='1';"
                    )
                    sql.append(
                        f"update tprnq02 set BUDGET_PRICE='{orig_budget}', OPS_REMARK='' where SECTION_NO IN ({inlist}) and ALIVE_FLAG='1';"
                    )
        
        # SQL合并模式 - PH主表回退
        contract_records_rb = [r for r in records if r.get('contract_bpo_id') and r.get('orig_budget')]
        if contract_records_rb:
            groups_contract_rb = merge_by_key(contract_records_rb, rb_key_part)
            sql.append("-- PH 回退：修改合同预算 TPHCT01")
            for orig_budget, recs in groups_contract_rb.items():
                if orig_budget is None:
                    continue
                bpo_ids = [r['contract_bpo_id'] for r in recs if r.get('contract_bpo_id')]
                for chunk in chunk_list(bpo_ids, int(cfg.get('MERGE_MAX_IN_SIZE', 500))):
                    inlist = format_in(chunk)
                    sql.append(
                        f"update TPHCT01 set BUDGETED_AMOUNT='{orig_budget}', OPS_REMARK='' WHERE bpo_id IN ({inlist}) and ALIVE_FLAG='1';"
                    )
    else:
        # 非合并模式
        part_rb = []
        contract_rb = []

        for r in records:
            if r.get('orig_budget') is None:
                continue
            if r.get('section_no'):
                cond_rb = f"SECTION_NO='{r['section_no']}' and ALIVE_FLAG='1'"
                if r.get('supplier_id'):
                    cond_rb += f" AND SUPPLIER_ID='{r['supplier_id']}'"
                part_rb.append(
                    f"update tprxj07 set BUDGET_PRICE='{r['orig_budget']}', OPS_REMARK='' where {cond_rb};"
                )
                part_rb.append(
                    f"update tprxj10 set BUDGET_PRICE='{r['orig_budget']}', OPS_REMARK='' where {cond_rb};"
                )
                part_rb.append(
                    f"update tprnq02 set BUDGET_PRICE='{r['orig_budget']}', OPS_REMARK='' where {cond_rb};"
                )
            if r.get('contract_bpo_id'):
                contract_rb.append(
                    f"update TPHCT01 set BUDGETED_AMOUNT='{r['orig_budget']}', OPS_REMARK='' WHERE bpo_id='{r['contract_bpo_id']}' and ALIVE_FLAG='1';"
                )

        if part_rb:
            sql.append("-- PR 回退：tprxj07 / tprxj10 / tprnq02")
            sql.extend(part_rb)
        if contract_rb:
            sql.append("-- PH 回退：修改合同预算 TPHCT01")
            sql.extend(contract_rb)

    sql.append("")
    sql.append("3、数据库")
    sql.append("ip:192.168.11.69")
    sql.append("库名：cnnc_pr")
    sql.append("")
    sql.append("ip:192.168.11.71")
    sql.append("库名：cnnc_ph")
    return "\n".join(sql)


def contract_budget_update_view(request):
    """合同预算修改视图"""
    saved_file = None
    validation_failure = None
    
    if request.method == 'POST':
        form = ContractBudgetUpdateForm(request.POST, request.FILES)
        if form.is_valid():
            try:
                cd = form.cleaned_data
                ops_remark = parse_ops_remark(cd.get('ops_remark', ''))

                if cd.get('excel_file'):
                    records = parse_budget_excel(cd['excel_file'])
                    
                    # 执行数据校验
                    validation_result = validate_budget_records(records)
                    
                    if not validation_result['valid']:
                        # 校验失败，生成失败文件
                        temp_file = generate_validation_failure_excel(
                            cd['excel_file'],
                            validation_result,
                            'contract_budget'
                        )
                        
                        if temp_file:
                            validation_failure = {
                                'total': validation_result['total'],
                                'passed': validation_result['passed'],
                                'failed': validation_result['failed'],
                                'filename': get_temp_filename_from_path(temp_file)
                            }
                            logger.info(f"校验失败: 总行数={validation_failure['total']}, "
                                      f"失败行数={validation_failure['failed']}")
                        
                        return render(request, 'modules/contract/contract_budget_form.html', {
                            'form': form,
                            'validation_failure': validation_failure,
                            'active_menu': 'contract_budget',
                            'sidebar_groups': get_sidebar_groups(),
                        })
                else:
                    rec = {
                        'contract_bpo_id': cd.get('contract_bpo_id'),
                        'section_no': cd.get('section_no'),
                        'supplier_id': cd.get('supplier_id'),
                        'new_budget': cd.get('new_budget'),
                        'orig_budget': cd.get('orig_budget'),
                    }
                    records = [rec]

                sql_content = generate_budget_sql_bulk(records, ops_remark)

                # 保存SQL到固定目录
                saved_file = save_sql_file(
                    sql_content, '修改合同预算', cd.get('dynamic_id'))

                # 保存会话数据
                session_data = {}
                for k, v in cd.items():
                    if k != 'excel_file':
                        if hasattr(v, 'to_eng_string'):
                            session_data[k] = v.to_eng_string()
                        else:
                            session_data[k] = v
                request.session['contract_budget_last'] = session_data
            except Exception as e:
                logger.error(f"[合同预算修改] 处理失败: {e}", exc_info=True)
                form.add_error(None, f"生成SQL失败: {str(e)}")
    else:
        if request.GET.get('clear'):
            request.session.pop('contract_budget_last', None)
            form = ContractBudgetUpdateForm()
        else:
            initial = request.session.get('contract_budget_last')
            form = ContractBudgetUpdateForm(initial=initial)

    return render(request, 'modules/contract/contract_budget_form.html', {
        'form': form,
        'saved_file': saved_file,
        'validation_failure': validation_failure,
        'active_menu': 'contract_budget',
        'sidebar_groups': get_sidebar_groups(),
    })


def download_budget_template(request):
    """下载预算修改模板"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '模板'
    ws.append(['合同号', '标段编号', '供应商ID(可选)', '新预算', '原预算'])
    ws.append(['BTRL-25-00046', 'YZN0-XJD-25-03316',
              'G00361071', '6100', '4881'])
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    response = FileResponse(bio, as_attachment=True,
                            filename='contract_budget_template.xlsx')
    response['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    return response


# 别名兼容
contract_budget_view = contract_budget_update_view

__all__ = [
    'contract_budget_view',
    'contract_budget_update_view',
    'download_budget_template',
]
